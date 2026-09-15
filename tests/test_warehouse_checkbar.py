from io import BytesIO
import pytest
from openpyxl import load_workbook
from test_warehouse_fulfillment import case
from app import warehouse_assistant_service as service
from app import warehouse_checkbar as checkbar


def activate(case, supplier='supplier', warehouse='karaj', brand='brand'):
    with service.warehouse_connection(case.settings) as conn:
        conn.execute('''INSERT INTO warehouse_supply_scope(warehouse_code,warehouse_name,supplier,
            brand,enabled,created_at,updated_at) VALUES(?,?,?,?,1,'test','test')
            ON CONFLICT(warehouse_code,supplier,brand) DO UPDATE SET enabled=1''',
            (warehouse,warehouse,supplier,brand))


@pytest.fixture(autouse=True)
def active_scope(case):
    activate(case)


def payload(preview, **changes):
    return dict(warehouse='karaj', supplier='supplier', order_ids=[1],
                expected_token=preview['expected_token'], request_id='request-one', metadata={'reference_no':'2222'},
                lines=[dict(preorder_id=1, product_code='00123', cartons=None, units=None)], **changes)


def test_supplier_without_orders_can_issue_standalone_checkbar(case):
    activate(case,'no prior orders')
    with service.warehouse_connection(case.settings) as conn:
        conn.execute('''INSERT INTO warehouse_snapshot_items(snapshot_id,source_row,warehouse_code,
            warehouse_name,product_code,product_name,conversion_rate,manufacturer,manufacturer_product_code,barcode)
            VALUES(1,20,'karaj','Karaj','NEW01','new supplier item',6,'no prior orders','0009','000123')''')
    context=checkbar.context(case.settings,'karaj')
    assert 'no prior orders' in context['suppliers']
    assert checkbar.context(case.settings,'karaj','no prior orders')['orders']==[]
    preview=checkbar.prepare(case.settings,'karaj','no prior orders',[])
    assert preview['lines']==[] and preview['order_ids']==[]
    body=dict(warehouse='karaj',supplier='no prior orders',order_ids=[],
              expected_token=preview['expected_token'],request_id='standalone',metadata={'reference_no':'2222'},
              lines=[dict(product_code='NEW01',cartons=2,units=1)])
    saved=checkbar.issue(case.settings,'tester',body)
    assert saved['lines'][0]['actual_qty']==13 and saved['order_ids']==[]
    assert saved['lines'][0]['remaining_qty'] is None
    assert checkbar.issue(case.settings,'tester',body)['id']==saved['id']
    sheet=load_workbook(BytesIO(checkbar.workbook(saved)))['چک بار']
    assert sheet['B5'].value=='NEW01' and sheet['Q5'].value is None
    with pytest.raises(service.WarehouseAssistantError):checkbar.prepare(case.settings,'tehran','no prior orders',[])
    with pytest.raises(service.WarehouseAssistantError):checkbar.prepare(case.settings,'karaj','invented',[])


def test_disabled_automatic_ordering_does_not_hide_checkbar_supplier(case):
    case.change('UPDATE warehouse_supplier_auto_order_settings SET enabled=0')
    assert 'supplier' in checkbar.context(case.settings,'karaj')['suppliers']
    assert checkbar.Selection(warehouse='karaj',supplier='supplier').order_ids==[]


def test_level3_survives_checkbar_catalog_source_document_and_receipt(case):
    from app.warehouse_fulfillment import fulfillment_detail
    case.change("UPDATE warehouse_snapshot_items SET group_level3='شوینده'")
    case.send()
    preview=checkbar.prepare(case.settings,'karaj','supplier',[1])
    assert preview['catalog'][0]['group_level3']=='شوینده'
    assert preview['lines'][0]['group_level3']=='شوینده'
    saved=checkbar.issue(case.settings,'tester',payload(preview))
    assert saved['lines'][0]['group_level3']=='شوینده'
    with service.warehouse_connection(case.settings) as conn:
        assert fulfillment_detail(conn,1)['lines'][0]['group_level3']=='شوینده'
        assert conn.execute('SELECT COUNT(*) FROM warehouse_fulfillment_receipts').fetchone()[0]==0


def test_inactive_supply_is_not_restored_by_old_orders_or_auto_settings(case):
    case.send()
    preview=checkbar.prepare(case.settings,'karaj','supplier',[1])
    case.change("UPDATE warehouse_supply_scope SET enabled=0 WHERE supplier='supplier'")
    assert checkbar.context(case.settings,'karaj')['suppliers']==[]
    with pytest.raises(service.WarehouseAssistantError,match='فعال'):
        checkbar.issue(case.settings,'tester',payload(preview))
    activate(case,warehouse='tehran')
    assert checkbar.context(case.settings,'karaj')['suppliers']==[]


def partial(case, quantity=5):
    case.change(f"""INSERT INTO warehouse_fulfillment_receipts
        (preorder_id,product_code,quantity,snapshot_id,reference,recorded_by,recorded_at)
        VALUES(1,'00123',{quantity},1,'test','tester','now')""")


def test_partial_receipts_and_unfilled_actual_are_distinct(case):
    case.send();partial(case)
    preview=checkbar.prepare(case.settings,'karaj','supplier',[1])
    assert preview['lines'][0]['remaining_qty']==19
    saved=checkbar.issue(case.settings,'tester',payload(preview))
    assert saved['lines'][0]['actual_qty'] is None
    sheet=load_workbook(BytesIO(checkbar.workbook(saved)))['چک بار']
    assert sheet['I5'].value is None and sheet['J5'].value is None
    assert sheet['K5'].value=='=IF(COUNT(I5:J5)=0,"",I5*F5+J5)'
    assert 'tax condition' not in [cell.value for cell in sheet[4]]
    assert sheet['R5'].value=='=IF(OR(K5="",Q5=""),"",K5-Q5)'
    assert sheet.auto_filter.ref == f'A4:R{sheet.max_row}'
    assert sheet['P5'].value==5 and sheet['Q5'].value==19


def test_remaining_quantity_change_does_not_block_initial_checkbar_save(case):
    case.send();preview=checkbar.prepare(case.settings,'karaj','supplier',[1]);partial(case)
    saved=checkbar.issue(case.settings,'tester',payload(preview))
    assert saved['lines'][0]['remaining_qty']==19
    assert saved['lines'][0]['actual_qty'] is None
    assert len(checkbar.documents(case.settings))==1


def test_excess_actual_is_inspection_only_and_excel_text_is_not_a_formula(case):
    case.change("UPDATE warehouse_automatic_preorder_lines SET product_name='=1+1' WHERE preorder_id=1")
    case.send()
    preview=checkbar.prepare(case.settings,'karaj','supplier',[1])
    body=payload(preview);body['lines'][0]['cartons']=3
    saved=checkbar.issue(case.settings,'tester',body)
    assert saved['lines'][0]['actual_qty']==36
    assert checkbar.prepare(case.settings,'karaj','supplier',[1])['lines'][0]['remaining_qty']==24
    sheet=load_workbook(BytesIO(checkbar.workbook(saved)))['چک بار']
    assert sheet['E5'].data_type=='s' and '1+1' in sheet['E5'].value
    assert sheet['I5'].value==3 and sheet['Q5'].value==24


def test_completed_or_duplicate_or_wrong_scope_cannot_be_selected(case):
    case.send()
    for warehouse,supplier,ids in [('tehran','supplier',[1]),('karaj','other',[1]),('karaj','supplier',[1,1]),('karaj','supplier',[999])]:
        with pytest.raises(service.WarehouseAssistantError):checkbar.prepare(case.settings,warehouse,supplier,ids)
    partial(case,24)
    with pytest.raises(service.WarehouseAssistantError):checkbar.prepare(case.settings,'karaj','supplier',[1])


def test_idempotent_retry_and_immutable_history(case):
    case.send();preview=checkbar.prepare(case.settings,'karaj','supplier',[1]);body=payload(preview)
    saved=checkbar.issue(case.settings,'tester',body)
    partial(case)
    assert checkbar.issue(case.settings,'tester',body)==saved
    assert len(checkbar.documents(case.settings))==1
    assert checkbar.get_document(case.settings,saved['id'])==saved
    body['metadata']={'driver':'changed','reference_no':'2222'}
    with pytest.raises(service.WarehouseAssistantError):checkbar.issue(case.settings,'tester',body)


def test_manual_items_are_supplier_and_warehouse_scoped(case):
    with service.warehouse_connection(case.settings) as conn:
        for code,warehouse,supplier in [('002','karaj','supplier'),('003','tehran','supplier'),('004','karaj','other')]:
            conn.execute('''INSERT INTO warehouse_snapshot_items(snapshot_id,source_row,warehouse_code,
                warehouse_name,product_code,product_name,conversion_rate,manufacturer,manufacturer_product_code,barcode)
                VALUES(1,2,?,?,?,'manual',6,?,'00009','00012345678901234567')''',(warehouse,warehouse,code,supplier))
    case.send();preview=checkbar.prepare(case.settings,'karaj','supplier',[1])
    body=payload(preview)
    for code in ['003','004','missing','00123']:
        body['lines']=[dict(preorder_id=None,product_code=code,cartons=1)]
        with pytest.raises(service.WarehouseAssistantError):checkbar.issue(case.settings,'tester',body)
    body['lines']=[dict(preorder_id=None,product_code='002',cartons=1)]
    saved=checkbar.issue(case.settings,'tester',body)
    assert saved['lines'][0]['actual_qty']==6 and saved['lines'][0]['remaining_qty'] is None
    sheet=load_workbook(BytesIO(checkbar.workbook(saved)))['چک بار']
    assert sheet['C5'].value=='00009' and sheet['D5'].value=='00012345678901234567'
    assert sheet['D5'].data_type=='s'


def test_duplicate_line_invalid_numbers_and_unknown_metadata_rejected(case):
    from pydantic import ValidationError
    case.send();preview=checkbar.prepare(case.settings,'karaj','supplier',[1])
    body=payload(preview);body['lines']*=2
    with pytest.raises(service.WarehouseAssistantError):checkbar.issue(case.settings,'tester',body)
    for field,value in [('cartons',.5),('units',-1),('units',float('inf')),('tax',101),('consumer_price_new',float('nan'))]:
        body=payload(preview);body['lines'][0][field]=value
        with pytest.raises(ValidationError):checkbar.issue(case.settings,'tester',body)
    body=payload(preview);body['metadata']={'unexpected':'bad'}
    with pytest.raises(service.WarehouseAssistantError):checkbar.issue(case.settings,'tester',body)
    assert checkbar.documents(case.settings)==[]


def test_two_orders_same_product_keep_distinct_source_rows(case):
    case.send()
    with service.warehouse_connection(case.settings) as conn:
        for table,updates in [('warehouse_automatic_preorders',dict(id=2,preorder_number='AUTO-TWO',generation_key='other')),
                              ('warehouse_automatic_preorder_lines',dict(preorder_id=2)),
                              ('warehouse_email_attempts',dict(preorder_id=2,message_id='fixture-two'))]:
            row=dict(conn.execute(f'SELECT * FROM {table} LIMIT 1').fetchone());row.pop('id');row.update(updates)
            conn.execute(f"INSERT INTO {table} ({','.join(row)}) VALUES ({','.join('?' for _ in row)})",list(row.values()))
    preview=checkbar.prepare(case.settings,'karaj','supplier',[1,2])
    assert [(l['preorder_id'],l['product_code']) for l in preview['lines']]==[(1,'00123'),(2,'00123')]
    body=payload(preview);body['order_ids']=[1,2];body['lines']=[dict(preorder_id=i,product_code='00123',cartons=1) for i in [1,2]]
    saved=checkbar.issue(case.settings,'tester',body)
    assert len(saved['lines'])==2


def test_routes_enforce_permissions_and_download_saved_document(case):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from unittest.mock import patch
    from app.routes import warehouse_assistant as routes
    app=FastAPI();app.state.settings=case.settings;app.include_router(routes.router)
    with TestClient(app) as client:
        assert client.get('/warehouse-assistant/api/checkbars').status_code==401
        app.dependency_overrides[routes.require_session_user]=lambda:'tester'
        with patch.object(routes,'_capabilities',return_value=set()):
            assert client.get('/warehouse-assistant/api/checkbars').status_code==403
            assert client.get('/warehouse-assistant/api/checkbars/1/document.xlsx').status_code==403
        with patch.object(routes,'_require',return_value='tester'):
            case.send()
            preview=client.post('/warehouse-assistant/api/checkbars/preview',json=dict(warehouse='karaj',supplier='supplier',order_ids=[1])).json()
            response=client.post('/warehouse-assistant/api/checkbars',json=payload(preview))
            assert response.status_code==201,response.text
            assert response.json()['receipt_recorded'] is False
            id=response.json()['document']['id']
            response=client.get(f'/warehouse-assistant/api/checkbars/{id}/document.xlsx')
            assert response.status_code==200
            assert load_workbook(BytesIO(response.content))['چک بار']['B5'].value=='00123'
            assert client.get('/warehouse-assistant/api/checkbars/context?warehouse=invalid').status_code==422


def test_only_approved_remaining_orders_in_selected_scope(case):
    initial = checkbar.context(case.settings, 'karaj', 'supplier')
    assert [o['id'] for o in initial['orders']] == [1]
    case.send()
    data = checkbar.context(case.settings, 'karaj', 'supplier')
    assert [o['id'] for o in data['orders']] == [1]
    assert data['orders'][0]['lines'][0]['remaining_qty'] == 24
    assert not checkbar.context(case.settings, 'tehran', 'supplier')['orders']
    assert not checkbar.context(case.settings, 'karaj', 'different')['orders']


def test_issued_checkbar_does_not_receive_or_change_in_transit(case):
    case.send()
    preview = checkbar.prepare(case.settings, 'karaj', 'supplier', [1])
    saved = checkbar.issue(case.settings, 'tester', dict(
        warehouse='karaj', supplier='supplier', order_ids=[1], expected_token=preview['expected_token'],
        request_id='test-one', metadata={'driver':'test','reference_no':'2222'},
        lines=[dict(preorder_id=1, product_code='00123', cartons=1, units=3)],
    ))
    assert saved['lines'][0]['actual_qty'] == 15
    assert saved['lines'][0]['remaining_qty'] == 24
    with service.warehouse_connection(case.settings) as conn:
        assert conn.execute('SELECT COUNT(*) FROM warehouse_fulfillment_receipts').fetchone()[0] == 0
    sheet = load_workbook(BytesIO(checkbar.workbook(saved)))['چک بار']
    assert sheet['B5'].value == '00123'
    assert sheet['I5'].value == 1 and sheet['J5'].value == 3
    assert sheet['N5'].value == 'AUTO-TEST'
    case.transport.assert_called_once()


def test_edit_preserves_number_source_and_original_with_retry_and_stale_guard(case):
    case.send()
    body=payload(checkbar.prepare(case.settings,'karaj','supplier',[1]))
    saved=checkbar.issue(case.settings,'tester',body)
    partial(case)
    edit=dict(expected_revision=0,request_id='edit-1',metadata={'driver':'updated','reference_no':'2222'},
              lines=[dict(preorder_id=1,product_code='00123',cartons=20,units=10)])
    updated=checkbar.edit_document(case.settings,'tester',saved['id'],edit)
    assert updated['number']==saved['number'] and updated['revision']==1
    assert updated['lines'][0]['actual_qty']==250
    assert updated['lines'][0]['remaining_qty']==24
    assert checkbar.edit_document(case.settings,'tester',saved['id'],edit)==updated
    assert checkbar.get_document(case.settings,saved['id'])==updated
    with pytest.raises(service.WarehouseAssistantError,match='تغییر'):
        checkbar.edit_document(case.settings,'tester',saved['id'],dict(edit,request_id='stale'))
    with service.warehouse_connection(case.settings) as conn:
        import json
        original=json.loads(conn.execute('SELECT document_json FROM warehouse_checkbars WHERE id=?',(saved['id'],)).fetchone()[0])
        assert original['lines'][0]['actual_qty'] is None
        assert conn.execute('SELECT SUM(quantity) FROM warehouse_fulfillment_receipts').fetchone()[0]==5
        assert conn.execute('SELECT COUNT(*) FROM warehouse_checkbar_revisions').fetchone()[0]==1
    sheet=load_workbook(BytesIO(checkbar.workbook(updated)))['چک بار']
    assert sheet['I5'].value==20 and sheet['J5'].value==10
    case.transport.assert_called_once()


def test_soft_delete_is_idempotent_and_old_issue_cannot_resurrect(case):
    case.send();body=payload(checkbar.prepare(case.settings,'karaj','supplier',[1]))
    saved=checkbar.issue(case.settings,'tester',body)
    deletion=dict(expected_revision=0,request_id='delete-1')
    result=checkbar.delete_document(case.settings,'tester',saved['id'],deletion)
    assert result['deleted'] is True
    assert checkbar.delete_document(case.settings,'tester',saved['id'],deletion)==result
    assert checkbar.documents(case.settings)==[]
    with pytest.raises(service.WarehouseAssistantError):checkbar.get_document(case.settings,saved['id'])
    with pytest.raises(service.WarehouseAssistantError):checkbar.issue(case.settings,'tester',body)
    with service.warehouse_connection(case.settings) as conn:
        assert conn.execute('SELECT COUNT(*) FROM warehouse_checkbars').fetchone()[0]==1
        assert conn.execute('SELECT COUNT(*) FROM warehouse_fulfillment_receipts').fetchone()[0]==0


def test_api_edit_delete_and_recover_history_only_change_checkbars(case, monkeypatch):
    from fastapi import FastAPI, Request
    from fastapi.testclient import TestClient
    from app.routes import warehouse_assistant as routes
    from app.routes.dependencies import require_session_user

    # A fresh app without production lifespan, jobs, network server or live settings.
    app = FastAPI()
    app.state.settings = case.settings
    app.include_router(routes.router)
    def signed_in(request: Request):
        request.state.username = 'tester'
    app.dependency_overrides[require_session_user] = signed_in
    monkeypatch.setattr(routes, '_capabilities', lambda *_: {'warehouse.order.draft'})
    case.send()
    def business_state():
        with service.warehouse_connection(case.settings) as conn:
            names = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")
                     if r[0] not in {'warehouse_checkbars', 'warehouse_checkbar_revisions', 'sqlite_sequence'}]
            return {name: [tuple(row) for row in conn.execute(f'SELECT * FROM "{name}"')] for name in names}
    before = business_state()
    with TestClient(app) as client:
        base = '/warehouse-assistant/api/checkbars'
        body = payload(checkbar.prepare(case.settings, 'karaj', 'supplier', [1]))
        body['lines'][0]['units'] = 0
        for reference in ['', '0', 'INV-22', '2147483648']:
            invalid = dict(body, metadata={'reference_no':reference})
            assert client.post(base, json=invalid).status_code == 422
        assert checkbar.documents(case.settings) == []
        issued = client.post(base, json=body)
        assert issued.status_code == 201
        doc = issued.json()['document']; url = f"{base}/{doc['id']}"
        assert client.get(url+'/edit-context').json()['document']['revision'] == 0
        edit = dict(expected_revision=0, request_id='api-edit', metadata={'note':'revised','reference_no':'2222'},
                    lines=[dict(preorder_id=1, product_code='00123', cartons=2, units=1)])
        assert client.put(url, json=dict(edit, metadata={})).status_code == 422
        assert checkbar.get_document(case.settings, doc['id'])['revision'] == 0
        response = client.put(url, json=edit)
        assert response.status_code == 200
        changed = response.json()['document']
        assert (changed['number'], changed['revision'], changed['lines'][0]['actual_qty']) == (doc['number'], 1, 25)
        assert response.json()['varanegar_write'] is False and response.json()['receipt_recorded'] is False
        assert client.put(url, json=edit).json()['document']['revision'] == 1
        assert client.put(url, json=dict(edit, request_id='stale')).status_code == 422
        assert client.request('DELETE', url, json={'expected_revision':0,'request_id':'stale-delete'}).status_code == 422
        deletion = dict(expected_revision=1, request_id='api-delete')
        assert client.request('DELETE', url, json=deletion).status_code == 200
        assert client.request('DELETE', url, json=deletion).status_code == 200
        assert client.get(base).json()['documents'] == []
        assert client.get(base+'?include_deleted=true').json()['documents'][0]['deleted'] is True
        assert client.get(url).status_code == 422
        assert client.get(url+'/edit-context').status_code == 422
        assert client.get(url+'/document.xlsx').status_code == 422
        history = client.get(url+'/history').json()
        assert [v['operation'] for v in history['versions']] == ['issue','edit','delete']
        assert [v['document']['lines'][0]['actual_qty'] for v in history['versions']] == [0,25,25]
        for revision, expected in [(0,0),(1,1),(2,1)]:
            result = client.get(url+f'/revisions/{revision}/document.xlsx')
            assert result.status_code == 200
            assert load_workbook(BytesIO(result.content))['چک بار']['J5'].value == expected
        assert client.get(url+'/revisions/99/document.xlsx').status_code == 422
        assert client.put(url, json=dict(edit, expected_revision=2, request_id='after-delete')).status_code == 422
        assert business_state() == before
        case.transport.assert_called_once()

        # All newly exposed operations retain the independent warehouse permission.
        monkeypatch.setattr(routes, '_capabilities', lambda *_: set())
        for method, path, data in [('GET',url+'/history',None),('GET',url+'/edit-context',None),
                                   ('GET',url+'/revisions/0/document.xlsx',None),
                                   ('PUT',url,edit),('DELETE',url,deletion)]:
            assert client.request(method,path,json=data).status_code == 403
        app.dependency_overrides.clear()
        assert client.get(url+'/history').status_code == 401


@pytest.mark.parametrize('reference', [None, '', ' ', '0', '-1', '1.5', 'INV-22', '2147483648', '00000000001'])
def test_required_reference_rejected_before_checkbar_write(case, reference):
    case.send()
    body = payload(checkbar.prepare(case.settings, 'karaj', 'supplier', [1]))
    body['metadata'] = {} if reference is None else {'reference_no': reference}
    with pytest.raises(service.WarehouseAssistantError, match='عطف'):
        checkbar.issue(case.settings, 'tester', body)
    assert checkbar.documents(case.settings) == []


@pytest.mark.parametrize('reference,expected', [(' ۲۲۲۲ ', '2222'), ('٢٢٢٢', '2222'), ('1', '1'), ('2147483647', '2147483647')])
def test_reference_normalization_persists_and_retries(case, reference, expected):
    case.send()
    body = payload(checkbar.prepare(case.settings, 'karaj', 'supplier', [1]))
    body['metadata']['reference_no'] = reference
    saved = checkbar.issue(case.settings, 'tester', body)
    assert saved['metadata']['reference_no'] == expected
    assert checkbar.issue(case.settings, 'tester', body)['id'] == saved['id']
    body['metadata']['reference_no'] = expected
    assert checkbar.issue(case.settings, 'tester', body)['id'] == saved['id']
    sheet = load_workbook(BytesIO(checkbar.workbook(saved)))['چک بار']
    assert sheet['G3'].value == 'شماره سند عطف: ' + expected


def test_legacy_reference_can_be_added_without_changing_original_or_number(case):
    import json
    case.send()
    body = payload(checkbar.prepare(case.settings, 'karaj', 'supplier', [1]))
    body['lines'][0]['units'] = 1
    saved = checkbar.issue(case.settings, 'tester', body)
    # Simulate an existing document from before the required-reference rule, in the temporary store.
    with service.warehouse_connection(case.settings) as conn:
        row = conn.execute('SELECT document_json FROM warehouse_checkbars WHERE id=?', (saved['id'],)).fetchone()
        legacy = json.loads(row[0]); legacy['metadata'].pop('reference_no')
        conn.execute('UPDATE warehouse_checkbars SET document_json=? WHERE id=?', (json.dumps(legacy), saved['id']))
    old = checkbar.get_document(case.settings, saved['id'])
    assert 'reference_no' not in old['metadata']
    assert load_workbook(BytesIO(checkbar.workbook(old)))['چک بار']['B5'].value == '00123'
    edit = dict(expected_revision=0, request_id='add-reference', metadata={'reference_no':'٢٢٢٢'}, lines=body['lines'])
    updated = checkbar.edit_document(case.settings, 'tester', saved['id'], edit)
    assert updated['number'] == saved['number'] and updated['revision'] == 1
    assert updated['metadata']['reference_no'] == '2222'
    assert 'reference_no' not in checkbar.historical_document(case.settings, saved['id'], 0)['metadata']
    checkbar.delete_document(case.settings, 'tester', saved['id'], dict(expected_revision=1, request_id='delete'))
    assert checkbar.historical_document(case.settings, saved['id'], 1)['metadata']['reference_no'] == '2222'
    assert 'reference_no' not in checkbar.historical_document(case.settings, saved['id'], 0)['metadata']
