"""Level-three taxonomy propagation; isolated SQLite and fake ERP only."""
from contextlib import contextmanager
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import warehouse_assistant_service as service
from test_warehouse_assistant import _inventory_workbook, _FakeWarehouseCursor
from test_warehouse_fulfillment import case


@pytest.mark.parametrize('header', ['group 3', 'Group3', 'گروه سطح ۳ کالا', None])
def test_excel_optional_level3_is_not_inferred_from_brand(settings, tmp_path, header):
    book = load_workbook(BytesIO(_inventory_workbook()))
    sheet = book.active
    if header:
        sheet.cell(1, 30, header)
        sheet.cell(2, 30, 'خمیر دندان')
    path = tmp_path / 'groups.xlsx'
    book.save(path)
    service.import_inventory_snapshot(settings, path, path.name, 'tester')
    rows = service.list_inventory_information(settings)['items']
    assert len(rows) == 3
    assert {row['group_level3'] for row in rows} == {'خمیر دندان' if header else ''}


def test_group_level3_migrates_old_snapshot_without_changing_data(case):
    with service.warehouse_connection(case.settings) as conn:
        conn.execute('ALTER TABLE warehouse_snapshot_items DROP COLUMN group_level3')
    service.init_warehouse_store(case.settings)
    row = service.list_inventory_information(case.settings)['items'][0]
    assert row['group_level3'] == ''
    assert (row['product_code'], row['brand'], row['on_hand_qty']) == ('00123', 'brand', 0)


def test_group_level3_inventory_filter_search_suggestion_and_layout(case):
    case.change("UPDATE warehouse_snapshot_items SET group_level3='شوینده ۱۲'")
    for filters in ({'column_filters': {'group_level3': '12'}}, {'search': 'شوینده 12'}):
        rows = service.list_inventory_information(case.settings, **filters)['items']
        assert len(rows) == 1
        assert rows[0]['group_level3'] == 'شوینده ۱۲'
    assert service.list_inventory_information(case.settings, column_filters={'group_level3':'missing'})['items'] == []
    suggestion = service.build_suggestions(case.settings, warehouse='karaj', only_needed=False)['items'][0]
    assert suggestion['group_level3'] == 'شوینده ۱۲'
    assert len(service.build_suggestions(case.settings, warehouse='karaj', only_needed=False,
        search='شوینده 12')['items']) == 1
    assert 'group_level3' in service.INVENTORY_COLUMN_KEYS
    assert 'group_level3' in service.TABLE_PREFERENCE_COLUMNS['ordering']
    saved = service.save_table_preference(case.settings, 'tester', 'ordering', ['product_code','group_level3'])
    assert saved['visible_columns'] == ['product_code','group_level3']


def test_group_level3_auto_and_manual_order_lines_use_source_snapshot(case):
    case.change("UPDATE warehouse_snapshot_items SET group_level3='شوینده'")
    assert case.order()['lines'][0]['group_level3'] == 'شوینده'
    order = service.create_supplier_orders(case.settings, 'tester', snapshot_id=1, warehouse='karaj',
        lines=[{'product_code':'00123', 'quantity':12}])[0]
    assert order['lines'][0]['group_level3'] == 'شوینده'
    with service.warehouse_connection(case.settings) as conn:
        conn.execute("INSERT INTO warehouse_snapshots(id,source_filename,source_sheet,content_sha256,product_count,item_count,imported_by,imported_at) VALUES(2,'x','x','other',1,1,'test','2099-02-01')")
        conn.execute("INSERT INTO warehouse_snapshot_items(snapshot_id,source_row,warehouse_code,warehouse_name,product_code,product_name,conversion_rate,group_level3) VALUES(2,1,'karaj','Karaj','00123','Test',12,'تغییر یافته')")
    assert case.order()['lines'][0]['group_level3'] == 'شوینده'
    with service.warehouse_connection(case.settings) as conn:
        saved = service._order_from_row(conn, conn.execute('SELECT * FROM supplier_orders WHERE id=?',(order['id'],)).fetchone())
    assert saved['lines'][0]['group_level3'] == 'شوینده'


def test_erp_level3_uses_legacy_nested_set_and_is_persisted(settings, monkeypatch):
    queries = []
    class Cursor(_FakeWarehouseCursor):
        def execute(self, sql):
            super().execute(sql)
            if 'FRU.StockGoodsModel' in sql and 'ranked_stock' in sql:
                self.description.append(('GroupLevel3',))
                self._rows = [(*row, 'خمیر دندان') for row in self._rows]
            return self
    class Connection:
        def cursor(self):
            return Cursor(queries)
    @contextmanager
    def fake_connection(_settings):
        yield Connection()
    monkeypatch.setattr(service, 'sql_connection', fake_connection)
    configured = replace(settings, sql_server='fake-host', sql_username='fake-reader', sql_password='fake')
    service.sync_varanegar_snapshot(configured, 'tester')
    assert service.list_inventory_information(configured)['items'][0]['group_level3'] == 'خمیر دندان'
    sql = next(sql for sql in queries if 'ranked_stock' in sql)
    assert 'GNR.tblGoodsGroup AS group_level3' in sql
    assert 'group_level3.NLevel = 3' in sql
    assert 'goods_group.NLeft >= group_level3.NLeft' in sql
    assert 'goods_group.NRight <= group_level3.NRight' in sql
    assert 'NGT.ProductGroups' not in sql
