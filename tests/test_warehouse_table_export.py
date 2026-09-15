from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.routes import warehouse_assistant as routes
from app.routes.dependencies import require_session_user


@pytest.fixture
def client(tmp_path, monkeypatch):
    app = FastAPI()
    app.state.settings = SimpleNamespace(sqlite_path=tmp_path / "main.db")

    async def identity(request: Request):
        request.state.username = request.headers.get("X-Test-User", "worker")

    app.dependency_overrides[require_session_user] = identity
    monkeypatch.setattr(routes, "_capabilities", lambda request, username: {"warehouse.assistant.view"})
    app.include_router(routes.router)
    with TestClient(app) as test_client:
        yield test_client


def test_visible_table_export_returns_real_safe_xlsx(client):
    response = client.post(
        "/warehouse-assistant/api/table-export.xlsx",
        headers={"X-Warehouse-Export": "1", "Origin": "http://testserver"},
        json={
            "title": "پیش‌نمایش سفارش",
            "columns": ["کد کالا", "مقدار", "توضیح"],
            "rows": [["00123", 12.5, "=HYPERLINK(\"bad\")"], ["00200", 0, "عادی"]],
        },
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["cache-control"] == "no-store"
    sheet = load_workbook(BytesIO(response.content))["گزارش"]
    assert [cell.value for cell in sheet[1]] == ["کد کالا", "مقدار", "توضیح"]
    assert sheet.cell(2, 1).value == "00123"
    assert sheet.cell(2, 2).value == 12.5
    assert sheet.cell(2, 3).value.startswith("'")


def test_ordering_breakdown_and_daily_audit_survive_xlsx_export(client):
    summary = "روزانه ۱۰ عدد؛ ۶ روز موجود؛ فروش ۷۲ عدد؛ حذف هیجانی ۱۲ عدد؛ مبنا ۶۰ عدد"
    response = client.post(
        "/warehouse-assistant/api/table-export.xlsx",
        headers={"X-Warehouse-Export": "1"},
        json={
            "title": "پیش‌نمایش سفارش",
            "columns": [
                "قابل استفاده انبار (بدون در راه)", "بار در راه",
                "رسیده، منتظر موجودی", "مبنای سفارش (جمع)",
                "خروج روزانه و مبنای فروش",
            ],
            "rows": [[24, 30, 6, 60, summary]],
        },
    )
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content), read_only=True)["گزارش"]
    assert [cell.value for cell in sheet[1]] == [
        "قابل استفاده انبار (بدون در راه)", "بار در راه",
        "رسیده، منتظر موجودی", "مبنای سفارش (جمع)",
        "خروج روزانه و مبنای فروش",
    ]
    assert [sheet.cell(2, column).value for column in range(1, 5)] == [24, 30, 6, 60]
    assert sheet.cell(2, 5).value == summary


@pytest.mark.parametrize(
    "headers",
    [{}, {"X-Warehouse-Export": "1", "Origin": "https://untrusted.example"}],
)
def test_export_rejects_missing_header_or_foreign_origin(client, headers):
    response = client.post(
        "/warehouse-assistant/api/table-export.xlsx", headers=headers,
        json={"title": "جدول", "columns": ["ستون"], "rows": [["مقدار"]]},
    )
    assert response.status_code == 403


def test_export_rejects_ragged_or_oversized_payload(client):
    headers = {"X-Warehouse-Export": "1"}
    ragged = client.post(
        "/warehouse-assistant/api/table-export.xlsx", headers=headers,
        json={"title": "جدول", "columns": ["یک", "دو"], "rows": [["فقط یک"]]},
    )
    assert ragged.status_code == 422
    oversized = client.post(
        "/warehouse-assistant/api/table-export.xlsx", headers=headers,
        content=b"{" + b"x" * (8 * 1024 * 1024) + b"}",
    )
    assert oversized.status_code == 413


def test_every_warehouse_table_gets_dynamic_export_control():
    html = Path("app/static/warehouse-assistant.html").read_text(encoding="utf-8")
    script = Path("app/static/warehouse-table-export.js").read_text(encoding="utf-8")
    assert "warehouse-table-export.js?v=2" in html
    assert "table:not([data-excel-export-ready])" in script
    assert "new MutationObserver" in script
    assert "X-Warehouse-Export" in script
    assert "خروجی اکسل همین جدول" in script
    assert "input[type=\"number\"]" in script
    assert "data-export-value" in script
    assert "/^(عملیات|انتخاب)$/" in script
