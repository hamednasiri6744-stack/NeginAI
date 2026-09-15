from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.excel_service import build_report_workbook
from app.warehouse_order_excel import build_supplier_workbook


@pytest.mark.parametrize('coverage,urgent', [
    (0, True), (-1, True), (3.9999, True), (4, False), (4.001, False),
    (None, False), ('', False), ('unknown', False), (True, False),
    (float('nan'), False), (float('inf'), False), (float('-inf'), False),
])
def test_priority_label_and_entire_row_style(coverage, urgent):
    book = load_workbook(BytesIO(build_supplier_workbook(
        ['کد کالا', 'تعداد'], [['00123', 12]], 'test', [{'coverage_days': coverage}],
    )))
    sheet = book['گزارش']
    assert sheet['C1'].value == 'اولویت ارسال'
    assert sheet['C2'].value == ('اولویت ارسال' if urgent else None)
    assert sheet['A2'].value == '00123'
    assert sheet['B2'].value == 12
    assert sheet.auto_filter.ref == 'A1:C2'
    for cell in sheet[2]:
        assert (cell.fill.patternType == 'solid') == urgent
        if urgent:
            assert cell.fill.fgColor.rgb == '00FEE2E2'
            assert cell.font.color.rgb == '00991B1B'


def test_generic_reports_unchanged_and_supplier_inputs_not_mutated():
    columns, rows = ['Code'], [['00123']]
    build_supplier_workbook(columns, rows, 'test', [{'coverage_days': 1}])
    assert columns == ['Code'] and rows == [['00123']]
    sheet = load_workbook(BytesIO(build_report_workbook(columns, rows)))['گزارش']
    assert sheet.max_column == 1 and sheet['A2'].fill.patternType is None


def test_supplier_line_alignment_is_required():
    with pytest.raises(ValueError):
        build_supplier_workbook(['Code'], [['00123']], 'test', [])
