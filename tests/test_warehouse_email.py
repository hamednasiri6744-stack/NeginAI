"""Isolated mail delivery checks: temporary SQLite and fake SMTP only."""
import json
import smtplib
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from app import warehouse_assistant_service as service
from app import warehouse_email as mail
from app.business_time import jalali_business_date, tehran_now

REAL_SMTP_SEND = mail._smtp_send


class WarehouseEmailTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.settings = SimpleNamespace(sqlite_path=Path(self.temp.name) / 'main.db')
        self.config = Path(self.temp.name) / 'mail.json'
        self.config.write_text(json.dumps(dict(username='orders@example.com', password='test-secret',
            host='smtp.example.com', port=465, tls='implicit', delivery_enabled=True)))
        self.addCleanup(patch.stopall)
        patch.object(mail, 'CONFIG_PATH', self.config).start()
        self.transport = patch.object(mail, '_smtp_send', return_value=None).start()
        service.init_warehouse_store(self.settings)
        with service.warehouse_connection(self.settings) as conn:
            conn.execute("INSERT INTO warehouse_snapshots(id,source_filename,source_sheet,content_sha256,product_count,item_count,imported_by,imported_at) VALUES(1,'test','test','abc',1,1,'test','now')")
            conn.execute("INSERT INTO warehouse_supplier_auto_order_settings(id,warehouse_code,warehouse_name,supplier,created_at,updated_at) VALUES(1,'karaj','Karaj','supplier','now','now')")
            conn.execute("UPDATE warehouse_supplier_auto_order_settings SET contact_email='buyer@example.com' WHERE id=1")
            conn.execute("""INSERT INTO warehouse_automatic_preorders(id,preorder_number,generation_key,snapshot_id,supplier_setting_id,warehouse_code,warehouse_name,supplier,status,reorder_coverage_days,target_days,item_count,total_quantity,total_cartons,contact_email,created_by,created_at,business_date,approved_at)
                VALUES(1,'AUTO-TEST','key',1,1,'karaj','Karaj','supplier','approved',10,20,1,24,2,'buyer@example.com','test','now',?,'approved-now')""", (jalali_business_date(tehran_now()),))
            conn.execute("""INSERT INTO warehouse_automatic_preorder_lines(preorder_id,warehouse_code,warehouse_name,product_code,product_name,brand,conversion_rate,order_quantity,cartons,manufacturer_price,consumer_price,buy_price,estimated_value)
                VALUES(1,'karaj','Karaj','00123','=unsafe','brand',12,24,2,90,120,70,1680)""")

    def order(self):
        return service.get_automatic_preorder(self.settings, 1)

    def send(self, token=None):
        return mail.send_preorder_email(self.settings, 'test-user', 1, token or self.order()['email_send_token'])

    def change(self, sql):
        with service.warehouse_connection(self.settings) as conn:
            conn.execute(sql)

    def save_contact(self, email='new@example.com', mobile='09121112233'):
        return service.save_auto_order_setting(self.settings, 'contact-editor', 1,
            enabled=True, reorder_coverage_days=10, target_days=20, minimum_cartons=0,
            contact_first_name='New', contact_last_name='Buyer',
            contact_email=email, contact_mobile=mobile)

    def test_contacts_change_immediately_without_rebuilding_or_unapproving(self):
        before = self.order()
        self.save_contact()
        after = self.order()
        self.assertEqual(after['contact_email'], 'new@example.com')
        self.assertEqual(after['contact_mobile'], '09121112233')
        self.assertEqual(after['contact_full_name'], 'New Buyer')
        for key in ['status', 'approved_at', 'generation_key', 'snapshot_id', 'lines', 'total_cartons']:
            self.assertEqual(after[key], before[key])
        self.assertNotEqual(after['email_send_token'], before['email_send_token'])
        with self.assertRaises(service.WarehouseAssistantError):
            self.send(before['email_send_token'])
        self.transport.assert_not_called()

    def test_revoke_approval_restores_editing_and_requires_new_approval(self):
        before = self.order()
        after = service.transition_automatic_preorder(self.settings, 'reviewer', 1, 'revoke_approval')
        self.assertEqual(after['status'], 'awaiting_approval')
        self.assertIsNone(after['approved_at'])
        self.assertIsNone(after['approved_by'])
        self.assertEqual(after['lines'], before['lines'])
        self.assertEqual(after['edited_by'], 'reviewer')
        with self.assertRaises(service.WarehouseAssistantError):
            self.send(before['email_send_token'])
        edited = service.update_automatic_preorder_lines(self.settings, 'reviewer', 1,
            [{'product_code':'00123', 'cartons':3}])
        self.assertEqual(edited['total_cartons'], 3)
        self.transport.assert_not_called()

    def test_sent_contacts_are_frozen_and_approval_cannot_be_revoked(self):
        self.send()
        self.save_contact()
        self.assertEqual(self.order()['contact_email'], 'buyer@example.com')
        with self.assertRaises(service.WarehouseAssistantError):
            service.transition_automatic_preorder(self.settings, 'reviewer', 1, 'revoke_approval')

    def test_current_recipient_is_used_and_frozen_when_send_is_claimed(self):
        self.save_contact()
        self.send()
        self.assertEqual(self.transport.call_args.args[1]['To'], 'new@example.com')
        self.save_contact(email='later@example.com', mobile='09129998877')
        self.assertEqual(self.order()['contact_email'], 'new@example.com')
        self.assertEqual(self.order()['contact_mobile'], '09121112233')

    def test_contact_can_be_cleared_and_is_scoped_to_its_setting(self):
        self.save_contact(email='', mobile='')
        self.assertEqual(self.order()['contact_email'], '')
        self.assertEqual(self.order()['contact_mobile'], '')
        with self.assertRaises(service.WarehouseAssistantError):
            self.send()
        self.transport.assert_not_called()
        self.change("INSERT INTO warehouse_supplier_auto_order_settings(id,warehouse_code,warehouse_name,supplier,contact_email,created_at,updated_at) VALUES(2,'tehran','Tehran','supplier','other@example.com','now','now')")
        self.assertEqual(self.order()['contact_email'], '')

    def test_failed_delivery_can_change_recipient_and_revoke(self):
        self.transport.side_effect = mail.MailDeliveryFailure('failed', 'rejected')
        self.send()
        self.save_contact()
        self.assertEqual(self.order()['contact_email'], 'new@example.com')
        self.assertEqual(self.order()['email_delivery']['recipient'], 'buyer@example.com')
        revoked = service.transition_automatic_preorder(self.settings, 'reviewer', 1, 'revoke_approval')
        self.assertEqual(revoked['status'], 'awaiting_approval')
        self.assertIsNone(revoked['send_requested_at'])

    def test_inflight_and_unknown_delivery_protect_contacts_and_approval(self):
        def during_send(config, message):
            self.save_contact()
            self.assertEqual(self.order()['contact_email'], 'buyer@example.com')
            with self.assertRaises(service.WarehouseAssistantError):
                service.transition_automatic_preorder(self.settings, 'reviewer', 1, 'revoke_approval')
            raise mail.MailDeliveryFailure('unknown', 'uncertain')
        self.transport.side_effect = during_send
        self.send()
        self.assertEqual(self.order()['contact_email'], 'buyer@example.com')
        with self.assertRaises(service.WarehouseAssistantError):
            service.transition_automatic_preorder(self.settings, 'reviewer', 1, 'revoke_approval')

    def test_confirmed_order_sends_exact_excel_once(self):
        self.change("""INSERT INTO warehouse_snapshot_items(snapshot_id,source_row,warehouse_code,
            warehouse_name,product_code,product_name,conversion_rate,stock,period_out)
            VALUES(1,1,'karaj','Karaj','00123','test',12,6,120)""")
        result = self.send()
        self.assertEqual(result['send_status'], 'sent')
        self.assertEqual(self.order()['email_delivery']['recipient'], 'buyer@example.com')
        msg = self.transport.call_args.args[1]
        self.assertEqual(msg['To'], 'buyer@example.com')
        self.assertEqual(msg['From'], 'Negin Pakhsh Orders <orders@example.com>')
        attachment = list(msg.iter_attachments())[0]
        book = load_workbook(BytesIO(attachment.get_payload(decode=True)))
        sheet = book['گزارش']
        self.assertEqual(sheet.cell(2, 4).value, '00123')
        self.assertEqual(sheet.cell(2, 9).value, 2)
        self.assertEqual(sheet.cell(2, 10).value, 24)
        self.assertEqual(sheet.cell(2, 12).value, 90)
        self.assertNotEqual(sheet.cell(2, 7).data_type, 'f')
        self.assertEqual(sheet.cell(2, 16).value, 'اولویت ارسال')
        self.assertTrue(all(cell.fill.fgColor.rgb == '00FEE2E2' for cell in sheet[2]))
        downloaded = load_workbook(BytesIO(mail.preorder_workbook(self.order())))['گزارش']
        self.assertEqual(list(downloaded.values), list(sheet.values))
        self.send()
        self.assertEqual(self.transport.call_count, 1)

    def test_manual_order_requires_approval_then_sends_exact_excel_once(self):
        with service.warehouse_connection(self.settings) as conn:
            conn.execute("""INSERT INTO supplier_orders
                (id,order_number,snapshot_id,warehouse_code,warehouse_name,supplier,total_quantity,created_by,created_at)
                VALUES(2,'SUP-MANUAL',1,'karaj','Karaj','supplier',24,'test-user','now')""")
            conn.execute("""INSERT INTO supplier_order_lines
                (order_id,product_code,product_name,brand,conversion_rate,requested_quantity,
                 order_quantity,cartons,manufacturer_price,consumer_price,buy_price,estimated_value,note)
                VALUES(2,'00123','manual item','brand',12,22,24,2,90,120,70,1680,'manual shortage')""")
        draft = service.get_supplier_order(self.settings, 2, 'test-user')
        with self.assertRaises(service.WarehouseAssistantError):
            mail.send_supplier_order_email(self.settings, 'test-user', 2, draft['email_send_token'])
        approved = service.transition_supplier_order(self.settings, 'test-user', 2, 'approve')
        result = mail.send_supplier_order_email(
            self.settings, 'test-user', 2, approved['email_send_token'])
        self.assertEqual(result['send_status'], 'sent')
        msg = self.transport.call_args.args[1]
        self.assertEqual(msg['To'], 'buyer@example.com')
        sheet = load_workbook(BytesIO(list(msg.iter_attachments())[0].get_payload(decode=True)))['گزارش']
        self.assertEqual(sheet.cell(2, 4).value, '00123')
        self.assertEqual(sheet.cell(2, 10).value, 24)
        mail.send_supplier_order_email(self.settings, 'test-user', 2,
                                       service.get_supplier_order(self.settings, 2, 'test-user')['email_send_token'])
        self.assertEqual(self.transport.call_count, 1)

    def test_approval_and_today_required(self):
        for sql in ["UPDATE warehouse_automatic_preorders SET status='awaiting_approval'", "UPDATE warehouse_automatic_preorders SET status='approved', business_date='old'"]:
            self.change(sql)
            with self.assertRaises(service.WarehouseAssistantError):
                self.send()
        self.transport.assert_not_called()

    def test_confirmation_rejects_changed_recipient(self):
        token = self.order()['email_send_token']
        self.change("UPDATE warehouse_supplier_auto_order_settings SET contact_email='different@example.com'")
        with self.assertRaises(service.WarehouseAssistantError):
            self.send(token)
        self.transport.assert_not_called()

    def test_rejects_multiple_or_injected_recipients(self):
        for address in ['', 'a@example.com,b@example.com', 'a@example.com\r\nBcc: b@example.com']:
            with service.warehouse_connection(self.settings) as conn:
                conn.execute('UPDATE warehouse_supplier_auto_order_settings SET contact_email=?', (address,))
            with self.assertRaises(service.WarehouseAssistantError):
                self.send()
        self.transport.assert_not_called()

    def test_disabled_config_never_claims_or_sends(self):
        self.config.write_text('{"delivery_enabled": false}')
        with self.assertRaises(service.WarehouseAssistantError):
            self.send()
        self.assertEqual(self.order()['status'], 'approved')
        self.transport.assert_not_called()

    def test_definite_failure_is_recorded_and_manual_retry_allowed(self):
        self.transport.side_effect = mail.MailDeliveryFailure('failed', 'رد شد')
        self.assertEqual(self.send()['send_status'], 'failed')
        self.transport.side_effect = None
        self.assertEqual(self.send()['send_status'], 'sent')
        self.assertEqual(self.order()['email_delivery']['attempts'], 2)

    def test_uncertain_failure_is_not_retried(self):
        self.transport.side_effect = mail.MailDeliveryFailure('unknown', 'نیازمند بررسی')
        result = self.send()
        self.assertEqual(result['send_status'], 'unknown')
        self.assertIsNone(result['external_delivery_performed'])
        with self.assertRaises(service.WarehouseAssistantError):
            self.send()
        self.assertEqual(self.transport.call_count, 1)

    def test_concurrent_clicks_deliver_once(self):
        token = self.order()['email_send_token']
        def attempt(_):
            try:
                return self.send(token)
            except service.WarehouseAssistantError:
                return None
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(attempt, range(2)))
        self.assertEqual(self.transport.call_count, 1)

    def test_smtp_tls_and_quit_error_do_not_resend(self):
        self.send()
        config, message = self.transport.call_args.args
        with patch.object(mail.smtplib, 'SMTP_SSL') as factory:
            client = factory.return_value
            client.send_message.return_value = {}
            client.quit.side_effect = OSError('disconnect after acceptance')
            REAL_SMTP_SEND(config, message)
            import ssl
            self.assertEqual(factory.call_args.kwargs['context'].verify_mode, ssl.CERT_REQUIRED)
            self.assertTrue(factory.call_args.kwargs['context'].check_hostname)
            client.send_message.assert_called_once()

    def test_transport_failures_are_classified_without_secret_leak(self):
        self.send()
        config, message = self.transport.call_args.args
        for phase, exc, expected in [
            ('login', smtplib.SMTPAuthenticationError(535, b'test-secret'), 'failed'),
            ('send_message', smtplib.SMTPDataError(550, b'test-secret'), 'failed'),
            ('send_message', OSError('test-secret'), 'unknown'),
        ]:
            with patch.object(mail.smtplib, 'SMTP_SSL') as factory:
                getattr(factory.return_value, phase).side_effect = exc
                with self.assertRaises(mail.MailDeliveryFailure) as caught:
                    REAL_SMTP_SEND(config, message)
                self.assertEqual(caught.exception.status, expected)
                self.assertNotIn('test-secret', str(caught.exception))

    def test_claim_survives_interruption_and_blocks_retry(self):
        self.transport.side_effect = KeyboardInterrupt()
        with self.assertRaises(KeyboardInterrupt):
            self.send()
        self.assertEqual(self.order()['email_delivery']['status'], 'sending')
        with self.assertRaises(service.WarehouseAssistantError):
            self.send()
        self.assertEqual(self.transport.call_count, 1)

    def test_history_preserves_failed_attempt(self):
        self.transport.side_effect = mail.MailDeliveryFailure('failed', 'رد شد')
        self.send()
        self.transport.side_effect = None
        self.send()
        with service.warehouse_connection(self.settings) as conn:
            rows = conn.execute('SELECT status,attachment_sha256 FROM warehouse_email_attempts ORDER BY id').fetchall()
        self.assertEqual([r['status'] for r in rows], ['failed', 'sent'])
        self.assertTrue(all(len(r['attachment_sha256']) == 64 for r in rows))

    def test_real_route_requires_session_permission_and_confirmation(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from app.routes.warehouse_assistant import router
        from app.auth_service import create_session, create_user
        from app.database import init_sqlite, sqlite_connection
        self.settings.action_api_key = 'isolated-test-key'
        self.settings.login_username = ''
        init_sqlite(self.settings.sqlite_path)
        create_user(self.settings, 'Admin', 'TestOnlyPass9')
        with sqlite_connection(self.settings.sqlite_path) as conn:
            conn.execute("UPDATE users SET role='Admin' WHERE username='Admin'")
        app = FastAPI()
        app.state.settings = self.settings
        app.include_router(router)
        url = '/warehouse-assistant/api/automatic-preorders/1/send-email'
        with TestClient(app) as client:
            self.assertEqual(client.post(url, json={}).status_code, 401)
            client.cookies.set('negin_session', create_session(self.settings, 'Admin'))
            self.assertEqual(client.post(url, json={}).status_code, 422)
            with patch('app.routes.warehouse_assistant._capabilities', return_value=set()):
                self.assertEqual(client.post(url, json={'expected_token':self.order()['email_send_token']}).status_code, 403)
            self.transport.assert_not_called()
            reply = client.post(url, json={'expected_token': self.order()['email_send_token']})
            self.assertEqual(reply.status_code, 200, reply.text)
            self.assertEqual(reply.json()['send_status'], 'sent')
            listing = client.get('/warehouse-assistant/api/automatic-preorders').json()
            self.assertEqual(listing['preorders'], [])
            inbox = client.get('/warehouse-assistant/api/fulfillment-orders').json()
            self.assertEqual(inbox['orders'][0]['email_delivery']['status'], 'sent')
            self.assertEqual(inbox['orders'][0]['fulfillment']['remaining_qty'], 24)
            download = client.get('/warehouse-assistant/api/automatic-preorders/1/document.xlsx')
            self.assertEqual(download.status_code, 200)
            self.assertEqual(load_workbook(BytesIO(download.content))['گزارش'].cell(2, 9).value, 2)


if __name__ == '__main__':
    unittest.main()
