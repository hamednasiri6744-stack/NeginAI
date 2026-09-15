from io import BytesIO
import pytest
from openpyxl import load_workbook
from test_warehouse_fulfillment import case
from app import warehouse_assistant_service as service
from app.warehouse_preorder_catalog import get_preorder_catalog
from app.warehouse_email import preorder_workbook


@pytest.fixture
def editable(case):
    case.change("UPDATE warehouse_automatic_preorders SET status='awaiting_approval'")
    with service.warehouse_connection(case.settings) as conn:
        for index, code, warehouse, supplier in [(2,'002','karaj','supplier'),(3,'003','karaj','different'),(4,'004','tehran','supplier')]:
            conn.execute('''INSERT INTO warehouse_snapshot_items(snapshot_id,source_row,warehouse_code,
                warehouse_name,product_code,product_name,manufacturer,brand,conversion_rate,
                manufacturer_product_code,barcode,buy_price,manufacturer_price,consumer_price,stock,period_out)
                VALUES(1,?,?,?,?,?,?,?,6,'00017','00012345678901234567',10,20,30,12,120)''',
                (index,warehouse,warehouse,code,'Additional','supplier' if supplier=='supplier' else supplier,'brand'))
    return case


def save(case, lines, token=None):
    return service.update_automatic_preorder_lines(case.settings,'tester',1,lines,
        expected_token=token or case.order()['email_send_token'])


def test_catalog_is_scoped_cached_and_has_decision_data(editable):
    data = get_preorder_catalog(editable.settings,1)
    assert {i['product_code'] for i in data['items']} == {'00123','002'}
    item = next(i for i in data['items'] if i['product_code']=='002')
    assert item['manufacturer_product_code']=='00017'
    assert item['barcode']=='00012345678901234567'
    assert item['average_daily_out']==2
    assert item['effective_procurement_qty']==12
    assert item['coverage_days']==6
    assert item['approximate_price']==10
    assert item['in_transit_qty']==0


def test_add_line_recalculates_and_download_contains_text_identifiers(editable):
    order=save(editable,[{'product_code':'00123','cartons':3},{'product_code':'002','cartons':2}])
    assert order['status']=='awaiting_approval'
    assert (order['item_count'],order['total_cartons'],order['total_quantity'])==(2,5,48)
    assert order['estimated_value']==36*70+12*10
    new=next(line for line in order['lines'] if line['product_code']=='002')
    assert new['system_suggested_cartons']==0
    assert new['manufacturer_product_code']=='00017'
    sheet=load_workbook(BytesIO(preorder_workbook(order)))['گزارش']
    cells=next(row for row in sheet.iter_rows(min_row=2) if row[3].value=='002')
    assert (cells[4].value,cells[5].value)==('00017','00012345678901234567')
    assert cells[4].data_type==cells[5].data_type=='s'
    editable.transport.assert_not_called()


def test_zero_cartons_removes_one_line_and_keeps_remaining_order(editable):
    added = save(editable, [
        {'product_code': '00123', 'cartons': 3},
        {'product_code': '002', 'cartons': 2},
    ])
    removed = save(editable, [
        {'product_code': '00123', 'cartons': 3},
        {'product_code': '002', 'cartons': 0},
    ], added['email_send_token'])
    assert removed['item_count'] == 1
    assert removed['total_cartons'] == 3
    assert [line['product_code'] for line in removed['lines']] == ['00123']
    with pytest.raises(service.WarehouseAssistantError):
        save(editable, [{'product_code': '00123', 'cartons': 0}],
             removed['email_send_token'])


@pytest.mark.parametrize('code',['003','004','missing'])
def test_reject_foreign_or_missing_items_atomically(editable,code):
    before=editable.order()
    with pytest.raises(service.WarehouseAssistantError):
        save(editable,[{'product_code':'00123','cartons':4},{'product_code':code,'cartons':2}])
    assert editable.order()['lines']==before['lines']


def test_stale_duplicate_missing_existing_and_approved_rejected(editable):
    old=editable.order()['email_send_token']
    save(editable,[{'product_code':'00123','cartons':3}])
    with pytest.raises(service.WarehouseAssistantError):
        save(editable,[{'product_code':'00123','cartons':2},{'product_code':'002','cartons':1}],old)
    for lines in [[{'product_code':'002','cartons':1}],
                  [{'product_code':'00123','cartons':1},{'product_code':'00123','cartons':2}]]:
        with pytest.raises(service.WarehouseAssistantError):save(editable,lines)
    editable.change("UPDATE warehouse_automatic_preorders SET status='approved'")
    with pytest.raises(service.WarehouseAssistantError):get_preorder_catalog(editable.settings,1)
    with pytest.raises(service.WarehouseAssistantError):save(editable,[{'product_code':'00123','cartons':2},{'product_code':'002','cartons':1}])


def test_catalog_route_is_permission_protected(editable):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.routes import warehouse_assistant as routes
    from unittest.mock import patch
    app=FastAPI();app.state.settings=editable.settings;app.include_router(routes.router)
    with TestClient(app) as client:
        assert client.get('/warehouse-assistant/api/automatic-preorders/1/catalog').status_code==401
        app.dependency_overrides[routes.require_session_user]=lambda:'tester'
        with patch.object(routes,'_capabilities',return_value=set()):
            assert client.get('/warehouse-assistant/api/automatic-preorders/1/catalog').status_code==403
        with patch.object(routes,'_require',return_value='tester'):
            response=client.get('/warehouse-assistant/api/automatic-preorders/1/catalog')
            assert response.status_code==200
            body=response.json()
            response=client.put('/warehouse-assistant/api/automatic-preorders/1/lines',json={
                'expected_token':body['expected_token'],'lines':[{'product_code':'00123','cartons':2},{'product_code':'002','cartons':1}]})
            assert response.status_code==200,response.text
            assert response.json()['preorder']['item_count']==2


def test_manual_supplier_order_also_resolves_identifiers(editable):
    order=service.create_supplier_orders(editable.settings,'tester',snapshot_id=1,warehouse='karaj',
        lines=[{'product_code':'002','quantity':6}],note='test')[0]
    assert order['lines'][0]['manufacturer_product_code']=='00017'
    assert order['lines'][0]['barcode']=='00012345678901234567'
    assert order['lines'][0]['coverage_days']==6


def test_manual_supplier_excel_priority_uses_saved_coverage(editable):
    from app.routes import warehouse_assistant as routes
    from unittest.mock import patch
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    editable.change("UPDATE warehouse_snapshot_items SET stock=6 WHERE product_code='002'")
    order=service.create_supplier_orders(editable.settings,'tester',snapshot_id=1,warehouse='karaj',
        lines=[{'product_code':'002','quantity':6}],note='test')[0]
    assert order['lines'][0]['coverage_days']==3
    app=FastAPI();app.include_router(routes.router)
    app.dependency_overrides[routes.require_session_user]=lambda:'tester'
    with patch.object(routes,'_order_for_document',return_value=order), TestClient(app) as client:
        response=client.get(f"/warehouse-assistant/api/supplier-orders/{order['id']}/document.xlsx")
    assert response.status_code==200
    sheet=load_workbook(BytesIO(response.content))['گزارش']
    assert sheet['Q2'].value=='اولویت ارسال'
    assert sheet['A2'].fill.fgColor.rgb=='00FEE2E2'
    editable.transport.assert_not_called()


def test_excel_formula_identifiers_are_inert_text(editable):
    order=editable.order()
    order['lines'][0].update(manufacturer_product_code='=1+1',barcode='+123')
    cells=list(load_workbook(BytesIO(preorder_workbook(order)))['گزارش'].iter_rows(min_row=2))[0]
    assert cells[4].value=="'=1+1" and cells[4].data_type=='s'
    assert cells[5].value=="'+123" and cells[5].data_type=='s'


@pytest.mark.parametrize('stock,urgent', [(7.99998, True), (8, False), (0, True)])
def test_automatic_excel_priority_preserves_strict_boundary(editable, stock, urgent):
    with service.warehouse_connection(editable.settings) as conn:
        conn.execute("UPDATE warehouse_snapshot_items SET stock=? WHERE product_code='00123'", (stock,))
    order=editable.order()
    sheet=load_workbook(BytesIO(preorder_workbook(order)))['گزارش']
    assert (sheet['P2'].value=='اولویت ارسال') == urgent
    assert order['lines'][0]['coverage_days'] == stock / 2
    editable.transport.assert_not_called()


def test_disabled_supply_scope_cannot_be_bypassed_by_direct_add(editable):
    editable.change("INSERT INTO warehouse_supply_scope_state(id,strict_enabled,source_kind,refreshed_at,refreshed_by) VALUES(1,1,'test','now','test')")
    data=get_preorder_catalog(editable.settings,1)
    assert all(not item['can_add'] for item in data['items'])
    with pytest.raises(service.WarehouseAssistantError):
        save(editable,[{'product_code':'00123','cartons':2},{'product_code':'002','cartons':1}])


def test_additions_require_token_even_from_legacy_client(editable):
    with pytest.raises(service.WarehouseAssistantError):
        service.update_automatic_preorder_lines(editable.settings,'test',1,
            [{'product_code':'00123','cartons':2},{'product_code':'002','cartons':1}])
