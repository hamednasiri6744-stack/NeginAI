from __future__ import annotations

from contextlib import contextmanager
from dataclasses import replace
from datetime import timedelta
from io import BytesIO
import sqlite3
import pytest

from openpyxl import Workbook, load_workbook

from app.auth_service import create_session, create_user
from app.control_service import control_snapshot
from app.database import sqlite_connection
from app.business_time import jalali_business_date, tehran_now
import app.warehouse_assistant_service as warehouse_service


@pytest.fixture(autouse=True)
def isolate_application_startup(settings, tmp_path, monkeypatch):
    # The shared client fixture starts lifespan before assigning test settings.
    # Keep startup writes and background workers away from the live application.
    import app.main as main
    isolated = replace(settings, metadata_sync_enabled=False, automation_enabled=False,
                       vapid_private_key_path=tmp_path / 'test-vapid.pem')
    monkeypatch.setattr(main, 'get_settings', lambda: isolated)
    monkeypatch.setattr(main, 'ensure_action_api_key', lambda: None)


def _session(settings, username: str, role: str = "کارمند انبار") -> dict[str, str]:
    create_user(settings, username, "StrongPass9")
    with sqlite_connection(settings.sqlite_path) as conn:
        conn.execute(
            "UPDATE users SET role=?, full_name=? WHERE username=?",
            (role, username, username),
        )
    return {"Cookie": f"negin_session={create_session(settings, username)}"}


def _inventory_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "فایل انبار"
    sheet.append(
        [
            "کد کالا",
            "item name",
            "conversion  rate",
            "Manufacturer",
            "Brand",
            "stock karaj",
            "reserved karaj",
            "two month out karaj",
            "sale price karaj",
            "manufacturer price karaj",
            "consumer price karaj",
            "karaj 30 days stock",
            "Buy Price Karaj",
            "stock tehran",
            "reserved tehran",
            "two month out tehran",
            "sale price tehran",
            "manufacturer price tehran",
            "consumer price tehran",
            "tehran 30 days stock",
            "Buy Price Tehran",
            "Stock Gilan",
            "Reserved Gilan",
            "Two Month Out Gilan",
            "gilan 30 days stock",
            "sale price gilan",
            "manufacturer price gilan",
            "consumer price gilan",
            "Buy Price gilan",
        ]
    )
    sheet.append(
        [
            4014,
            "خمیر دندان آزمایشی",
            12,
            "تامین‌کننده نمونه",
            "میسویک",
            10,
            2,
            60,
            100,
            90,
            120,
            8,
            70,
            30,
            0,
            30,
            100,
            90,
            120,
            30,
            70,
            0,
            0,
            0,
            0,
            100,
            90,
            120,
            70,
        ]
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _coverage_workbook() -> bytes:
    workbook = load_workbook(BytesIO(_inventory_workbook()))
    sheet = workbook["فایل انبار"]
    base_row = [cell.value for cell in sheet[2]]

    twenty_days = list(base_row)
    twenty_days[0] = 4020
    twenty_days[1] = "کالای با پوشش بیست روز"
    twenty_days[5] = 20
    twenty_days[6] = 0
    sheet.append(twenty_days)

    fifteen_days = list(base_row)
    fifteen_days[0] = 4015
    fifteen_days[1] = "کالای با پوشش دقیق پانزده روز"
    fifteen_days[5] = 15
    fifteen_days[6] = 0
    sheet.append(fifteen_days)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_warehouse_assistant_page_is_separate_from_main_assistant(client):
    page = client.get("/warehouse-assistant")
    script = client.get("/static/warehouse-assistant.js?v=84")
    stylesheet = client.get("/static/warehouse-assistant.css?v=37")
    assistant = client.get("/assistant")

    assert page.status_code == 200
    assert 'id="warehouseAssistantApp"' in page.text
    assert "/static/warehouse-assistant.js?v=84" in page.text
    assert "/static/warehouse-assistant.css?v=37" in page.text
    assert 'id="inventoryDisplayOptions"' in page.text
    assert 'data-guide="inventory"' in page.text
    assert 'aria-labelledby="previewPreorderTitle"' in page.text
    assert script.status_code == 200
    assert stylesheet.status_code == 200
    assert "/warehouse-assistant/api/bootstrap" in script.text
    assert 'id="loginDialog"' in page.text
    assert "/auth/login" in script.text
    assert "/warehouse-assistant/api/sync/varanegar" in script.text
    assert "/warehouse-assistant/api/inventory" in script.text
    assert 'id="inventoryView"' in page.text
    assert 'id="inventoryViewSelect"' in page.text
    assert 'data-filter="manufacturer_product_code"' in page.text
    assert 'data-filter="barcode"' in page.text
    assert "/warehouse-assistant/api/inventory/views" in script.text
    assert "/warehouse-assistant/api/catalog" in script.text
    assert 'id="supplyView"' in page.text
    assert 'id="supplyWarehouseSelect"' in page.text
    assert "/warehouse-assistant/api/supply-scope" in script.text
    assert "بازخوانی رسیدهای ۹۰ روز" in page.text
    assert '<select id="manufacturerFilter">' in page.text
    assert '<select id="brandFilter"' in page.text
    assert 'id="reorderCoverageDays"' in page.text
    assert "میانگین فروش روزانه تعدیل‌شده" in page.text
    assert 'data-order-column="stockout_days"' in page.text
    assert 'data-order-column="sales_window_start"' in page.text
    assert 'data-column="ordering_cycle_active"' in page.text
    assert 'data-column="sale_price"' in page.text
    assert 'data-column="manufacturer_price"' in page.text
    assert 'data-column="consumer_price"' in page.text
    assert "inventory-cycle-select" in script.text
    assert "/warehouse-assistant/api/inventory/order-cycle" in script.text
    assert 'id="orderingColumnChoices"' in page.text
    assert 'data-order-column="manufacturer"' in page.text
    assert 'data-order-column="brand"' in page.text
    assert 'id="automaticView"' in page.text
    assert 'id="automaticRows"' in page.text
    assert 'id="saveAutomaticSettingsButton"' in page.text
    assert "/warehouse-assistant/api/automatic-settings" in script.text
    assert "حداقل کل سفارش" in page.text
    assert "هر تأمین‌کننده در هر انبار تنظیم مستقل دارد" in page.text
    assert 'id="automaticPreordersList"' in page.text
    assert 'id="preorderPreviewDialog"' in page.text
    assert 'id="previewPreorderLines"' in page.text
    assert "قابل استفاده انبار (بدون در راه)" in page.text
    assert "مبنای سفارش (جمع)" in page.text
    assert "اعضای گروه مشتری امیران کامل لحاظ می‌شوند" in page.text
    assert "حداقل ۸۰٪ فروش همان روز" in page.text
    assert "انتقال قطعی به انبار آنلاین" in page.text
    assert "بدون تعدیل هیجان" in page.text
    assert "line.unadjusted_suggested_cartons" in script.text
    assert "demandAuditNote(item)" in script.text
    assert "خروج روزانه" in page.text
    assert "پوشش موجودی (روز)" in page.text
    assert "قیمت تولیدکننده" in page.text
    assert "قیمت مصرف‌کننده" in page.text
    assert "قیمت حدودی" in page.text
    assert "line.approximate_price" in script.text
    assert "previewDataValue(line.physical_procurement_qty" in script.text
    assert "قابل استفاده انبار (بدون در راه)" in page.text
    assert "مبنای سفارش (جمع)" in page.text
    assert "مشاهده اقلام" in script.text
    assert "/lines" in script.text
    assert 'class="app-sidebar"' in page.text
    assert 'id="preordersView"' in page.text
    assert 'id="preorderWarehouseFilter"' in page.text
    assert 'id="automaticColumnChoices"' in page.text
    assert 'data-order-filter="product_code"' in page.text
    assert "/warehouse-assistant/api/table-preferences" in script.text
    assert "/warehouse-assistant/api/automatic-preorders/run" in script.text
    assert "/warehouse-assistant/api/automatic-preorders/refresh-status" in script.text
    assert "بازسازی دستی همه پیش‌سفارش‌ها" in page.text
    assert 'id="manualRefreshPreordersButton"' in page.text
    assert "به‌روزرسانی دستی سفارش‌های اتوماتیک" in page.text
    assert "آخرین به‌روزرسانی" in page.text
    assert "preorder.last_refreshed_at||preorder.created_at" in script.text
    assert "updatePreparedOrderCounts()" in script.text
    assert "fitTableWrapsToViewport" in script.text
    assert "initializeTableScrollPersistence" in script.text
    assert 'data-scroll-key="inventory"' in page.text
    assert "--wa-topbar-height:40px" in stylesheet.text
    assert "--wa-font-table:11px" in stylesheet.text
    assert "automaticDrafts" in script.text
    assert "ابتدا تغییرات تنظیمات را ذخیره کنید" in script.text
    assert "auto-dirty" in stylesheet.text
    assert "viewport-fitted" in stylesheet.text
    assert "هر یک ساعت" in page.text
    workflow_script = client.get('/static/warehouse-order-workflow.js?v=5')
    assert workflow_script.status_code == 200
    assert "قرار دادن در کارتابل تأمین‌کننده" in workflow_script.text
    assert 'data-order-sms-kind="supplier-order"' not in script.text
    assert "مبنای سفارش (با در راه)" in page.text
    assert "warehouse-assistant.js" not in assistant.text
    assert "warehouseAssistantApp" not in assistant.text


def test_warehouse_assistant_api_requires_explicit_permission(client, settings):
    assert client.get("/warehouse-assistant/api/bootstrap").status_code == 401

    denied_headers = _session(settings, "warehouse.denied")
    denied = client.get(
        "/warehouse-assistant/api/bootstrap", headers=denied_headers
    )
    assert denied.status_code == 403

    admin_headers = _session(settings, "Admin", "Admin")
    allowed = client.get(
        "/warehouse-assistant/api/bootstrap", headers=admin_headers
    )
    assert allowed.status_code == 200
    assert allowed.json()["separated_from_assistant"] is True
    assert allowed.json()["commit_enabled"] is False
    assert allowed.json()["output_mode"] == "supplier_documents"
    assert allowed.json()["varanegar_integration"] == "not_applicable"


def test_table_column_preferences_are_scoped_to_each_user(client, settings):
    first_user = _session(settings, "warehouse.columns.one", "Admin")
    second_user = _session(settings, "warehouse.columns.two", "Admin")

    saved = client.put(
        "/warehouse-assistant/api/table-preferences/ordering",
        headers=first_user,
        json={"visible_columns": ["product_code", "product_name", "final_order_cartons"]},
    )
    assert saved.status_code == 200
    assert saved.json()["visible_columns"] == [
        "product_code", "product_name", "final_order_cartons"
    ]

    first = client.get(
        "/warehouse-assistant/api/table-preferences", headers=first_user
    )
    second = client.get(
        "/warehouse-assistant/api/table-preferences", headers=second_user
    )
    assert first.json()["preferences"]["ordering"] == [
        "product_code", "product_name", "final_order_cartons"
    ]
    assert "ordering" not in second.json()["preferences"]

    invalid = client.put(
        "/warehouse-assistant/api/table-preferences/ordering",
        headers=first_user,
        json={"visible_columns": ["unknown_column"]},
    )
    assert invalid.status_code == 422


def test_permission_catalog_contains_warehouse_capabilities(settings):
    create_user(settings, "Admin", "StrongPass9")
    snapshot = control_snapshot(settings, "Admin")
    keys = {item["key"] for item in snapshot["permissions"]}

    assert {
        "warehouse.assistant.view",
        "warehouse.order.suggest",
        "warehouse.order.draft",
        "warehouse.data.refresh",
    } <= keys


def test_inventory_snapshot_import_and_deterministic_suggestion_are_isolated(
    client, settings
):
    headers = _session(settings, "Admin", "Admin")
    content = _inventory_workbook()

    imported = client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={
            "file": (
                "inventory.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert imported.status_code == 201
    body = imported.json()
    assert body["product_count"] == 1
    assert body["item_count"] == 3
    assert body["duplicate"] is False
    assert body["source_sheet"] == "فایل انبار"

    duplicate = client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("inventory.xlsx", content)},
    )
    assert duplicate.status_code == 200
    assert duplicate.json()["duplicate"] is True
    assert duplicate.json()["id"] == body["id"]

    inventory = client.get(
        "/warehouse-assistant/api/inventory",
        headers=headers,
        params={"warehouse": "karaj", "search": "4014"},
    )
    assert inventory.status_code == 200
    assert inventory.json()["items"][0]["sale_price"] == 100
    assert inventory.json()["items"][0]["manufacturer_price"] == 90
    assert inventory.json()["items"][0]["consumer_price"] == 120
    assert "approximate_price" not in inventory.json()["items"][0]
    assert "buy_price" not in inventory.json()["items"][0]

    response = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "brand": "میسویک",
            "target_days": 30,
            "safety_days": 0,
            "period_days": 60,
        },
    )

    assert response.status_code == 200
    suggestion = response.json()["items"][0]
    assert suggestion["product_code"] == "4014"
    assert suggestion["manufacturer_product_code"] == ""
    assert suggestion["barcode"] == ""
    assert suggestion["available_quantity"] == 12
    assert suggestion["effective_procurement_qty"] == 12
    assert suggestion["open_order_qty"] == 0
    assert suggestion["average_daily_out"] == 1
    assert suggestion["suggested_quantity"] == 24
    assert suggestion["suggested_cartons"] == 2
    assert suggestion["effective_cartons"] == 1
    assert suggestion["calculation"] == {
        "target_quantity": 30,
        "raw_requirement": 18,
        "rounded_to_conversion_rate": 12,
    }

    warehouse_db = settings.sqlite_path.with_name("warehouse-assistant.db")
    assert warehouse_db.exists()
    with sqlite3.connect(settings.sqlite_path) as conn:
        assert conn.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='warehouse_snapshots'"
        ).fetchone()[0] == 0


def test_reorder_coverage_triggers_below_15_days_and_refills_to_30(client, settings):
    headers = _session(settings, "Admin", "Admin")
    imported = client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("coverage.xlsx", _coverage_workbook())},
    )
    assert imported.status_code == 201

    response = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "reorder_coverage_days": 15,
            "target_days": 30,
            "safety_days": 0,
            "period_days": 60,
            "only_needed": True,
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["parameters"]["reorder_coverage_days"] == 15
    assert [item["product_code"] for item in body["items"]] == ["4014"]
    item = body["items"][0]
    assert item["coverage_days"] == 12
    assert item["needs_reorder"] is True
    assert item["calculation"]["target_quantity"] == 30
    assert item["calculation"]["raw_requirement"] == 18
    assert item["suggested_quantity"] == 24

    ten_day_order = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "reorder_coverage_days": 10,
            "target_days": 30,
            "safety_days": 0,
            "period_days": 60,
            "only_needed": True,
        },
    )
    assert ten_day_order.status_code == 200
    assert ten_day_order.json()["parameters"]["reorder_coverage_days"] == 10
    assert ten_day_order.json()["items"] == []

    diagnostic = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "reorder_coverage_days": 15,
            "target_days": 30,
            "safety_days": 0,
            "period_days": 60,
            "only_needed": False,
        },
    )
    assert diagnostic.status_code == 200
    by_code = {item["product_code"]: item for item in diagnostic.json()["items"]}
    assert by_code["4015"]["coverage_days"] == 15
    assert by_code["4015"]["needs_reorder"] is False
    assert by_code["4015"]["suggested_quantity"] == 0
    assert by_code["4020"]["coverage_days"] == 20
    assert by_code["4020"]["needs_reorder"] is False
    assert by_code["4020"]["suggested_quantity"] == 0


def test_automatic_order_settings_default_to_10_20_and_enforce_supplier_minimum(
    client, settings
):
    headers = _session(settings, "Admin", "Admin")
    imported = client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("inventory.xlsx", _inventory_workbook())},
    )
    assert imported.status_code == 201

    listing = client.get(
        "/warehouse-assistant/api/automatic-settings", headers=headers
    )
    assert listing.status_code == 200
    body = listing.json()
    assert body["defaults"] == {
        "reorder_coverage_days": 10,
        "target_days": 20,
        "minimum_cartons": 0,
    }
    assert body["scope"] == "supplier_per_warehouse"
    assert body["varanegar_write"] is False
    assert len(body["settings"]) == 3
    assert {item["warehouse_code"] for item in body["settings"]} == {
        "karaj", "tehran", "gilan"
    }
    setting = next(
        item for item in body["settings"] if item["warehouse_code"] == "karaj"
    )
    assert setting["warehouse_name"] == "انبار مرکزی کرج"
    assert setting["supplier"] == "تامین‌کننده نمونه"
    assert setting["enabled"] is True
    assert setting["reorder_coverage_days"] == 10
    assert setting["target_days"] == 20
    assert setting["minimum_cartons"] == 0

    saved = client.put(
        f"/warehouse-assistant/api/automatic-settings/{setting['id']}",
        headers=headers,
        json={
            "enabled": True,
            "reorder_coverage_days": 15,
            "target_days": 20,
            "minimum_cartons": 200,
            "contact_first_name": "علی",
            "contact_last_name": "محمدی",
            "contact_email": "BUY@EXAMPLE.COM",
            "contact_mobile": "+۹۸۹۱۲۳۴۵۶۷۸۹",
        },
    )
    assert saved.status_code == 200
    assert saved.json()["contact_first_name"] == "علی"
    assert saved.json()["contact_last_name"] == "محمدی"
    assert saved.json()["contact_email"] == "buy@example.com"
    assert saved.json()["contact_mobile"] == "+989123456789"
    assert saved.json()["minimum_cartons"] == 200

    below_minimum = client.get(
        "/warehouse-assistant/api/automatic-settings/preview/karaj",
        headers=headers,
    )
    assert below_minimum.status_code == 200
    preview = below_minimum.json()["suppliers"][0]
    assert preview["suggested_items"] == 1
    assert preview["suggested_cartons"] == 1
    assert preview["minimum_shortfall_cartons"] == 199
    assert preview["ready_to_prepare"] is False
    assert preview["readiness_status"] == "below_minimum"
    assert below_minimum.json()["documents_created"] is False

    ready_payload = {
        "enabled": True,
        "reorder_coverage_days": 15,
        "target_days": 20,
        "minimum_cartons": 1,
        "contact_first_name": "علی",
        "contact_last_name": "محمدی",
        "contact_email": "buy@example.com",
        "contact_mobile": "+989123456789",
    }
    assert client.put(
        f"/warehouse-assistant/api/automatic-settings/{setting['id']}",
        headers=headers,
        json=ready_payload,
    ).status_code == 200
    ready = client.get(
        "/warehouse-assistant/api/automatic-settings/preview/karaj",
        headers=headers,
    ).json()
    assert ready["ready_suppliers"] == 1
    assert ready["suppliers"][0]["ready_to_prepare"] is True
    assert ready["suppliers"][0]["readiness_status"] == "ready"

    invalid = dict(ready_payload, reorder_coverage_days=20, target_days=20)
    rejected = client.put(
        f"/warehouse-assistant/api/automatic-settings/{setting['id']}",
        headers=headers,
        json=invalid,
    )
    assert rejected.status_code == 422
    assert "بیشتر" in rejected.json()["detail"]


def test_automatic_preorders_are_separate_per_warehouse_and_wait_for_user_send(
    client, settings
):
    headers = _session(settings, "Admin", "Admin")
    assert client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("inventory.xlsx", _inventory_workbook())},
    ).status_code == 201
    settings_body = client.get(
        "/warehouse-assistant/api/automatic-settings", headers=headers
    ).json()
    assert all(item["minimum_cartons"] == 0 for item in settings_body["settings"])
    karaj_setting = next(
        item
        for item in settings_body["settings"]
        if item["warehouse_code"] == "karaj"
    )
    saved = client.put(
        f"/warehouse-assistant/api/automatic-settings/{karaj_setting['id']}",
        headers=headers,
        json={
            "enabled": True,
            "reorder_coverage_days": 15,
            "target_days": 20,
            "minimum_cartons": 0,
            "contact_first_name": "علی",
            "contact_last_name": "محمدی",
            "contact_email": "buy@example.com",
            "contact_mobile": "09123456789",
        },
    )
    assert saved.status_code == 200
    assert saved.json()["warehouse_code"] == "karaj"

    prepared = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    )
    assert prepared.status_code == 200
    body = prepared.json()
    assert body["evaluated_settings"] == 3
    assert body["eligible_settings"] == 1
    assert body["created_count"] == 1
    assert body["external_delivery_performed"] is False
    preorder = body["preorders"][0]
    assert preorder["warehouse_code"] == "karaj"
    assert preorder["warehouse_name"] == "انبار مرکزی کرج"
    assert preorder["supplier"] == "تامین‌کننده نمونه"
    assert preorder["minimum_cartons"] == 0
    assert preorder["total_cartons"] == 1
    assert preorder["status"] == "awaiting_approval"

    repeated = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    )
    assert repeated.status_code == 200
    assert repeated.json()["created_count"] == 0
    assert repeated.json()["reused_count"] == 1
    assert len(repeated.json()["preorders"]) == 1

    line = preorder["lines"][0]
    assert line["manufacturer_price"] == 90
    assert line["consumer_price"] == 120
    assert line["approximate_price"] == 70
    assert line["estimated_value"] == 840
    assert line["system_suggested_cartons"] == 1
    assert line["unadjusted_suggested_cartons"] == 1
    assert line["effective_procurement_qty"] == 12
    assert line["physical_procurement_qty"] == 12
    assert line["in_transit_qty"] == 0
    assert line["pending_receipt_qty"] == 0
    assert line["inventory_position_qty"] == 12
    assert line["average_daily_out"] == 1
    assert line["sales_rate_days"] == 60
    assert line["period_out_qty"] == 60
    assert line["coverage_days"] == 12
    assert line["raw_period_out"] == 60
    assert line["adjusted_period_out"] == 60
    assert line["amiran_period_out"] == 0
    assert line["exceptional_period_out"] == 0
    assert line["demand_anomaly_days"] == 0
    assert line["demand_recurring_pattern"] is False
    assert line["demand_recurring_pattern_days"] == 0
    edited = client.put(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/lines",
        headers=headers,
        json={
            "lines": [
                {"product_code": line["product_code"], "cartons": 2}
            ]
        },
    )
    assert edited.status_code == 200
    edited_preorder = edited.json()["preorder"]
    assert edited_preorder["total_cartons"] == 2
    assert edited_preorder["total_quantity"] == 2 * line["conversion_rate"]
    assert edited_preorder["lines"][0]["cartons"] == 2
    assert edited_preorder["lines"][0]["system_suggested_cartons"] == 1
    assert edited_preorder["edited_by"] == "Admin"
    assert edited.json()["external_delivery_performed"] is False

    unchanged_after_edit = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    )
    assert unchanged_after_edit.status_code == 200
    assert unchanged_after_edit.json()["reused_count"] == 1
    assert unchanged_after_edit.json()["preorders"][0]["total_cartons"] == 2
    assert unchanged_after_edit.json()["preorders"][0]["lines"][0]["cartons"] == 2

    premature_send = client.post(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/request-send",
        headers=headers,
    )
    assert premature_send.status_code == 409
    assert "ابتدا" in premature_send.json()["detail"]

    approved = client.post(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/approve",
        headers=headers,
    )
    assert approved.status_code == 200
    assert approved.json()["preorder"]["status"] == "approved"
    assert approved.json()["external_delivery_performed"] is False

    locked_edit = client.put(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/lines",
        headers=headers,
        json={
            "lines": [
                {"product_code": line["product_code"], "cartons": 3}
            ]
        },
    )
    assert locked_edit.status_code == 422
    assert "قبل از تأیید" in locked_edit.json()["detail"]

    revoke_url = f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/revoke-approval"
    assert client.post(revoke_url).status_code == 401
    old_token = approved.json()["preorder"]["email_send_token"]
    revoked = client.post(revoke_url, headers=headers)
    assert revoked.status_code == 200
    assert revoked.json()["preorder"]["status"] == "awaiting_approval"
    assert revoked.json()["preorder"]["approved_at"] is None
    assert revoked.json()["preorder"]["total_cartons"] == 2
    assert revoked.json()["external_delivery_performed"] is False
    reapproved = client.post(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/approve", headers=headers
    )
    assert reapproved.status_code == 200
    assert reapproved.json()["preorder"]["email_send_token"] != old_token
    fresh = client.get(f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}", headers=headers)
    assert fresh.status_code == 200
    assert fresh.json()["preorder"]["can_revoke_approval"] is True

    send_requested = client.post(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/request-send",
        headers=headers,
    )
    assert send_requested.status_code == 200
    assert send_requested.json()["preorder"]["status"] == "send_requested"
    assert send_requested.json()["send_status"] == "requested"
    assert send_requested.json()["external_delivery_performed"] is False

    protected = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    )
    assert protected.status_code == 200
    assert protected.json()["created_count"] == 0
    # Internal approval already freezes the order and includes it in transit;
    # requesting delivery no longer creates a second protected-sent bucket.
    assert protected.json()["protected_sent_count"] == 0

    listing = client.get(
        "/warehouse-assistant/api/automatic-preorders", headers=headers
    )
    assert listing.status_code == 200
    assert listing.json()["preorders"][0]["status"] == "send_requested"
    assert client.get(
        "/warehouse-assistant/api/supplier-orders", headers=headers
    ).json()["orders"] == []

    text_document = client.get(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/document.txt",
        headers=headers,
    )
    assert text_document.status_code == 200
    assert "پیش‌سفارش خرید نگین پخش" in text_document.text
    assert "انبار مرکزی کرج" in text_document.text
    assert "قیمت تولیدکننده" in text_document.text
    assert "قیمت مصرف‌کننده" in text_document.text
    assert "قیمت حدودی" in text_document.text

    excel_document = client.get(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/document.xlsx",
        headers=headers,
    )
    assert excel_document.status_code == 200
    workbook = load_workbook(BytesIO(excel_document.content), read_only=True)
    assert workbook["گزارش"]["C2"].value == "انبار مرکزی کرج"
    assert workbook["گزارش"]["I2"].value == 2
    assert workbook["گزارش"]["L1"].value == "قیمت تولیدکننده"
    assert workbook["گزارش"]["M1"].value == "قیمت مصرف‌کننده"
    assert workbook["گزارش"]["N1"].value == "قیمت حدودی"
    assert workbook["گزارش"]["L2"].value == 90
    assert workbook["گزارش"]["M2"].value == 120
    assert workbook["گزارش"]["N2"].value == 70
    workbook.close()


def test_hourly_refresh_keeps_same_day_approved_order_frozen_in_transit(
    client, settings
):
    headers = _session(settings, "Admin", "Admin")
    assert client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("inventory.xlsx", _inventory_workbook())},
    ).status_code == 201
    setting = next(
        item
        for item in client.get(
            "/warehouse-assistant/api/automatic-settings", headers=headers
        ).json()["settings"]
        if item["warehouse_code"] == "karaj"
    )
    payload = {
        "enabled": True,
        "reorder_coverage_days": 15,
        "target_days": 20,
        "minimum_cartons": 0,
        "contact_first_name": "",
        "contact_last_name": "",
        "contact_email": "",
        "contact_mobile": "",
    }
    assert client.put(
        f"/warehouse-assistant/api/automatic-settings/{setting['id']}",
        headers=headers,
        json=payload,
    ).status_code == 200
    first = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    ).json()
    preorder = first["preorders"][0]
    assert first["created_count"] == 1
    assert client.post(
        f"/warehouse-assistant/api/automatic-preorders/{preorder['id']}/approve",
        headers=headers,
    ).status_code == 200

    changed_payload = dict(payload, target_days=30)
    assert client.put(
        f"/warehouse-assistant/api/automatic-settings/{setting['id']}",
        headers=headers,
        json=changed_payload,
    ).status_code == 200
    refreshed = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    )
    assert refreshed.status_code == 200
    refreshed_body = refreshed.json()
    assert refreshed_body["created_count"] == 0
    assert refreshed_body["updated_count"] == 0
    assert len(refreshed_body["preorders"]) == 1
    updated = refreshed_body["preorders"][0]
    assert updated["id"] == preorder["id"]
    assert updated["status"] == "approved"
    assert updated["approved_at"] is not None
    assert updated["target_days"] == 20

    warehouse_db = warehouse_service.warehouse_database_path(settings)
    with sqlite3.connect(warehouse_db) as conn:
        conn.execute(
            """UPDATE warehouse_snapshot_items SET stock=1000000
               WHERE snapshot_id=? AND warehouse_code='karaj'""",
            (refreshed_body["snapshot_id"],),
        )
    no_need = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    )
    assert no_need.status_code == 200
    assert no_need.json()["removed_count"] == 0
    listing = client.get(
        "/warehouse-assistant/api/automatic-preorders", headers=headers
    ).json()["preorders"]
    assert any(item["id"] == preorder["id"] for item in listing)

    status_response = client.get(
        "/warehouse-assistant/api/automatic-preorders/refresh-status",
        headers=headers,
    )
    assert status_response.status_code == 200
    assert status_response.json()["last_trigger"] == "manual"
    assert status_response.json()["last_success_at"] is not None
    assert status_response.json()["next_run_at"] is not None


def test_refresh_preserves_previous_day_unsent_approved_orders_in_drafts(
    client, settings, monkeypatch
):
    current_business_date = {"value": "1405/06/12"}
    monkeypatch.setattr(
        warehouse_service,
        "jalali_business_date",
        lambda _value=None: current_business_date["value"],
    )
    headers = _session(settings, "Admin", "Admin")
    assert client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("inventory.xlsx", _inventory_workbook())},
    ).status_code == 201
    setting = next(
        item
        for item in client.get(
            "/warehouse-assistant/api/automatic-settings", headers=headers
        ).json()["settings"]
        if item["warehouse_code"] == "karaj"
    )
    assert client.put(
        f"/warehouse-assistant/api/automatic-settings/{setting['id']}",
        headers=headers,
        json={
            "enabled": True,
            "reorder_coverage_days": 15,
            "target_days": 20,
            "minimum_cartons": 0,
            "contact_first_name": "",
            "contact_last_name": "",
            "contact_email": "buy@example.com",
            "contact_mobile": "",
        },
    ).status_code == 200
    first = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    ).json()
    stale_id = first["preorders"][0]["id"]
    assert client.post(
        f"/warehouse-assistant/api/automatic-preorders/{stale_id}/approve",
        headers=headers,
    ).status_code == 200

    warehouse_db = warehouse_service.warehouse_database_path(settings)
    with sqlite3.connect(warehouse_db) as conn:
        conn.execute(
            """UPDATE warehouse_snapshot_items
               SET manufacturer='تامین‌کننده روز جدید'
               WHERE snapshot_id=? AND warehouse_code='karaj'""",
            (first["snapshot_id"],),
        )

    current_business_date["value"] = "1405/06/13"
    listing = client.get(
        "/warehouse-assistant/api/automatic-preorders", headers=headers
    )
    assert listing.status_code == 200
    assert [order['id'] for order in listing.json()['preorders']] == [stale_id]
    assert listing.json()['preorders'][0]['order_stage'] == 'draft'
    stale_send = client.post(
        f"/warehouse-assistant/api/automatic-preorders/{stale_id}/request-send",
        headers=headers,
    )
    assert stale_send.status_code == 422
    assert "روز جاری" in stale_send.json()["detail"]

    refreshed = client.post(
        "/warehouse-assistant/api/automatic-preorders/run", headers=headers
    )
    assert refreshed.status_code == 200
    body = refreshed.json()
    assert body["removed_count"] == 0
    assert any(item['id'] == stale_id and item['order_stage'] == 'draft' for item in body['preorders'])
    assert all(item["business_date"] == "1405/06/13" or item['id'] == stale_id for item in body["preorders"])
    with sqlite3.connect(warehouse_db) as conn:
        status = conn.execute(
            "SELECT status FROM warehouse_automatic_preorders WHERE id=?",
            (stale_id,),
        ).fetchone()[0]
    assert status == "approved"


def test_inventory_import_rejects_non_excel_files(client, settings):
    headers = _session(settings, "Admin", "Admin")
    response = client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("inventory.txt", b"not an excel workbook", "text/plain")},
    )

    assert response.status_code == 415


def test_admin_creates_supplier_order_and_downloads_ready_documents(client, settings):
    headers = _session(settings, "Admin", "Admin")
    imported = client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("inventory.xlsx", _inventory_workbook())},
    ).json()

    created = client.post(
        "/warehouse-assistant/api/supplier-orders",
        headers=headers,
        json={
            "snapshot_id": imported["id"],
            "warehouse": "karaj",
            "note": "تحویل در انبار مرکزی",
            "lines": [
                {
                    "product_code": "4014",
                    "quantity": 22,
                    "note": "مقدار پیشنهادی سیستم",
                }
            ],
        },
    )

    assert created.status_code == 201
    body = created.json()
    assert len(body["orders"]) == 1
    order = body["orders"][0]
    assert order["supplier"] == "تامین‌کننده نمونه"
    assert order["warehouse_code"] == "karaj"
    assert order["status"] == "prepared"
    assert order["delivery_mode"] == "supplier_document"
    assert order["lines"][0]["requested_quantity"] == 22
    assert order["lines"][0]["order_quantity"] == 24
    assert order["lines"][0]["cartons"] == 2
    assert order["lines"][0]["manufacturer_price"] == 90
    assert order["lines"][0]["consumer_price"] == 120
    assert order["lines"][0]["approximate_price"] == 70
    assert order["lines"][0]["estimated_value"] == 1680

    listing = client.get(
        "/warehouse-assistant/api/supplier-orders", headers=headers
    )
    assert listing.status_code == 200
    assert listing.json()["orders"][0]["order_number"] == order["order_number"]

    text_document = client.get(
        f"/warehouse-assistant/api/supplier-orders/{order['id']}/document.txt",
        headers=headers,
    )
    assert text_document.status_code == 200
    assert "سفارش خرید نگین پخش" in text_document.text
    assert "تامین‌کننده نمونه" in text_document.text
    assert "خمیر دندان آزمایشی" in text_document.text
    assert "24" in text_document.text
    assert "قیمت تولیدکننده" in text_document.text
    assert "قیمت مصرف‌کننده" in text_document.text
    assert "قیمت حدودی" in text_document.text

    excel_document = client.get(
        f"/warehouse-assistant/api/supplier-orders/{order['id']}/document.xlsx",
        headers=headers,
    )
    assert excel_document.status_code == 200
    workbook = load_workbook(BytesIO(excel_document.content), read_only=True)
    assert workbook["گزارش"]["A2"].value == order["order_number"]
    assert workbook["گزارش"]["D2"].value == "4014"
    assert workbook["گزارش"]["J2"].value == 24
    assert workbook["گزارش"]["L1"].value == "قیمت تولیدکننده"
    assert workbook["گزارش"]["M1"].value == "قیمت مصرف‌کننده"
    assert workbook["گزارش"]["N1"].value == "قیمت حدودی"
    assert workbook["گزارش"]["L2"].value == 90
    assert workbook["گزارش"]["M2"].value == 120
    assert workbook["گزارش"]["N2"].value == 70
    assert workbook["اطلاعات"]["B1"].value.startswith("سفارش خرید")


def test_manual_order_edit_api_checks_permission_token_and_approval(client, settings):
    headers = _session(settings, 'Admin', 'Admin')
    imported = client.post('/warehouse-assistant/api/snapshots/import', headers=headers,
                           files={'file': ('inventory.xlsx', _inventory_workbook())}).json()
    order = client.post('/warehouse-assistant/api/supplier-orders', headers=headers, json={
        'snapshot_id': imported['id'], 'warehouse':'karaj', 'lines':[{'product_code':'4014','quantity':24}]
    }).json()['orders'][0]
    path = f"/warehouse-assistant/api/supplier-orders/{order['id']}"
    assert client.get(path).status_code == 401
    payload = {'expected_token':order['email_send_token'], 'lines':[{'product_code':'4014','cartons':3}]}
    assert client.post(path+'/lines', json=payload).status_code == 401
    edited = client.post(path+'/lines', headers=headers, json=payload)
    assert edited.status_code == 200
    assert edited.json()['total_quantity'] == 36
    assert client.post(path+'/lines', headers=headers, json=payload).status_code == 422
    payload['expected_token'] = edited.json()['email_send_token']
    assert client.post(path+'/approve', headers=headers).status_code == 200
    assert client.post(path+'/lines', headers=headers, json=payload).status_code == 422
    listed = client.get('/warehouse-assistant/api/supplier-orders?stage=draft', headers=headers).json()['orders']
    assert listed[0]['id'] == order['id'] and listed[0]['can_send_portal']


def test_delivery_date_creation_and_edit_api_permissions(client, settings):
    headers = _session(settings, 'Admin', 'Admin')
    imported = client.post('/warehouse-assistant/api/snapshots/import', headers=headers,
        files={'file':('inventory.xlsx',_inventory_workbook())}).json()
    payload = dict(snapshot_id=imported['id'],warehouse='karaj',delivery_date='۱۴۰۵/۰۶/۲۵',
                   lines=[dict(product_code='4014',quantity=22)])
    created = client.post('/warehouse-assistant/api/supplier-orders',headers=headers,json=payload).json()['orders'][0]
    assert created['delivery_date'] == '1405/06/25'
    path = f"/warehouse-assistant/api/supplier-orders/{created['id']}/delivery-date"
    change = dict(delivery_date='1405/06/27', expected_token=created['email_send_token'])
    assert client.post(path,json=change).status_code == 401
    denied = _session(settings, 'delivery.denied')
    assert client.post(path,headers=denied,json=change).status_code == 403
    edited = client.post(path,headers=headers,json=change)
    assert edited.status_code == 200 and edited.json()['delivery_date'] == '1405/06/27'
    assert client.post(path,headers=headers,json=change).status_code == 422
    payload['delivery_date'] = '1405/07/31'
    assert client.post('/warehouse-assistant/api/supplier-orders',headers=headers,json=payload).status_code == 422
    assert len(client.get('/warehouse-assistant/api/supplier-orders',headers=headers).json()['orders']) == 1
    assert client.put('/warehouse-assistant/api/table-preferences/automatic_preorders',headers=headers,
                      json={'visible_columns':['supplier','delivery_date']}).status_code == 200


def test_manual_order_delete_requires_confirmation_retains_history_and_owner_scope(client, settings):
    headers=_session(settings,'Admin','Admin')
    imported=client.post('/warehouse-assistant/api/snapshots/import',headers=headers,
        files={'file':('inventory.xlsx',_inventory_workbook())}).json()
    created=client.post('/warehouse-assistant/api/supplier-orders',headers=headers,json=dict(
        snapshot_id=imported['id'],warehouse='karaj',lines=[dict(product_code='4014',quantity=22)])).json()['orders'][0]
    path=f"/warehouse-assistant/api/supplier-orders/{created['id']}"
    with warehouse_service.warehouse_connection(settings) as conn:
        snapshot=[tuple(r) for r in conn.execute('SELECT * FROM warehouse_snapshot_items')]
    with pytest.raises(warehouse_service.WarehouseAssistantError):
        warehouse_service.delete_supplier_order(settings,'different-user',created['id'])
    assert not warehouse_service.list_supplier_orders(settings,'different-user',include_deleted=True)
    assert client.request('DELETE',path,headers=headers,json={}).status_code==400
    assert len(client.get('/warehouse-assistant/api/supplier-orders',headers=headers).json()['orders'])==1
    for _ in range(2):
        assert client.request('DELETE',path,headers=headers,json={'confirmed':True}).status_code==200
    assert client.get('/warehouse-assistant/api/supplier-orders',headers=headers).json()['orders']==[]
    archived=client.get('/warehouse-assistant/api/supplier-orders?include_deleted=true',headers=headers).json()['orders'][0]
    assert archived['deleted'] and archived['deletion']['deleted_by']=='Admin'
    assert archived['order_number']==created['order_number'] and archived['lines']==created['lines']
    assert client.get(path+'/document.xlsx',headers=headers).status_code==200
    with warehouse_service.warehouse_connection(settings) as conn:
        assert [tuple(r) for r in conn.execute('SELECT * FROM warehouse_snapshot_items')]==snapshot


def test_manual_order_edit_requires_explicit_receipt_review_override(client, settings, monkeypatch):
    headers = _session(settings, 'Admin', 'Admin')
    imported = client.post('/warehouse-assistant/api/snapshots/import', headers=headers,
        files={'file': ('inventory.xlsx', _coverage_workbook())}).json()
    created = client.post('/warehouse-assistant/api/supplier-orders', headers=headers, json={
        'snapshot_id': imported['id'], 'warehouse': 'karaj',
        'lines': [{'product_code': '4014', 'quantity': 22}],
    }).json()['orders'][0]
    monkeypatch.setattr(warehouse_service, 'pending_stock',
                        lambda conn, warehouse: ({'4020': 5}, {'4020'}))
    with pytest.raises(warehouse_service.WarehouseAssistantError):
        warehouse_service.add_supplier_order_lines(settings, 'Admin', created['id'],
            [{'product_code': '4020', 'quantity': 13}])
    edited = warehouse_service.add_supplier_order_lines(settings, 'Admin', created['id'],
        [{'product_code': '4020', 'quantity': 13}], receipt_review_confirmed=True,
        reason='کاربر با وجود رسید در انتظار بررسی، سفارش را تأیید کرد.')
    assert len(edited['lines']) == 2
    assert edited['receipt_review_override_by'] == 'Admin'
    assert edited['receipt_review_override_reason']
    assert edited['is_approved'] is False


def test_supplier_order_rejects_unknown_product(client, settings):
    headers = _session(settings, "Admin", "Admin")
    imported = client.post(
        "/warehouse-assistant/api/snapshots/import",
        headers=headers,
        files={"file": ("inventory.xlsx", _inventory_workbook())},
    ).json()

    response = client.post(
        "/warehouse-assistant/api/supplier-orders",
        headers=headers,
        json={
            "snapshot_id": imported["id"],
            "warehouse": "karaj",
            "lines": [{"product_code": "UNKNOWN", "quantity": 12}],
        },
    )

    assert response.status_code == 422
    assert "UNKNOWN" in response.json()["detail"]


class _FakeWarehouseCursor:
    def __init__(self, queries: list[str]):
        self.queries = queries
        self.description = []
        self._rows: list[tuple[object, ...]] = []

    def execute(self, sql: str):
        self.queries.append(sql)
        if 'WITH native_transfer_credits AS' in sql or sql.startswith('SET TRANSACTION ISOLATION LEVEL'):
            columns = ['voucher_id'];self._rows = []
        elif "ICA.TblSupInvoiceHdr AS purchase_header" in sql:
            columns = [
                "GoodsRef", "ApproximatePrice", "PurchaseDate", "InvoiceId",
            ]
            self._rows = [(1, 75, "1405/06/12", 9001)]
        elif "Inv.tblVocherHdr AS receipt" in sql:
            columns = [
                "StockDCRef", "SupplierRef", "SupplierName",
                "ManufacturerName", "BrandName", "ReceiptCount",
                "ReceiptQuantity", "FirstReceiptDate", "LastReceiptDate",
            ]
            current_date = jalali_business_date(tehran_now())
            self._rows = [
                (
                    1, 17, "تامین‌کننده نمونه", "تامین‌کننده نمونه",
                    "میسویک", 2, 120, current_date, current_date,
                )
            ]
        elif "FRU.StockGoodsModel" in sql:
            columns = [
                "GoodsRef", "GoodsCode", "GoodsName", "ConversionRate",
                "ManufacturerName", "BrandName", "ManufacturerGoodsCode",
                "Barcode1", "Barcode2", "DefaultGoodsBarcode",
                "GoodsBarcodeNameList", "StockDCRef", "StockDCName",
                "OnHandQty", "ReservedQty", "DamagedQty", "UnDeliveredQty",
                "OpenCustomerOrderQty", "UnconfirmedFreeInvoiceQty",
                "LegacyOpenOrderQty", "OpenOrderQty", "PendingSaleVoucherQty",
                "SalePrice", "ManufacturerPrice", "ConsumerPrice",
            ]
            self._rows = [
                (
                    1, "4014", "خمیر دندان آزمایشی", 12,
                    "تامین‌کننده نمونه", "میسویک", "MFG-4014",
                    "62600004014", "62600004015", "62600004014",
                    "62600004014, 62600004015", 1, "انبار مرکزي",
                    10, 20, 1, 0, 5, 3, 2, 10, 4, 100, 90, 120,
                ),
            ]
        elif "inv.vwGoodsCardex" in sql:
            columns = ["GoodsRef", "StockDCRef", "VocherDate", "NetMovementQty"]
            self._rows = []
        else:
            columns = [
                "GoodsRef", "StockDCRef", "ReportDate", "CustomerId",
                "CustomerCategoryId", "GrossOutQty",
                "ReturnQty", "NetOutQty", "AmiranNetOutQty",
                "OtherCustomerNetOutQty", "ExcludedSellerNetQty",
                "AllGrossSalesQty", "OnlineTransferOutQty",
            ]
            self._rows = [
                (
                    1, 1, jalali_business_date(tehran_now()),
                    101, 4, 60, 6, 54, 0, 54, 12, 72, 0,
                ),
                (
                    1, 1, jalali_business_date(tehran_now()),
                    None, None, 0, 0, 0, 0, 0, 0, 0, 30,
                ),
            ]
        self.description = [(column,) for column in columns]
        return self

    def fetchall(self):
        return self._rows


class _FakeWarehouseConnection:
    def __init__(self, queries: list[str]):
        self.queries = queries

    def cursor(self):
        return _FakeWarehouseCursor(self.queries)


def test_admin_syncs_inventory_and_outflow_from_varanegar_read_only(
    client, settings, monkeypatch
):
    headers = _session(settings, "Admin", "Admin")
    configured = replace(
        settings,
        sql_server="read-only-host",
        sql_username="warehouse_reader",
        sql_password="secret",
    )
    client.app.state.settings = configured
    queries: list[str] = []

    @contextmanager
    def fake_sql_connection(_settings):
        yield _FakeWarehouseConnection(queries)

    monkeypatch.setattr(warehouse_service, "sql_connection", fake_sql_connection)

    synced = client.post(
        "/warehouse-assistant/api/sync/varanegar",
        headers=headers,
        json={"period_days": 60},
    )

    assert synced.status_code == 201
    snapshot = synced.json()
    assert snapshot["source_kind"] == "varanegar"
    assert snapshot["period_days"] == 60
    assert snapshot["demand_basis"] == "net_sales_last_stock_window"
    assert snapshot["product_count"] == 1
    assert snapshot["item_count"] == 1
    assert snapshot["varanegar_write"] is False
    assert "FRU.StockGoodsModel" in snapshot["sources"]

    suggestions = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "reorder_coverage_days": 29,
            "target_days": 30,
            "safety_days": 0,
            "period_days": 60,
        },
    )
    assert suggestions.status_code == 200
    policy = suggestions.json()["parameters"]["demand_outlier_policy"]
    assert policy["protected_customer_category_id"] is None
    assert policy["protected_customer_group_id"] == 2
    assert policy["protected_customer_membership_source"] == "GNR.tblCust.CustGroupRef"
    item = suggestions.json()["items"][0]
    assert item["stock"] == 10
    assert item["reserved"] == 20
    assert item["open_order_qty"] == 10
    assert item["open_customer_order_qty"] == 5
    assert item["unconfirmed_free_invoice_qty"] == 3
    assert item["legacy_open_order_qty"] == 2
    assert item["pending_sale_voucher_qty"] == 4
    assert item["effective_procurement_qty"] == 20
    assert item["available_quantity"] == 20
    assert item["period_out"] == 84
    assert item["raw_period_out"] == 84
    assert item["gross_out"] == 60
    assert item["period_return"] == 6
    assert item["online_transfer_out_qty"] == 30
    assert item["excluded_seller_qty"] == 12
    assert item["sales_rate_days"] == 60
    assert item["stockout_days"] == 0
    assert item["days_since_last_stock"] == 0
    assert item["ordering_cycle_active"] is True
    assert item["sale_price"] == 100
    assert item["manufacturer_price"] == 90
    assert item["consumer_price"] == 120
    assert item["approximate_price"] == 75
    assert item["estimated_value"] == 1800
    assert item["sales_window_end"] == jalali_business_date(tehran_now())
    assert item["manufacturer_product_code"] == "MFG-4014"
    assert item["barcode"] == "62600004014"
    assert item["effective_cartons"] == 1

    assert len(queries) == 7
    assert queries[0] == 'SET TRANSACTION ISOLATION LEVEL SNAPSHOT'
    native_query = next(q for q in queries if 'WITH native_transfer_credits AS' in q)
    assert 'NOLOCK' not in native_query
    assert 'R.DocRef=C.voucher_id' in native_query
    queries = [q for q in queries if q != native_query and not q.startswith('SET TRANSACTION ISOLATION LEVEL')]
    assert all(query.lstrip().upper().startswith(("SELECT", "WITH")) for query in queries)
    forbidden = ("INSERT ", "UPDATE ", "DELETE ", "MERGE ", "EXEC ", "CREATE ", "ALTER ", "DROP ")
    assert all(not any(token in query.upper() for token in forbidden) for query in queries)
    assert "FRU.StockGoodsModel" in queries[0]
    assert "COALESCE(sale.DealerId, -1) <> 7" in queries[1]
    assert "LEFT JOIN GNR.tblCust AS customer ON customer.ID = sale.CustomerId" in queries[1]
    assert "customer.CustGroupRef = 2" in queries[1]
    assert "CustomerCategoryId = 3" not in queries[1]
    assert "AmiranNetOutQty" in queries[1]
    assert "OtherCustomerNetOutQty" in queries[1]
    assert "sale.CustomerId" in queries[1]
    assert "sale.CustomerCategoryId" in queries[1]
    assert "VocherTypeCode = 65" in queries[1]
    assert "TStockDCRef = 7" in queries[1]
    assert "OnlineTransferOutQty" in queries[1]
    assert "FROM NGT.ContractPrices AS contract_price" in queries[0]
    assert "TRY_CONVERT(int, contract_price.OrderTypeRef) IN (2, 10, 13)" in queries[0]
    assert "WHEN 1 THEN 2" in queries[0]
    assert "WHEN 2 THEN 10" in queries[0]
    assert "WHEN 9 THEN 13" in queries[0]
    assert "ISNULL(contract_price.IsRemoved, 0) = 0" in queries[0]
    assert "contract_price.StartDate <=" in queries[0]
    assert "NULLIF(contract_price.CustRef, '') IS NULL" in queries[0]
    assert "ISNULL(contract_price.MinQty, 0) = 0" in queries[0]
    assert "ISNULL(contract_price.MaxQty, 0) = 0" in queries[0]
    assert "source_price.ManufacturerPrice" in queries[0]
    assert "ICA.TblSupInvoiceHdr" not in queries[0]
    purchase_query = next(
        query for query in queries if "ICA.TblSupInvoiceHdr AS purchase_header" in query
    )
    assert "ICA.TblSupInvoiceItm" in purchase_query
    assert "purchase_header.Status = 1" in purchase_query
    assert "purchase_header.ConfirmDate IS NOT NULL" in purchase_query
    assert "purchase_item.Price" in purchase_query
    assert "model.SalePrice" not in queries[0]
    assert "inv.vwGoodsCardex" in queries[2]
    assert "Inv.tblVocherHdr AS receipt" in queries[3]

    cached = client.post(
        "/warehouse-assistant/api/sync/varanegar",
        headers=headers,
        json={"period_days": 60},
    )
    assert cached.status_code == 200
    assert cached.json()["duplicate"] is True
    assert len(queries) == 11  # cached refresh adds SNAPSHOT setup and native in-transit read
    assert sum(
        "ICA.TblSupInvoiceHdr AS purchase_header" in query for query in queries
    ) == 1

    with sqlite3.connect(warehouse_service.warehouse_database_path(configured)) as conn:
        cache = conn.execute(
            "SELECT approximate_price FROM warehouse_purchase_price_cache "
            "WHERE goods_ref=1"
        ).fetchone()
        cache_state = conn.execute(
            "SELECT source_row_count FROM warehouse_purchase_price_cache_state "
            "WHERE id=1"
        ).fetchone()
    assert cache == (75.0,)
    assert cache_state == (1,)


def test_demand_spike_is_capped_but_amiran_sales_are_fully_preserved():
    history_dates = [f"day-{index:02d}" for index in range(60)]
    trends = []
    for index, report_date in enumerate(history_dates):
        other = 1000 if index == 59 else 10
        amiran = 500 if index == 59 else 0
        net = other + amiran
        trends.append(
            {
                "ReportDate": report_date,
                "CustomerId": 10,
                "CustomerCategoryId": 4,
                "GrossOutQty": other,
                "ReturnQty": 0,
                "NetOutQty": other,
                "AmiranNetOutQty": 0,
                "OtherCustomerNetOutQty": other,
                "ExcludedSellerNetQty": 0,
                "AllGrossSalesQty": other,
            }
        )
        if amiran:
            trends.append(
                {
                    "ReportDate": report_date,
                    "CustomerId": 20,
                    "CustomerCategoryId": 3,
                    "GrossOutQty": amiran,
                    "ReturnQty": 0,
                    "NetOutQty": amiran,
                    "AmiranNetOutQty": amiran,
                    "OtherCustomerNetOutQty": 0,
                    "ExcludedSellerNetQty": 0,
                    "AllGrossSalesQty": amiran,
                }
            )

    demand = warehouse_service._last_stock_window_metrics(
        current_on_hand=100000,
        history_dates=history_dates,
        daily_movement={},
        trend_rows=trends,
        period_days=60,
    )

    assert demand["raw_net_out"] == 2090
    assert demand["daily_cap"] == 10
    assert demand["exceptional_out"] == 990
    assert demand["amiran_net_out"] == 500
    assert demand["net_out"] == 1100
    assert demand["anomaly_days"] == 1
    assert demand["recurring_pattern"] is False
    assert demand["recurring_pattern_days"] == 0


def test_online_warehouse_transfer_is_fully_counted_as_demand():
    history_dates = [f"day-{index:02d}" for index in range(60)]
    trends = [
        {
            "ReportDate": report_date,
            "CustomerId": 10,
            "CustomerCategoryId": 4,
            "GrossOutQty": 10,
            "ReturnQty": 0,
            "NetOutQty": 10,
            "AmiranNetOutQty": 0,
            "OtherCustomerNetOutQty": 10,
            "ExcludedSellerNetQty": 0,
            "AllGrossSalesQty": 10,
            "OnlineTransferOutQty": 0,
        }
        for report_date in history_dates
    ]
    trends.append(
        {
            "ReportDate": history_dates[-1],
            "CustomerId": None,
            "CustomerCategoryId": None,
            "GrossOutQty": 0,
            "ReturnQty": 0,
            "NetOutQty": 0,
            "AmiranNetOutQty": 0,
            "OtherCustomerNetOutQty": 0,
            "ExcludedSellerNetQty": 0,
            "AllGrossSalesQty": 0,
            "OnlineTransferOutQty": 1000,
        }
    )

    demand = warehouse_service._last_stock_window_metrics(
        current_on_hand=100000,
        history_dates=history_dates,
        daily_movement={},
        trend_rows=trends,
        period_days=60,
    )

    assert demand["gross_out"] == 600
    assert demand["online_transfer_out"] == 1000
    assert demand["raw_net_out"] == 1600
    assert demand["net_out"] == 1600
    assert demand["anomaly_days"] == 0


def test_customer_spike_requires_large_quantity_and_eighty_percent_daily_share():
    history_dates = [f"day-{index:02d}" for index in range(60)]
    trends = []
    for index, report_date in enumerate(history_dates[:57]):
        trends.append(
            {
                "ReportDate": report_date,
                "CustomerId": 1,
                "CustomerCategoryId": 4,
                "GrossOutQty": 10,
                "ReturnQty": 0,
                "NetOutQty": 10,
                "AmiranNetOutQty": 0,
                "OtherCustomerNetOutQty": 10,
                "ExcludedSellerNetQty": 0,
                "AllGrossSalesQty": 10,
            }
        )
    for customer_id, quantity in ((2, 1000), (3, 300)):
        trends.append(
            {
                "ReportDate": history_dates[57],
                "CustomerId": customer_id,
                "CustomerCategoryId": 4,
                "GrossOutQty": quantity,
                "ReturnQty": 0,
                "NetOutQty": quantity,
                "AmiranNetOutQty": 0,
                "OtherCustomerNetOutQty": quantity,
                "ExcludedSellerNetQty": 0,
                "AllGrossSalesQty": quantity,
            }
        )
    trends.append(
        {
            "ReportDate": history_dates[58],
            "CustomerId": 4,
            "CustomerCategoryId": 4,
            "GrossOutQty": 20,
            "ReturnQty": 0,
            "NetOutQty": 20,
            "AmiranNetOutQty": 0,
            "OtherCustomerNetOutQty": 20,
            "ExcludedSellerNetQty": 0,
            "AllGrossSalesQty": 20,
        }
    )
    for customer_id, quantity in ((5, 1000), (6, 10)):
        trends.append(
            {
                "ReportDate": history_dates[59],
                "CustomerId": customer_id,
                "CustomerCategoryId": 4,
                "GrossOutQty": quantity,
                "ReturnQty": 0,
                "NetOutQty": quantity,
                "AmiranNetOutQty": 0,
                "OtherCustomerNetOutQty": quantity,
                "ExcludedSellerNetQty": 0,
                "AllGrossSalesQty": quantity,
            }
        )

    demand = warehouse_service._last_stock_window_metrics(
        current_on_hand=100000,
        history_dates=history_dates,
        daily_movement={},
        trend_rows=trends,
        period_days=60,
    )

    assert demand["raw_net_out"] == 2900
    assert demand["net_out"] == 1910
    assert demand["exceptional_out"] == 990
    assert demand["anomaly_days"] == 1


def test_frequent_spikes_from_one_customer_are_not_treated_as_normal_demand():
    history_dates = [f"day-{index:02d}" for index in range(49)]
    trends = []
    for index, report_date in enumerate(history_dates):
        demand = 1000 if index >= 44 else 10
        trends.append(
            {
                "ReportDate": report_date,
                "CustomerId": 99,
                "CustomerCategoryId": 4,
                "GrossOutQty": demand,
                "ReturnQty": 0,
                "NetOutQty": demand,
                "AmiranNetOutQty": 0,
                "OtherCustomerNetOutQty": demand,
                "ExcludedSellerNetQty": 0,
                "AllGrossSalesQty": demand,
            }
        )

    demand = warehouse_service._last_stock_window_metrics(
        current_on_hand=100000,
        history_dates=history_dates,
        daily_movement={},
        trend_rows=trends,
        period_days=49,
    )

    assert demand["raw_net_out"] == 5440
    assert demand["net_out"] < demand["raw_net_out"]
    assert demand["exceptional_out"] > 0
    assert demand["anomaly_days"] == 5
    assert demand["recurring_pattern"] is False
    assert demand["recurring_pattern_days"] == 0


def test_frequent_spikes_across_multiple_customers_are_treated_as_normal_demand():
    history_dates = [f"day-{index:02d}" for index in range(49)]
    trends = []
    for index, report_date in enumerate(history_dates):
        demand = 1000 if index >= 44 else 10
        trends.append(
            {
                "ReportDate": report_date,
                "CustomerId": 100 + index if index >= 44 else 1,
                "CustomerCategoryId": 4,
                "GrossOutQty": demand,
                "ReturnQty": 0,
                "NetOutQty": demand,
                "AmiranNetOutQty": 0,
                "OtherCustomerNetOutQty": demand,
                "ExcludedSellerNetQty": 0,
                "AllGrossSalesQty": demand,
            }
        )

    demand = warehouse_service._last_stock_window_metrics(
        current_on_hand=100000,
        history_dates=history_dates,
        daily_movement={},
        trend_rows=trends,
        period_days=49,
    )

    assert demand["raw_net_out"] == 5440
    assert demand["net_out"] == 5440
    assert demand["exceptional_out"] == 0
    assert demand["anomaly_days"] == 0
    assert demand["recurring_pattern"] is True
    assert demand["recurring_pattern_days"] == 5


def test_exactly_ten_percent_spike_days_are_still_adjusted():
    history_dates = [f"day-{index:02d}" for index in range(40)]
    trends = []
    for index, report_date in enumerate(history_dates):
        demand = 1000 if index >= 36 else 10
        trends.append(
            {
                "ReportDate": report_date,
                "CustomerId": 100 + index if index >= 36 else 1,
                "CustomerCategoryId": 4,
                "GrossOutQty": demand,
                "ReturnQty": 0,
                "NetOutQty": demand,
                "AmiranNetOutQty": 0,
                "OtherCustomerNetOutQty": demand,
                "ExcludedSellerNetQty": 0,
                "AllGrossSalesQty": demand,
            }
        )

    demand = warehouse_service._last_stock_window_metrics(
        current_on_hand=100000,
        history_dates=history_dates,
        daily_movement={},
        trend_rows=trends,
        period_days=40,
    )

    assert demand["net_out"] < demand["raw_net_out"]
    assert demand["exceptional_out"] > 0
    assert demand["anomaly_days"] == 4
    assert demand["recurring_pattern"] is False
    assert demand["recurring_pattern_days"] == 0


def test_sales_window_is_anchored_to_the_last_day_product_was_available(
    client, settings, monkeypatch
):
    headers = _session(settings, "Admin", "Admin")
    configured = replace(
        settings,
        sql_server="read-only-host",
        sql_username="warehouse_reader",
        sql_password="secret",
    )
    client.app.state.settings = configured
    last_in_stock_date = jalali_business_date(tehran_now() - timedelta(days=20))
    expected_window_start = jalali_business_date(
        tehran_now() - timedelta(days=79)
    )
    older_sale_date = jalali_business_date(tehran_now() - timedelta(days=70))
    queries: list[str] = []

    class StockoutCursor(_FakeWarehouseCursor):
        def execute(self, sql: str):
            super().execute(sql)
            if "ranked_stock AS ranked" in sql:
                row = list(self._rows[0])
                row[13] = 0
                row[14] = 0
                row[17] = 0
                row[18] = 0
                row[19] = 0
                row[20] = 0
                self._rows = [tuple(row)]
            elif "inv.vwGoodsCardex" in sql:
                self._rows = [(1, 1, last_in_stock_date, -110)]
            elif "dbo.SalesReviewFast" in sql:
                self._rows = [
                    (
                        1, 1, older_sale_date, 70, 10,
                        70, 10, 60, 0, 60, 0, 70, 0,
                    ),
                    (
                        1, 1, last_in_stock_date, 40, 10,
                        40, 10, 30, 0, 30, 20, 60, 0,
                    ),
                ]
            return self

    class StockoutConnection:
        def cursor(self):
            return StockoutCursor(queries)

    @contextmanager
    def fake_sql_connection(_settings):
        yield StockoutConnection()

    monkeypatch.setattr(warehouse_service, "sql_connection", fake_sql_connection)
    synced = client.post(
        "/warehouse-assistant/api/sync/varanegar",
        headers=headers,
        json={"period_days": 60},
    )
    assert synced.status_code == 201

    response = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "reorder_coverage_days": 15,
            "target_days": 30,
            "period_days": 60,
        },
    )
    assert response.status_code == 200
    item = response.json()["items"][0]
    assert response.json()["parameters"]["excluded_seller"] == {
        "dealer_id": 7,
        "name": "ایمان شریف پور",
    }
    assert item["gross_out"] == 110
    assert item["period_return"] == 20
    assert item["period_out"] == 90
    assert item["excluded_seller_qty"] == 20
    assert item["sales_rate_days"] == 60
    assert item["stockout_days"] == 0
    assert item["last_in_stock_date"] == last_in_stock_date
    assert item["days_since_last_stock"] == 20
    assert item["ordering_cycle_active"] is True
    assert item["sales_window_start"] == expected_window_start
    assert item["sales_window_end"] == last_in_stock_date
    assert item["average_daily_out"] == 1.5
    assert item["suggested_quantity"] == 48


def test_product_last_available_more_than_90_days_ago_is_out_of_order_cycle(
    client, settings, monkeypatch
):
    headers = _session(settings, "Admin", "Admin")
    configured = replace(
        settings,
        sql_server="read-only-host",
        sql_username="warehouse_reader",
        sql_password="secret",
    )
    client.app.state.settings = configured
    stale_date = jalali_business_date(tehran_now() - timedelta(days=91))
    queries: list[str] = []

    class StaleCursor(_FakeWarehouseCursor):
        def execute(self, sql: str):
            super().execute(sql)
            if "ranked_stock AS ranked" in sql:
                row = list(self._rows[0])
                row[13] = 0
                row[14] = 0
                row[17] = 0
                row[18] = 0
                row[19] = 0
                row[20] = 0
                self._rows = [tuple(row)]
            elif "inv.vwGoodsCardex" in sql:
                self._rows = [(1, 1, stale_date, -110)]
            elif "dbo.SalesReviewFast" in sql:
                self._rows = [
                    (1, 1, stale_date, 101, 4, 100, 10, 90, 0, 90, 0, 100, 0)
                ]
            return self

    class StaleConnection:
        def cursor(self):
            return StaleCursor(queries)

    @contextmanager
    def fake_sql_connection(_settings):
        yield StaleConnection()

    monkeypatch.setattr(warehouse_service, "sql_connection", fake_sql_connection)
    synced = client.post(
        "/warehouse-assistant/api/sync/varanegar",
        headers=headers,
        json={"period_days": 60},
    )
    assert synced.status_code == 201

    response = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "reorder_coverage_days": 15,
            "target_days": 30,
            "period_days": 60,
        },
    )
    assert response.status_code == 200
    assert response.json()["items"] == []
    assert response.json()["summary"]["excluded_stale_items"] == 1
    assert response.json()["parameters"]["stale_after_days"] == 90

    direct_order = client.post(
        "/warehouse-assistant/api/supplier-orders",
        headers=headers,
        json={
            "snapshot_id": synced.json()["id"],
            "warehouse": "karaj",
            "lines": [{"product_code": "4014", "quantity": 12}],
        },
    )
    assert direct_order.status_code == 422
    assert "خارج از چرخه سفارش" in direct_order.json()["detail"]

    inventory = client.get(
        "/warehouse-assistant/api/inventory",
        headers=headers,
        params={"warehouse": "karaj", "search": "4014"},
    )
    assert inventory.status_code == 200
    inventory_item = inventory.json()["items"][0]
    assert inventory_item["system_ordering_cycle_active"] is False
    assert inventory_item["order_cycle_forced_active"] is False
    assert inventory_item["ordering_cycle_active"] is False

    enabled = client.put(
        "/warehouse-assistant/api/inventory/order-cycle",
        headers=headers,
        json={
            "warehouse": "karaj",
            "product_code": "4014",
            "mode": "force_active",
        },
    )
    assert enabled.status_code == 200
    assert enabled.json()["system_ordering_cycle_active"] is False
    assert enabled.json()["order_cycle_forced_active"] is True
    assert enabled.json()["ordering_cycle_active"] is True

    included = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "reorder_coverage_days": 15,
            "target_days": 30,
            "period_days": 60,
        },
    )
    assert included.status_code == 200
    assert included.json()["items"][0]["product_code"] == "4014"
    assert included.json()["items"][0]["system_ordering_cycle_active"] is False
    assert included.json()["items"][0]["order_cycle_forced_active"] is True
    assert included.json()["items"][0]["ordering_cycle_active"] is True
    assert included.json()["summary"]["excluded_stale_items"] == 0

    restored = client.put(
        "/warehouse-assistant/api/inventory/order-cycle",
        headers=headers,
        json={
            "warehouse": "karaj",
            "product_code": "4014",
            "mode": "system",
        },
    )
    assert restored.status_code == 200
    assert restored.json()["order_cycle_forced_active"] is False
    assert restored.json()["ordering_cycle_active"] is False


def test_ordering_catalog_returns_manufacturers_with_their_brands(
    client, settings, monkeypatch
):
    headers = _session(settings, "Admin", "Admin")
    configured = replace(
        settings,
        sql_server="read-only-host",
        sql_username="warehouse_reader",
        sql_password="secret",
    )
    client.app.state.settings = configured
    queries: list[str] = []

    @contextmanager
    def fake_sql_connection(_settings):
        yield _FakeWarehouseConnection(queries)

    monkeypatch.setattr(warehouse_service, "sql_connection", fake_sql_connection)
    assert client.post(
        "/warehouse-assistant/api/sync/varanegar",
        headers=headers,
        json={"period_days": 60},
    ).status_code == 201

    response = client.get("/warehouse-assistant/api/catalog", headers=headers)

    assert response.status_code == 200
    assert response.json()["manufacturers"] == [
        {
            "name": "تامین‌کننده نمونه",
            "product_count": 1,
            "brands": [{"name": "میسویک", "product_count": 1}],
        }
    ]


def test_receipt_supply_scope_filters_every_ordering_path_and_preserves_override(
    client, settings, monkeypatch
):
    headers = _session(settings, "Admin", "Admin")
    configured = replace(
        settings,
        sql_server="read-only-host",
        sql_username="warehouse_reader",
        sql_password="secret",
    )
    client.app.state.settings = configured
    queries: list[str] = []

    @contextmanager
    def fake_sql_connection(_settings):
        yield _FakeWarehouseConnection(queries)

    monkeypatch.setattr(warehouse_service, "sql_connection", fake_sql_connection)
    synced = client.post(
        "/warehouse-assistant/api/sync/varanegar",
        headers=headers,
        json={"period_days": 60},
    )
    assert synced.status_code == 201

    scope = client.get("/warehouse-assistant/api/supply-scope", headers=headers)
    assert scope.status_code == 200
    assert scope.json()["strict_enabled"] is True
    row = scope.json()["rows"][0]
    assert row["warehouse_code"] == "karaj"
    assert row["supplier"] == "تامین‌کننده نمونه"
    assert row["brand"] == "میسویک"
    assert row["enabled"] is True
    assert row["observed_in_recent_receipts"] is True
    assert row["receipt_count"] == 2

    disabled = client.put(
        f"/warehouse-assistant/api/supply-scope/{row['id']}",
        headers=headers,
        json={"enabled": False},
    )
    assert disabled.status_code == 200
    assert disabled.json()["enabled"] is False
    assert disabled.json()["is_user_override"] is True

    catalog = client.get(
        "/warehouse-assistant/api/catalog",
        headers=headers,
        params={"warehouse": "karaj"},
    )
    assert catalog.status_code == 200
    assert catalog.json()["manufacturers"] == []

    suggestions = client.get(
        "/warehouse-assistant/api/suggestions",
        headers=headers,
        params={
            "warehouse": "karaj",
            "reorder_coverage_days": 29,
            "target_days": 30,
        },
    )
    assert suggestions.status_code == 200
    assert suggestions.json()["items"] == []

    direct_order = client.post(
        "/warehouse-assistant/api/supplier-orders",
        headers=headers,
        json={
            "snapshot_id": synced.json()["id"],
            "warehouse": "karaj",
            "lines": [{"product_code": "4014", "quantity": 12}],
        },
    )
    assert direct_order.status_code == 422
    assert "برای انبار فعال نیست" in direct_order.json()["detail"]

    refreshed = client.post(
        "/warehouse-assistant/api/supply-scope/refresh", headers=headers
    )
    assert refreshed.status_code == 200
    scope_after_refresh = client.get(
        "/warehouse-assistant/api/supply-scope", headers=headers
    ).json()["rows"][0]
    assert scope_after_refresh["enabled"] is False
    assert scope_after_refresh["is_user_override"] is True

    enabled = client.put(
        "/warehouse-assistant/api/supply-scope/supplier",
        headers=headers,
        json={
            "warehouse": "karaj",
            "supplier": "تامین‌کننده نمونه",
            "enabled": True,
        },
    )
    assert enabled.status_code == 200
    assert enabled.json()["updated_count"] == 1


def test_inventory_information_exposes_varanegar_stock_components(
    client, settings, monkeypatch
):
    headers = _session(settings, "Admin", "Admin")
    configured = replace(
        settings,
        sql_server="read-only-host",
        sql_username="warehouse_reader",
        sql_password="secret",
    )
    client.app.state.settings = configured
    queries: list[str] = []

    @contextmanager
    def fake_sql_connection(_settings):
        yield _FakeWarehouseConnection(queries)

    monkeypatch.setattr(warehouse_service, "sql_connection", fake_sql_connection)
    assert client.post(
        "/warehouse-assistant/api/sync/varanegar",
        headers=headers,
        json={"period_days": 60},
    ).status_code == 201

    response = client.get(
        "/warehouse-assistant/api/inventory",
        headers=headers,
        params={"warehouse": "karaj", "search": "4014"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["summary"]["total_items"] == 1
    assert body["summary"]["effective_procurement_qty"] == 20
    assert body["summary"]["pending_sale_voucher_qty"] == 4
    assert body["formula"] == "on_hand + reserved - open_order + in_transit + pending_receipt"
    assert body["pending_sale_voucher_already_in_on_hand"] is True
    assert body["items"][0] == {
        "warehouse": "karaj",
        "warehouse_name": "انبار مرکزی کرج",
        "product_code": "4014",
        "product_name": "خمیر دندان آزمایشی",
        "manufacturer": "تامین‌کننده نمونه",
        "brand": "میسویک",
        "group_level3": "",
        "manufacturer_product_code": "MFG-4014",
        "tax_rate": None,
        "tax_status": "unknown",
        "tax_updated_at": None,
        "barcode": "62600004014",
        "barcode2": "62600004015",
        "barcode_list": "62600004014, 62600004015",
        "conversion_rate": 12,
        "on_hand_qty": 10,
        "reserved_qty": 20,
        "damaged_qty": 1,
        "undelivered_qty": 0,
        "open_customer_order_qty": 5,
        "unconfirmed_free_invoice_qty": 3,
        "legacy_open_order_qty": 2,
        "open_order_qty": 10,
        "pending_sale_voucher_qty": 4,
        "owned_procurement_qty": 30,
        "effective_procurement_qty": 20,
        "physical_procurement_qty": 20,
        "in_transit_qty": 0,
        "pending_receipt_qty": 0,
        "receipt_review_required": False,
        "period_out_qty": 84,
        "raw_period_out_qty": 84,
        "amiran_period_out_qty": 0,
        "exceptional_period_out_qty": 0,
        "demand_anomaly_days": 0,
        "demand_daily_cap": None,
        "demand_recurring_pattern": False,
        "demand_recurring_pattern_days": 0,
            "gross_out_qty": 60,
            "period_return_qty": 6,
            "sales_rate_days": 60,
            "online_transfer_out_qty": 30,
            "purchase_price_validation": {
                "message": "",
                "on_date": jalali_business_date(tehran_now()),
                "required_bases": [],
                "status": "not_required",
            },
        "sale_price": 100,
        "manufacturer_price": 90,
        "consumer_price": 120,
        "last_in_stock_date": jalali_business_date(tehran_now()),
        "days_since_last_stock": 0,
        "sales_window_start": jalali_business_date(
            tehran_now() - timedelta(days=59)
        ),
        "sales_window_end": jalali_business_date(tehran_now()),
        "system_ordering_cycle_active": True,
        "order_cycle_forced_active": False,
        "order_cycle_forced_inactive": False,
        "ordering_cycle_active": True,
    }

    filtered = client.get(
        "/warehouse-assistant/api/inventory",
        headers=headers,
        params={
            "filters": '{"brand":"میسویک","barcode":"62600004014","manufacturer_product_code":"MFG-4014"}'
        },
    )
    assert filtered.status_code == 200
    assert filtered.json()["summary"]["total_items"] == 1

    missing = client.get(
        "/warehouse-assistant/api/inventory",
        headers=headers,
        params={"filters": '{"brand":"برند دیگر"}'},
    )
    assert missing.status_code == 200
    assert missing.json()["summary"]["total_items"] == 0


def test_inventory_views_are_named_user_scoped_and_have_one_default(client, settings):
    first_user = _session(settings, "warehouse.views.one", "Admin")
    second_user = _session(settings, "warehouse.views.two", "Admin")

    first = client.post(
        "/warehouse-assistant/api/inventory/views",
        headers=first_user,
        json={
            "name": "نمایش خرید تهران",
            "is_default": True,
            "visible_columns": [
                "product_code", "product_name", "warehouse_name", "manufacturer",
                "brand", "barcode", "manufacturer_product_code",
                "effective_procurement_qty",
            ],
            "filters": {"warehouse_name": "تهران", "brand": "میسویک"},
        },
    )
    assert first.status_code == 201
    first_id = first.json()["id"]

    updated_same_name = client.post(
        "/warehouse-assistant/api/inventory/views",
        headers=first_user,
        json={
            "name": "نمایش خرید تهران",
            "is_default": False,
            "visible_columns": ["product_code", "brand", "barcode"],
            "filters": {"brand": "میسویک"},
        },
    )
    assert updated_same_name.status_code == 200
    assert updated_same_name.json()["id"] == first_id

    second_default = client.post(
        "/warehouse-assistant/api/inventory/views",
        headers=first_user,
        json={
            "name": "طرح خلاصه",
            "is_default": True,
            "visible_columns": ["product_code", "product_name", "on_hand_qty"],
            "filters": {},
        },
    )
    assert second_default.status_code == 201

    other = client.post(
        "/warehouse-assistant/api/inventory/views",
        headers=second_user,
        json={
            "name": "طرح شخص دیگر",
            "is_default": True,
            "visible_columns": ["product_code", "product_name"],
            "filters": {},
        },
    )
    assert other.status_code == 201

    first_views = client.get(
        "/warehouse-assistant/api/inventory/views", headers=first_user
    ).json()["views"]
    assert [item["name"] for item in first_views] == ["طرح خلاصه", "نمایش خرید تهران"]
    assert [item["is_default"] for item in first_views] == [True, False]
    assert len(first_views) == 2

    second_views = client.get(
        "/warehouse-assistant/api/inventory/views", headers=second_user
    ).json()["views"]
    assert [item["name"] for item in second_views] == ["طرح شخص دیگر"]

    cannot_delete_other_users_view = client.delete(
        f"/warehouse-assistant/api/inventory/views/{first_id}", headers=second_user
    )
    assert cannot_delete_other_users_view.status_code == 404
