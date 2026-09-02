import json
from io import BytesIO

from openpyxl import load_workbook

from app.database import sqlite_connection
from app.excel_service import build_report_workbook


def test_report_workbook_is_rtl_formatted_and_formula_safe():
    content = build_report_workbook(
        ["مشتری", "فروش"],
        [["جامبو", 1234567], ["=HYPERLINK(\"bad\")", 50]],
        "گزارش فروش",
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["گزارش"]

    assert sheet.sheet_view.rightToLeft is True
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].value == "مشتری"
    assert sheet["B2"].value == 1234567
    assert sheet["A3"].value.startswith("'=")
    assert workbook["اطلاعات"]["B1"].value == "گزارش فروش"


def test_latest_chat_report_can_be_downloaded_as_xlsx(client, auth, settings):
    response_payload = {
        "conversation_id": "excel-report",
        "answer": "گزارش آماده است.",
        "columns": ["BusinessDate", "NetSales"],
        "rows": [["1405/05/18", 1234567]],
        "row_count": 1,
    }
    with sqlite_connection(settings.sqlite_path) as conn:
        conn.execute(
            """INSERT INTO chat_messages
               (conversation_id, role, content, sources_json, response_json, created_at)
               VALUES (?, 'assistant', ?, '[]', ?, '2026-08-09T12:00:00')""",
            ("excel-report", "گزارش آماده است.", json.dumps(response_payload, ensure_ascii=False)),
        )

    response = client.get("/chat/conversations/excel-report/export.xlsx", headers=auth)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook["گزارش"]["A2"].value == "1405/05/18"
    assert workbook["گزارش"]["B2"].value == 1234567


def test_explicit_excel_request_can_download_answer_only_report(client, auth, settings):
    response_payload = {
        "conversation_id": "excel-answer-only",
        "answer": "خلاصه فروش امروز\nفروش خالص: ۱۲۳٬۴۵۶ ریال",
        "row_count": 0,
    }
    with sqlite_connection(settings.sqlite_path) as conn:
        conn.execute(
            """INSERT INTO chat_messages
               (conversation_id, role, content, sources_json, response_json, created_at)
               VALUES (?, 'assistant', ?, '[]', ?, '2026-08-09T12:00:00')""",
            (
                "excel-answer-only",
                response_payload["answer"],
                json.dumps(response_payload, ensure_ascii=False),
            ),
        )

    response = client.get(
        "/chat/conversations/excel-answer-only/export.xlsx",
        headers=auth,
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook["گزارش"]["A1"].value == "گزارش"
    assert workbook["گزارش"]["A2"].value == "خلاصه فروش امروز"
    assert "۱۲۳" in workbook["گزارش"]["A3"].value
