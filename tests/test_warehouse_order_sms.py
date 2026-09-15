from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import warehouse_order_sms as sms
from app import warehouse_assistant_service as service
from app import warehouse_fulfillment as fulfillment
from app.routes import warehouse_assistant as routes
from app.routes.dependencies import require_session_user


@pytest.fixture
def store(tmp_path, monkeypatch):
    settings = SimpleNamespace(sqlite_path=Path(tmp_path) / "main.db")
    service.init_warehouse_store(settings)
    with service.warehouse_connection(settings) as conn:
        conn.execute("""INSERT INTO warehouse_snapshots
            (id,source_filename,source_sheet,content_sha256,product_count,item_count,imported_by,imported_at)
            VALUES(1,'test','test','sms-test',1,1,'tester','now')""")
        conn.execute("""INSERT INTO supplier_orders
            (id,order_number,snapshot_id,warehouse_code,warehouse_name,supplier,total_quantity,created_by,created_at)
            VALUES(1,'SUP-TEST',1,'karaj','انبار کرج','تأمین‌کننده',24,'tester','now')""")
        conn.execute("""INSERT INTO supplier_order_lines
            (order_id,product_code,product_name,brand,conversion_rate,requested_quantity,
             order_quantity,cartons,manufacturer_price,consumer_price,buy_price,estimated_value,note)
            VALUES(1,'00123','کالای تست','برند',12,24,24,2,90,120,70,1680,'توضیح')""")
    monkeypatch.setattr(sms, "_provider_config", lambda: {
        "username": "test", "password": "secret", "source": "9821", "send_url": "https://example.test"
    })
    transport = []
    monkeypatch.setattr(sms, "_send_provider", lambda config, mobile, message: transport.append((mobile, message)) or "42")
    return settings, transport


def _order(settings):
    return service.get_supplier_order(settings, 1, "tester")


def test_sms_stores_exact_workbook_and_short_lived_download(store):
    settings, transport = store
    order = _order(settings)
    content = sms.supplier_order_workbook(order)
    result = sms.send_document_link(
        settings, "tester", document_kind="supplier_order", document_id=1,
        order_number=order["order_number"], supplier=order["supplier"], mobile="+۹۸۹۱۲۳۴۵۶۷۸۹",
        filename="supplier-order-SUP-TEST.xlsx", content=content,
    )
    assert result["send_status"] == "sent"
    assert result["recipient_masked"] == "0912***6789"
    assert len(transport) == 1 and "https://ai.neginpakhsh.com/warehouse-download/" in transport[0][1]
    token = transport[0][1].split("/warehouse-download/", 1)[1].splitlines()[0]
    downloaded = sms.consume_download(settings, token)
    assert downloaded["content"] == content
    assert load_workbook(BytesIO(downloaded["content"]))["گزارش"].cell(2, 4).value == "00123"
    for _ in range(4):
        assert sms.consume_download(settings, token) is not None
    assert sms.consume_download(settings, token) is None
    with service.warehouse_connection(settings) as conn:
        row = conn.execute("SELECT token_hash,content_sha256,status,downloaded_count FROM warehouse_sms_downloads").fetchone()
    assert token not in row["token_hash"] and len(row["token_hash"]) == 64
    assert row["status"] == "sent" and row["downloaded_count"] == 5


@pytest.mark.parametrize("mobile", ["", "912345", "091234567890", "0912\n3456789", "0019123456789"])
def test_invalid_mobile_never_sends(store, mobile):
    settings, transport = store
    order = _order(settings)
    with pytest.raises(service.WarehouseAssistantError):
        sms.send_document_link(
            settings, "tester", document_kind="supplier_order", document_id=1,
            order_number="SUP-TEST", supplier="supplier", mobile=mobile,
            filename="test.xlsx", content=sms.supplier_order_workbook(order),
        )
    assert transport == []


def test_routes_require_permission_confirmation_and_public_token(store, monkeypatch):
    settings, transport = store
    app = FastAPI()
    app.state.settings = settings
    async def identity(request: Request):
        request.state.username = "tester"
    app.dependency_overrides[require_session_user] = identity
    monkeypatch.setattr(routes, "_capabilities", lambda request, username: {"warehouse.order.draft"})
    monkeypatch.setattr(routes, "_is_admin", lambda request, username: False)
    app.include_router(routes.router)
    app.include_router(routes.public_router)
    with TestClient(app) as client:
        url = "/warehouse-assistant/api/supplier-orders/1/send-sms-link"
        assert client.post(url, json={"mobile": "09123456789", "confirmed": True}).status_code == 403
        blocked = client.post(url, headers={"X-Warehouse-Sms": "1", "Origin": "http://testserver"},
                              json={"mobile": "09123456789", "confirmed": True})
        assert blocked.status_code == 409
        approved = client.post("/warehouse-assistant/api/supplier-orders/1/approve")
        assert approved.status_code == 200 and approved.json()["is_approved"] is True
        in_transit = fulfillment.list_fulfillment_orders(settings)
        assert len(in_transit) == 1
        assert in_transit[0]["source_kind"] == "manual"
        assert in_transit[0]["order_number"] == "SUP-TEST"
        sent = client.post(url, headers={"X-Warehouse-Sms": "1", "Origin": "http://testserver"},
                           json={"mobile": "09123456789", "confirmed": True})
        assert sent.status_code == 200, sent.text
        token = transport[0][1].split("/warehouse-download/", 1)[1].splitlines()[0]
        landing = client.get(f"/warehouse-download/{token}")
        assert landing.status_code == 200
        assert landing.headers["content-type"].startswith("text/html")
        assert "SUP-TEST" in landing.text and "00123" in landing.text
        assert f'/warehouse-download/{token}/file' in landing.text
        download = client.get(f"/warehouse-download/{token}/file")
        assert download.status_code == 200
        assert download.headers["cache-control"] == "no-store, private"
        assert load_workbook(BytesIO(download.content))["گزارش"].cell(2, 4).value == "00123"
        assert download.headers["x-content-type-options"] == "nosniff"
        assert client.get("/warehouse-download/not-a-token").status_code == 404


def test_ui_has_explicit_sms_dialog_and_buttons():
    html = Path("app/static/warehouse-assistant.html").read_text(encoding="utf-8")
    js = Path("app/static/warehouse-assistant.js").read_text(encoding="utf-8")
    sms_js = Path("app/static/warehouse-order-sms.js").read_text(encoding="utf-8")
    assert 'id="orderSmsDialog"' in html and 'id="orderSmsMobile"' in html
    assert 'data-order-sms-kind="supplier-order"' in js
    assert 'data-order-sms-kind="automatic-preorder"' in js
    assert "تأیید سفارش" in js and "ارسال ایمیل و اکسل" in js and "ارسال پیامک و اکسل" in js
    assert "X-Warehouse-Sms" in sms_js and "confirmed:true" in sms_js
