from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import warehouse_supplier_portal as portal
from app import warehouse_assistant_service as service
from app.routes import warehouse_supplier_portal as routes
from app.routes.dependencies import require_session_user


@pytest.fixture
def store(tmp_path):
    settings = SimpleNamespace(sqlite_path=Path(tmp_path) / "main.db")
    service.init_warehouse_store(settings)
    with service.warehouse_connection(settings) as conn:
        conn.execute("""INSERT INTO warehouse_snapshots
            (id,source_filename,source_sheet,content_sha256,product_count,item_count,imported_by,imported_at)
            VALUES(1,'test','test','portal-test',2,2,'tester','now')""")
        for source_row, code, name, maker_code, barcode in (
            (1, "00123", "کالای اول", "SUP-01", "626000000001"),
            (2, "00456", "کالای دوم", "SUP-02", "626000000002"),
        ):
            conn.execute("""INSERT INTO warehouse_snapshot_items
                (snapshot_id,source_row,warehouse_code,warehouse_name,product_code,product_name,
                 conversion_rate,manufacturer,brand,manufacturer_product_code,barcode,buy_price)
                VALUES(1,?,'karaj','انبار کرج',?,?,12,'تأمین‌کننده نمونه','برند',?,?,70)""",
                (source_row, code, name, maker_code, barcode))
        conn.execute("""INSERT INTO supplier_orders
            (id,order_number,snapshot_id,warehouse_code,warehouse_name,supplier,total_quantity,
             estimated_value,created_by,created_at)
            VALUES(1,'SUP-TEST',1,'karaj','انبار کرج','تأمین‌کننده نمونه',36,2520,'buyer','2026-01-01')""")
        conn.executemany("""INSERT INTO supplier_order_lines
            (order_id,product_code,product_name,brand,conversion_rate,requested_quantity,
             order_quantity,cartons,buy_price,estimated_value,note)
            VALUES(1,?,?, 'برند',12,?,?,?,70,?, '')""", [
                ("00123", "کالای اول", 24, 24, 2, 1680),
                ("00456", "کالای دوم", 12, 12, 1, 840),
            ])
    service.transition_supplier_order(settings, "buyer", 1, "approve", include_all=True)
    return settings


def test_supplier_scope_proposal_approval_and_excel(store):
    account = portal.create_account(store, "buyer", username="supplier.one",
        password="Temporary-123", supplier_name="تأمین‌کننده نمونه", mobile="09120000000")
    portal.create_account(store, "buyer", username="supplier.two",
        password="Temporary-456", supplier_name="شرکت دیگر")
    assignment = portal.create_assignment(store, "buyer", document_kind="supplier_order",
        document_id=1, requested_delivery_date="1405/06/25")

    auth = portal.authenticate(store, "supplier.one", "Temporary-123")
    assert auth is not None and auth[1]["must_change_password"] is False
    token, profile = auth
    assert portal.session_profile(store, token)["supplier_name"] == "تأمین‌کننده نمونه"
    portal.change_password(store, account["id"], "Temporary-123", "A-new-password-123")
    profile = portal.session_profile(store, token)
    assert profile["must_change_password"] is False
    assert len(portal.list_assignments(store, supplier_key=profile["supplier_key"])) == 1
    other = portal.authenticate(store, "supplier.two", "Temporary-456")[1]
    assert portal.list_assignments(store, supplier_key=other["supplier_key"]) == []

    submitted = portal.save_response(store, profile, assignment["id"],
        expected_revision=assignment["revision"], proposed_delivery_date="1405/06/27",
        supplier_comment="یک قلم موجود نیست", submit=True, lines=[
            {"product_code": "00123", "proposed_cartons": 3, "line_status": "changed", "supplier_note": "سه کارتن قابل تأمین"},
            {"product_code": "00456", "proposed_cartons": 0, "line_status": "unavailable", "supplier_note": "ناموجود"},
        ])
    assert submitted["status"] == "submitted"
    # Original order remains untouched until staff accepts.
    original = service.get_supplier_order(store, 1, "", include_all=True)
    assert [(line["product_code"], line["cartons"]) for line in original["lines"]] == [("00123", 2), ("00456", 1)]

    accepted = portal.decide(store, "buyer", assignment["id"], decision="accept",
                             expected_revision=submitted["revision"], manager_comment="تأیید شد")
    assert accepted["status"] == "accepted"
    final = service.get_supplier_order(store, 1, "", include_all=True)
    assert final["is_approved"] is True
    assert [(line["product_code"], line["cartons"]) for line in final["lines"]] == [("00123", 3)]
    assert accepted["workflow_status"] == "awaiting_delivery"
    with service.warehouse_connection(store) as conn:
        projection_id = conn.execute(
            "SELECT id FROM warehouse_automatic_preorders WHERE source_supplier_order_id=1"
        ).fetchone()["id"]
    projected = service.get_automatic_preorder(store, projection_id)
    assert [(line["product_code"], line["cartons"]) for line in projected["lines"]] == [("00123", 3)]
    workbook = load_workbook(BytesIO(portal.response_workbook(store, assignment["id"], staff=True)))
    assert workbook["پاسخ سفارش"].cell(2, 8).value == "626000000001"


def test_supplier_cannot_omit_lines_or_reuse_stale_revision(store):
    account = portal.create_account(store, "buyer", username="supplier.one",
        password="Temporary-123", supplier_name="تأمین‌کننده نمونه")
    assignment = portal.create_assignment(store, "buyer", document_kind="supplier_order",
        document_id=1, requested_delivery_date="1405/06/25")
    profile = portal.authenticate(store, "supplier.one", "Temporary-123")[1]
    portal.change_password(store, account["id"], "Temporary-123", "A-new-password-123")
    profile["must_change_password"] = False
    one_line = [{"product_code": "00123", "proposed_cartons": 2,
                 "line_status": "confirmed", "supplier_note": ""}]
    with pytest.raises(service.WarehouseAssistantError, match="همه اقلام"):
        portal.save_response(store, profile, assignment["id"], expected_revision=0,
            proposed_delivery_date="1405/06/25", supplier_comment="", lines=one_line, submit=False)


def test_unchanged_overall_confirmation_is_final_without_manager_step(store):
    account = portal.create_account(store, "buyer", username="supplier.one",
        password="Temporary-123", supplier_name="تأمین‌کننده نمونه")
    assignment = portal.create_assignment(store, "buyer", document_kind="supplier_order",
        document_id=1, requested_delivery_date="1405/06/25")
    profile = portal.authenticate(store, "supplier.one", "Temporary-123")[1]
    portal.change_password(store, account["id"], "Temporary-123", "A-new-password-123")
    confirmed = portal.save_response(store, profile, assignment["id"],
        expected_revision=assignment["revision"], proposed_delivery_date="1405/06/25",
        supplier_comment="", submit=True, confirmation_mode="confirm", lines=[
            {"product_code": "00123", "proposed_cartons": 2, "line_status": "confirmed", "supplier_note": ""},
            {"product_code": "00456", "proposed_cartons": 1, "line_status": "confirmed", "supplier_note": ""},
        ])
    assert confirmed["status"] == "accepted"
    assert confirmed["workflow_status"] == "supplier_confirmed"
    # Internal approval precedes the supplier portal and remains active.
    assert service.get_supplier_order(store, 1, "", include_all=True)["is_approved"] is True


def test_overall_confirm_rejects_changed_date(store):
    portal.create_account(store, "buyer", username="supplier.one",
        password="Temporary-123", supplier_name="تأمین‌کننده نمونه")
    assignment = portal.create_assignment(store, "buyer", document_kind="supplier_order",
        document_id=1, requested_delivery_date="1405/06/25")
    profile = portal.authenticate(store, "supplier.one", "Temporary-123")[1]
    with pytest.raises(service.WarehouseAssistantError, match="درخواست تغییر"):
        portal.save_response(store, profile, assignment["id"], expected_revision=assignment["revision"],
            proposed_delivery_date="1405/06/26", supplier_comment="", submit=True,
            confirmation_mode="confirm", lines=[
                {"product_code": "00123", "proposed_cartons": 2, "line_status": "confirmed", "supplier_note": ""},
                {"product_code": "00456", "proposed_cartons": 1, "line_status": "confirmed", "supplier_note": ""},
            ])


def test_supplier_http_session_isolated_and_password_change_optional(store, monkeypatch):
    app = FastAPI()
    app.state.settings = store
    async def staff_identity(request: Request):
        request.state.username = "buyer"
    app.dependency_overrides[require_session_user] = staff_identity
    import app.routes.warehouse_assistant as warehouse_routes
    monkeypatch.setattr(warehouse_routes, "_capabilities", lambda request, username: {"warehouse.order.draft"})
    app.include_router(routes.staff_router)
    app.include_router(routes.public_router)
    with TestClient(app) as client:
        created = client.post("/warehouse-assistant/api/supplier-portal/accounts", json={
            "username": "supplier.one", "temporary_password": "Temporary-123",
            "supplier_name": "تأمین‌کننده نمونه", "mobile": "09120000000",
        })
        assert created.status_code == 201
        assigned = client.post("/warehouse-assistant/api/supplier-portal/orders", json={
            "document_kind": "supplier_order", "document_id": 1,
            "requested_delivery_date": "1405/06/25",
        })
        assert assigned.status_code == 201
        logged_in = client.post("/supplier-portal/api/login", headers={"Origin": "http://testserver"},
                                json={"username": "09120000000", "password": "Temporary-123"})
        assert logged_in.status_code == 200
        assert client.get("/supplier-portal/api/orders").status_code == 200
        changed = client.post("/supplier-portal/api/change-password", headers={"Origin": "http://testserver"},
            json={"current_password": "Temporary-123", "new_password": "A-new-password-123"})
        assert changed.status_code == 200
        listing = client.get("/supplier-portal/api/orders")
        assert listing.status_code == 200 and len(listing.json()["orders"]) == 1
        excel = client.get(f"/supplier-portal/api/orders/{assigned.json()['id']}/document.xlsx")
        assert excel.status_code == 200
        assert load_workbook(BytesIO(excel.content))["پاسخ سفارش"].cell(2, 6).value == "00123"


def test_portal_sms_contains_login_not_password(store, monkeypatch):
    from app import warehouse_order_sms as sms
    portal.create_account(store, "buyer", username="supplier.one", password="Temporary-123",
                          supplier_name="تأمین‌کننده نمونه", mobile="09120000000")
    assignment = portal.create_assignment(store, "buyer", document_kind="supplier_order",
        document_id=1, requested_delivery_date="1405/06/25")
    monkeypatch.setattr(sms, "_provider_config", lambda: {"username": "x"})
    messages = []
    monkeypatch.setattr(sms, "_send_provider", lambda config, mobile, message: messages.append(message) or "42")
    result = portal.send_invitation(store, "buyer", assignment["id"], "")
    assert result["send_status"] == "sent"
    assert portal.get_assignment(store, assignment["id"], staff=True)["workflow_status"] == "awaiting_supplier"
    assert "/supplier-portal" in messages[0] and "supplier.one" in messages[0]
    assert "Temporary-123" not in messages[0]


def test_portal_assignment_requires_internal_approval(store):
    service.transition_supplier_order(store, "buyer", 1, "revoke_approval", include_all=True)
    with pytest.raises(service.WarehouseAssistantError, match="ابتدا سفارش را"):
        portal.create_assignment(
            store, "buyer", document_kind="supplier_order", document_id=1,
            requested_delivery_date="1405/06/25",
        )


def test_supplier_portal_ui_disables_full_confirm_and_exposes_line_filters():
    script = (Path(__file__).parents[1] / "app" / "static" / "supplier-portal.js").read_text(
        encoding="utf-8"
    )
    for element_id in (
        "confirmWholeOrder",
        "requestOrderChange",
        "makerCodeFilter",
        "barcodeFilter",
        "packFilter",
        "brandFilter",
    ):
        assert element_id in script
    assert "confirm.disabled = changed" in script
    assert "request.disabled = !changed" in script
    assert "data-maker" in script
    assert "data-barcode" in script
    assert "data-pack" in script
    assert "data-brand" in script

    stylesheet = (Path(__file__).parents[1] / "app" / "static" / "supplier-portal.css").read_text(
        encoding="utf-8"
    )
    assert "height:calc(100dvh - 68px)" in stylesheet
    assert "#responseForm{flex:1" in stylesheet
    assert ".compact-order-table{flex:1" in stylesheet
    assert ".conversation-section.is-collapsed" in stylesheet
    assert "#orderDetail .sp-meta{display:flex" in stylesheet
    assert ".compact-order-table{min-height:320px" in stylesheet
    assert "#responseForm{display:flex;flex:1;min-height:0" in stylesheet
    assert ".compact-order-table{flex:1;min-height:0!important" in stylesheet
    assert ".conversation-section:not(.is-collapsed){position:fixed" in stylesheet
    assert "@media(min-width:1280px)" in stylesheet
    assert "#orderDetail .sp-meta{display:none}" in stylesheet
    assert "grid-template-rows:32px minmax(0,1fr) 34px" in stylesheet
