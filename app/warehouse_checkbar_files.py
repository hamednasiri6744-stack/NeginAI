"""Read delivery-note attachments as ordered, bounded AI inputs. Never execute workbook content."""
import base64
import json
import math
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile
from xml.etree.ElementTree import iterparse

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, coordinate_to_tuple
from pypdf import PdfReader, PdfWriter
import xlrd

from app.warehouse_assistant_service import WarehouseAssistantError

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TOTAL_BYTES = 30 * 1024 * 1024
MAX_PAGES = 10
MAX_ROWS = 1500
MAX_COLUMNS = 64
MAX_CELLS = 20000
MAX_TEXT = 200000


def file_kind(data):
    if not data or len(data) > MAX_FILE_BYTES:
        raise WarehouseAssistantError('هر فایل باید غیرخالی و حداکثر ۱۰ مگابایت باشد.')
    if data.startswith(b'%PDF-'):
        return 'pdf'
    if data.startswith(b'PK\x03\x04'):
        return 'xlsx'
    if data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        return 'xls'
    from app.warehouse_checkbar_image import image_type
    try:
        return image_type(data)
    except WarehouseAssistantError as exc:
        raise WarehouseAssistantError('فقط عکس JPG، PNG، WebP، فایل PDF یا اکسل XLSX و XLS انتخاب کنید.') from exc


def cell_value(value, number_format=''):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(value):
            raise ValueError('Non-finite cell')
        # Excel identifiers often store 12 with display format 0000. Preserve 0012.
        if re.fullmatch(r'0{2,30}', number_format) and value >= 0 and int(value) == value:
            return str(int(value)).zfill(len(number_format))
    if isinstance(value, str) and len(value) > 2000:
        raise WarehouseAssistantError('متن یک سلول بیش از حد طولانی است؛ فایل حواله را کوچک‌تر کنید.')
    return value


def sheet_input(name, rows):
    text = json.dumps({'sheet': name, 'cells': rows}, ensure_ascii=False, allow_nan=False)
    if len(text) > MAX_TEXT:
        raise WarehouseAssistantError('حجم اطلاعات اکسل زیاد است؛ فقط برگه‌های حوالهٔ موردنظر را انتخاب کنید.')
    return {'type': 'input_text', 'text': text}


def dimensions(rows, columns):
    if rows > MAX_ROWS or columns > MAX_COLUMNS or rows * columns > MAX_CELLS:
        raise WarehouseAssistantError('محدودهٔ اکسل بیش از حد بزرگ است؛ حداکثر ۱۵۰۰ ردیف، ۶۴ ستون و ۲۰٬۰۰۰ خانه در هر برگه.')


def xlsx_pages(data):
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > MAX_TOTAL_BYTES:
            raise WarehouseAssistantError('فایل اکسل پس از باز شدن بیش از حد بزرگ است.')
        names = {e.filename for e in entries}
        if 'xl/workbook.xml' not in names:
            raise ValueError('Not an XLSX workbook')
        if any('vbaproject' in name.lower() or name.startswith(('xl/embeddings/', 'xl/media/')) for name in names):
            raise WarehouseAssistantError('اکسلِ دارای تصویر، شیء پیوست یا ماکرو را به PDF تبدیل کنید تا محتوا جا نماند.')
        for name in names:
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                with archive.open(name) as stream:
                    count = 0
                    for _, element in iterparse(stream, events=('end',)):
                        if element.tag.rsplit('}', 1)[-1] == 'c':
                            row, column = coordinate_to_tuple(element.attrib['r'])
                            count += 1
                            if row > MAX_ROWS or column > MAX_COLUMNS or count > MAX_CELLS:
                                raise WarehouseAssistantError('محدودهٔ سلول‌های اکسل بیش از حد مجاز است؛ فایل را کوچک‌تر کنید.')
                        element.clear()
    book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        cached = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        if len(book.sheetnames) > MAX_PAGES:
            raise WarehouseAssistantError('مجموع صفحه‌های PDF، عکس‌ها و برگه‌های اکسل حداکثر ۱۰ مورد است.')
        if len(book.worksheets) != len(book.sheetnames):
            raise WarehouseAssistantError('اکسل دارای برگهٔ غیرجدولی است؛ فایل را به PDF تبدیل کنید.')
        pages = []
        for sheet in book.worksheets:
            dimensions(sheet.max_row or 0, sheet.max_column or 0)
            rows = []
            # Iterate bounded dimensions explicitly; do not trust a missing or understated dimension tag.
            for raw, saved in zip(sheet.iter_rows(max_row=MAX_ROWS+1, max_col=MAX_COLUMNS+1),
                                  cached[sheet.title].iter_rows(max_row=MAX_ROWS+1, max_col=MAX_COLUMNS+1)):
                cells = {}
                for cell, cache in zip(raw, saved):
                    value = cell.value
                    if value is None:
                        continue
                    if cell.row > MAX_ROWS or cell.column > MAX_COLUMNS:
                        dimensions(cell.row, cell.column)
                    if cell.data_type == 'f':
                        if cache.value is None:
                            raise WarehouseAssistantError('فرمول اکسل نتیجهٔ ذخیره‌شده ندارد؛ در Excel محاسبه و ذخیره کنید یا PDF بفرستید.')
                        value = cache.value
                    if cell.data_type == 'e' or cache.data_type == 'e':
                        raise WarehouseAssistantError('اکسل دارای سلول خطادار است؛ خطا را اصلاح کنید یا PDF بفرستید.')
                    cells[cell.coordinate] = cell_value(value, cell.number_format)
                if cells:
                    rows.append(cells)
            if rows:
                pages.append((f'برگه {sheet.title}', sheet_input(sheet.title, rows)))
        if not pages:
            raise WarehouseAssistantError('اکسل هیچ سلول پُری ندارد.')
        return pages
    finally:
        book.close()
        if cached:
            cached.close()


def xls_pages(data):
    book = xlrd.open_workbook(file_contents=data, on_demand=True, formatting_info=True)
    try:
        if book.nsheets > MAX_PAGES:
            raise WarehouseAssistantError('مجموع صفحه‌های PDF، عکس‌ها و برگه‌های اکسل حداکثر ۱۰ مورد است.')
        pages = []
        for sheet in book.sheets():
            dimensions(sheet.nrows, sheet.ncols)
            rows = []
            for r in range(sheet.nrows):
                cells = {}
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        raise WarehouseAssistantError('اکسل دارای سلول خطادار است؛ خطا را اصلاح کنید یا PDF بفرستید.')
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    cells[f'{get_column_letter(c+1)}{r+1}'] = cell_value(value, fmt)
                if cells:
                    rows.append(cells)
            if rows:
                pages.append((f'برگه {sheet.name}', sheet_input(sheet.name, rows)))
        if not pages:
            raise WarehouseAssistantError('اکسل هیچ سلول پُری ندارد.')
        return pages
    finally:
        book.release_resources()


def pdf_pages(data):
    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise WarehouseAssistantError('PDF رمزدار است؛ نسخهٔ بدون رمز را انتخاب کنید.')
    if not 1 <= len(reader.pages) <= MAX_PAGES:
        raise WarehouseAssistantError('PDF باید از ۱ تا ۱۰ صفحه داشته باشد.')
    pages = []
    total = 0
    for index, page in enumerate(reader.pages, 1):
        writer = PdfWriter()
        writer.add_page(page)
        output = BytesIO()
        writer.write(output)
        encoded = output.getvalue()
        total += len(encoded)
        if total > MAX_TOTAL_BYTES:
            raise WarehouseAssistantError('حجم صفحه‌های آماده‌شدهٔ PDF بیش از حد مجاز است.')
        pages.append((f'صفحه {index}', {'type': 'input_file', 'filename': f'page-{index}.pdf',
            'file_data': 'data:application/pdf;base64,' + base64.b64encode(encoded).decode('ascii')}))
    return pages


def prepare_files(contents, filenames):
    if not 1 <= len(contents) <= MAX_PAGES or len(contents) != len(filenames):
        raise WarehouseAssistantError('از ۱ تا ۱۰ فایل حواله انتخاب کنید.')
    if sum(map(len, contents)) > MAX_TOTAL_BYTES:
        raise WarehouseAssistantError('حجم مجموع فایل‌ها باید حداکثر ۳۰ مگابایت باشد.')
    pages = []
    for data, filename in zip(contents, filenames):
        label = PurePosixPath((filename or 'حواله').replace('\\', '/')).name[:150]
        try:
            kind = file_kind(data)
            if kind == 'pdf':
                parts = pdf_pages(data)
            elif kind == 'xlsx':
                parts = xlsx_pages(data)
            elif kind == 'xls':
                parts = xls_pages(data)
            else:
                parts = [('عکس', {'type': 'input_image', 'detail': 'high',
                    'image_url': f'data:{kind};base64,' + base64.b64encode(data).decode('ascii')})]
            pages.extend((f'{label} · {part}', content) for part, content in parts)
            if len(pages) > MAX_PAGES:
                raise WarehouseAssistantError('مجموع صفحه‌های PDF، عکس‌ها و برگه‌های اکسل حداکثر ۱۰ مورد است.')
        except WarehouseAssistantError:
            raise
        except Exception as exc:
            raise WarehouseAssistantError(f'فایل «{label}» قابل خواندن نیست؛ فایل سالم و بدون رمز انتخاب کنید.') from exc
    if sum(len(json.dumps(content, ensure_ascii=False)) for _, content in pages) > 42 * 1024 * 1024:
        raise WarehouseAssistantError('حجم فایل‌های آماده‌شده بیش از حد مجاز است.')
    return pages
