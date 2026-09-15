"""User-confirmed SMS delivery of short-lived, immutable order workbooks."""
from __future__ import annotations

import hashlib
import json
import re
import secrets
from datetime import datetime, timedelta, timezone
from io import BytesIO
from urllib.parse import quote

import httpx
from dotenv import dotenv_values
from openpyxl import load_workbook

from app.config import ENV_PATH
from app.warehouse_assistant_service import (
    WarehouseAssistantError,
    _now,
    init_warehouse_store,
    warehouse_connection,
)
from app.warehouse_order_excel import build_supplier_workbook
from app.warehouse_order_receipts import ensure_cycle_allowed


PUBLIC_BASE_URL = "https://ai.neginpakhsh.com"
LINK_LIFETIME_HOURS = 24
MAX_DOWNLOADS = 5
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_DIGIT_TRANSLATION = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


class SmsDeliveryFailure(Exception):
    def __init__(self, status: str, message: str):
        super().__init__(message)
        self.status = status


def normalize_mobile(value: str) -> str:
    if not isinstance(value, str):
        raise WarehouseAssistantError("شماره موبایل گیرنده معتبر نیست.")
    mobile = value.translate(_DIGIT_TRANSLATION).strip().replace(" ", "").replace("-", "")
    if mobile.startswith("+98"):
        mobile = "0" + mobile[3:]
    elif mobile.startswith("0098"):
        mobile = "0" + mobile[4:]
    elif mobile.startswith("98"):
        mobile = "0" + mobile[2:]
    if not re.fullmatch(r"09\d{9}", mobile):
        raise WarehouseAssistantError("شماره موبایل باید یک شمارهٔ معتبر ایران مانند 09123456789 باشد.")
    return mobile


def supplier_order_workbook(order: dict[str, object]) -> bytes:
    columns = [
        "شماره سفارش", "تأمین‌کننده", "انبار مقصد", "کد کالا",
        "کد کالای تولیدکننده", "بارکد", "نام کالا", "برند", "کارتن",
        "تعداد سفارش", "ضریب کارتن", "قیمت تولیدکننده", "قیمت مصرف‌کننده",
        "قیمت حدودی", "ارزش تخمینی", "توضیحات قلم",
    ]
    rows = [[
        order["order_number"], order["supplier"], order["warehouse_name"],
        line["product_code"], str(line.get("manufacturer_product_code") or ""),
        str(line.get("barcode") or ""), line["product_name"], line["brand"],
        line["cartons"], line["order_quantity"], line["conversion_rate"],
        line["manufacturer_price"], line["consumer_price"], line["approximate_price"],
        line["estimated_value"], line["note"],
    ] for line in order["lines"]]
    return build_supplier_workbook(
        columns, rows, f"سفارش خرید {order['order_number']} برای {order['supplier']}", order["lines"]
    )


def _init_store(settings) -> None:
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS warehouse_sms_downloads (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              token_hash TEXT NOT NULL UNIQUE,
              document_kind TEXT NOT NULL CHECK(document_kind IN ('supplier_order','automatic_preorder')),
              document_id INTEGER NOT NULL,
              order_number TEXT NOT NULL,
              supplier TEXT NOT NULL,
              recipient TEXT NOT NULL,
              filename TEXT NOT NULL,
              content BLOB NOT NULL,
              content_sha256 TEXT NOT NULL,
              status TEXT NOT NULL CHECK(status IN ('sending','sent','failed','unknown','revoked')),
              provider_message_id TEXT,
              requested_by TEXT NOT NULL,
              created_at TEXT NOT NULL,
              expires_at TEXT NOT NULL,
              completed_at TEXT,
              downloaded_count INTEGER NOT NULL DEFAULT 0,
              max_downloads INTEGER NOT NULL DEFAULT 5,
              error TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS idx_warehouse_sms_download_expiry
              ON warehouse_sms_downloads(expires_at, status);
        """)


def _provider_config() -> dict[str, str]:
    try:
        values = dotenv_values(ENV_PATH, interpolate=False, encoding="utf-8-sig")
    except (OSError, UnicodeError):
        raise WarehouseAssistantError("خواندن تنظیمات پیامک انجام نشد.") from None
    config = {
        "username": str(values.get("ASANAK_USERNAME") or "").strip(),
        "password": str(values.get("ASANAK_PASSWORD") or ""),
        "source": str(values.get("ASANAK_SOURCE") or "").strip(),
        "send_url": str(values.get("ASANAK_SEND_URL") or "https://panel.asanak.com/webservice/v2rest/sendsms").strip(),
    }
    if not config["username"] or not config["password"] or not config["source"]:
        raise WarehouseAssistantError("اتصال پیامک فعال یا کامل تنظیم نشده است؛ مدیر سیستم بررسی کند.")
    if config["send_url"] not in {
        "https://panel.asanak.com/webservice/v2rest/sendsms",
        "https://sms.asanak.ir/webservice/v2rest/sendsms",
    }:
        raise WarehouseAssistantError("آدرس امن سرویس پیامک معتبر نیست؛ مدیر سیستم بررسی کند.")
    return config


def _send_provider(config: dict[str, str], mobile: str, message: str) -> str:
    payload = {
        "username": config["username"], "password": config["password"],
        "source": config["source"], "destination": mobile, "message": message,
    }
    try:
        response = httpx.post(
            config["send_url"], json=payload,
            headers={"Accept": "application/json"}, timeout=20.0,
        )
    except (httpx.ConnectError, httpx.ConnectTimeout):
        raise SmsDeliveryFailure("failed", "اتصال به سرویس پیامک برقرار نشد؛ دوباره تلاش کنید.") from None
    except httpx.TransportError:
        raise SmsDeliveryFailure(
            "unknown", "نتیجهٔ ارسال پیامک نامشخص است؛ پیش از تلاش دوباره گزارش سرویس بررسی شود."
        ) from None
    try:
        body = response.json()
    except (ValueError, json.JSONDecodeError):
        body = {}
    meta = body.get("meta") if isinstance(body, dict) else None
    data = body.get("data") if isinstance(body, dict) else None
    if response.is_success and isinstance(meta, dict) and int(meta.get("status", 0)) == 200 and data:
        return str(data[0])[:120]
    provider_status = str(meta.get("status", "")) if isinstance(meta, dict) else ""
    provider_message = str(meta.get("message", "")) if isinstance(meta, dict) else ""
    provider_message = re.sub(r"[\x00-\x1f\x7f]+", " ", provider_message).strip()[:120]
    if "can not send link" in provider_message.casefold():
        provider_message = "خط ارسال فعلی مجوز ارسال لینک ندارد"
    detail = " · ".join(part for part in (provider_status, provider_message) if part)
    suffix = f" (پاسخ آسانک: {detail})" if detail else ""
    raise SmsDeliveryFailure(
        "failed", f"سرویس آسانک پیامک را نپذیرفت؛ اعتبار حساب، خط و شماره بررسی شود.{suffix}"
    )


def send_document_link(
    settings, username: str, *, document_kind: str, document_id: int,
    order_number: str, supplier: str, mobile: str, filename: str, content: bytes,
) -> dict[str, object]:
    recipient = normalize_mobile(mobile)
    if document_kind not in {"supplier_order", "automatic_preorder"}:
        raise WarehouseAssistantError("نوع سند برای پیامک معتبر نیست.")
    if not content or len(content) > MAX_WORKBOOK_BYTES:
        raise WarehouseAssistantError("حجم فایل اکسل سفارش برای ارسال لینک معتبر نیست.")
    config = _provider_config()
    _init_store(settings)
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode("ascii")).hexdigest()
    created = datetime.now(timezone.utc).replace(microsecond=0)
    expires = created + timedelta(hours=LINK_LIFETIME_HOURS)
    safe_filename = re.sub(r"[^A-Za-z0-9_.-]", "-", filename)[:180]
    with warehouse_connection(settings) as conn:
        conn.execute("BEGIN IMMEDIATE")
        if document_kind == "supplier_order":
            order = conn.execute("SELECT warehouse_code FROM supplier_orders WHERE id=?", (document_id,)).fetchone()
            codes = [row[0] for row in conn.execute("SELECT product_code FROM supplier_order_lines WHERE order_id=?", (document_id,))]
        else:
            order = conn.execute("SELECT warehouse_code FROM warehouse_automatic_preorders WHERE id=?", (document_id,)).fetchone()
            codes = [row[0] for row in conn.execute("SELECT product_code FROM warehouse_automatic_preorder_lines WHERE preorder_id=?", (document_id,))]
        if order:
            ensure_cycle_allowed(conn, order["warehouse_code"], codes)
        conn.execute(
            "UPDATE warehouse_sms_downloads SET content=X'' WHERE expires_at<=? AND length(content)>0",
            (created.isoformat(),),
        )
        pending = conn.execute("""SELECT 1 FROM warehouse_sms_downloads
            WHERE document_kind=? AND document_id=? AND recipient=?
              AND status IN ('sending','unknown') AND expires_at>? LIMIT 1""",
            (document_kind, int(document_id), recipient, created.isoformat())).fetchone()
        if pending:
            raise WarehouseAssistantError(
                "نتیجهٔ ارسال قبلی این سفارش هنوز مشخص نیست؛ برای جلوگیری از پیامک تکراری دوباره ارسال نشد."
            )
        attempt_id = conn.execute("""INSERT INTO warehouse_sms_downloads
            (token_hash,document_kind,document_id,order_number,supplier,recipient,filename,
             content,content_sha256,status,requested_by,created_at,expires_at,max_downloads)
            VALUES(?,?,?,?,?,?,?,?,?,'sending',?,?,?,?)""",
            (token_hash, document_kind, int(document_id), str(order_number)[:100], str(supplier)[:200],
             recipient, safe_filename, content, hashlib.sha256(content).hexdigest(),
             username[:100], created.isoformat(), expires.isoformat(), MAX_DOWNLOADS)).lastrowid
    url = f"{PUBLIC_BASE_URL}/warehouse-download/{token}"
    message = (
        f"نگین پخش\nفایل اکسل سفارش {order_number}:\n{url}\n"
        f"اعتبار لینک: {LINK_LIFETIME_HOURS} ساعت"
    )
    outcome, provider_id, error = "sent", "", ""
    try:
        provider_id = _send_provider(config, recipient, message)
    except SmsDeliveryFailure as exc:
        outcome, error = exc.status, str(exc)
    except Exception:
        outcome, error = "unknown", "نتیجهٔ ارسال پیامک نیازمند بررسی مدیر سیستم است."
    with warehouse_connection(settings) as conn:
        conn.execute("""UPDATE warehouse_sms_downloads SET status=?,provider_message_id=?,
            completed_at=?,error=? WHERE id=?""", (outcome, provider_id or None, _now(), error, attempt_id))
    if outcome != "sent":
        raise WarehouseAssistantError(error)
    return {
        "send_status": "sent", "recipient_masked": recipient[:4] + "***" + recipient[-4:],
        "order_number": order_number, "expires_at": expires.isoformat(),
        "download_url": url, "external_delivery_performed": True,
    }


def _valid_download_row(settings, token: str, *, consume: bool):
    if not re.fullmatch(r"[A-Za-z0-9_-]{40,60}", token or ""):
        return None
    _init_store(settings)
    token_hash = hashlib.sha256(token.encode("ascii")).hexdigest()
    now = datetime.now(timezone.utc).replace(microsecond=0)
    with warehouse_connection(settings) as conn:
        if consume:
            conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT * FROM warehouse_sms_downloads WHERE token_hash=?", (token_hash,)).fetchone()
        if row is None or row["status"] not in ("sent", "unknown"):
            return None
        try:
            expires = datetime.fromisoformat(row["expires_at"])
        except (TypeError, ValueError):
            return None
        if expires <= now or row["downloaded_count"] >= row["max_downloads"]:
            return None
        content = bytes(row["content"])
        if not content:
            return None
        if consume:
            conn.execute(
                "UPDATE warehouse_sms_downloads SET downloaded_count=downloaded_count+1 WHERE id=?",
                (row["id"],),
            )
        return dict(row, content=content)


def preview_download(settings, token: str) -> dict[str, object] | None:
    """Return safe display metadata without consuming a download allowance."""
    row = _valid_download_row(settings, token, consume=False)
    if row is None:
        return None
    try:
        workbook = load_workbook(BytesIO(row["content"]), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        lines = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not values or not values[3]:
                continue
            lines.append({
                "product_code": str(values[3] or ""),
                "manufacturer_product_code": str(values[4] or ""),
                "barcode": str(values[5] or ""),
                "product_name": str(values[6] or ""),
                "brand": str(values[7] or ""),
                "cartons": values[8] or 0,
                "quantity": values[9] or 0,
            })
        workbook.close()
    except Exception:
        return None
    return {
        "order_number": row["order_number"],
        "supplier": row["supplier"],
        "filename": row["filename"],
        "expires_at": row["expires_at"],
        "remaining_downloads": max(0, int(row["max_downloads"]) - int(row["downloaded_count"])),
        "lines": lines,
    }


def consume_download(settings, token: str) -> dict[str, object] | None:
    row = _valid_download_row(settings, token, consume=True)
    if row is None:
        return None
    return {"content": row["content"], "filename": row["filename"]}
