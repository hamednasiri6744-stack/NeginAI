from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _excel_safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def build_report_workbook(
    columns: list[str],
    rows: list[list[Any]],
    title: str = "گزارش نگین AI",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "گزارش"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)
    for column_index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column_index, value=_excel_safe(name))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_excel_safe(value))
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = "#,##0.####"

    if columns:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    sheet.row_dimensions[1].height = 25
    for column_index, name in enumerate(columns, start=1):
        values = [str(name), *(str(row[column_index - 1]) for row in rows[:500] if len(row) >= column_index)]
        width = min(45, max(11, max((len(value) for value in values), default=10) + 3))
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    info = workbook.create_sheet("اطلاعات")
    info.sheet_view.rightToLeft = True
    info.append(["عنوان گزارش", _excel_safe(title)])
    info.append(["تعداد ردیف", len(rows)])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 55

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
