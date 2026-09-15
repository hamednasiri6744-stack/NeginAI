"""Versioned inspection documents and confirmed local receipts; no stock/ERP writes."""
import hashlib
import json
import math
import re
from copy import deepcopy
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from app.excel_service import _excel_safe
from app.warehouse_assistant_service import (
    WAREHOUSES, WarehouseAssistantError, init_warehouse_store, warehouse_connection, _now,
)
from app.warehouse_fulfillment import supply_rows


class Selection(BaseModel):
    warehouse: str = Field(pattern=r'^(karaj|tehran|gilan)$')
    supplier: str = Field(min_length=1, max_length=200)
    order_ids: list[int] = Field(default_factory=list, max_length=100)
    remaining_order_id: int | None = Field(default=None, gt=0)


class InspectionLine(BaseModel):
    preorder_id: int | None = Field(default=None, gt=0)
    product_code: str = Field(min_length=1, max_length=80)
    cartons: int | None = Field(default=None, ge=0, le=1000000)
    units: float | None = Field(default=None, ge=0, le=1000000000, allow_inf_nan=False)
    manufacturer_price_new: float | None = Field(default=None, ge=0, le=1e15, allow_inf_nan=False)
    consumer_price_new: float | None = Field(default=None, ge=0, le=1e15, allow_inf_nan=False)
    tax: float | None = Field(default=None, ge=0, le=100, allow_inf_nan=False)


class CheckbarAllocation(BaseModel):
    model_config = {'extra': 'forbid'}
    source_row: int = Field(ge=1,le=500)
    preorder_id: int = Field(gt=0)
    quantity: float = Field(ge=0,le=1e9,allow_inf_nan=False)


class ConfirmationFields(BaseModel):
    confirm_receipt: bool = False
    skip_order_matching: bool = False
    allocations: list[CheckbarAllocation] | None = Field(default=None,max_length=5000)
    allocation_revision: str = Field(default='',max_length=64)
    accept_unallocated: bool = False


class InspectionRequest(Selection, ConfirmationFields):
    worksheet_workflow: bool = False
    split_by_brand: bool = False
    expected_token: str = Field(pattern=r'^[a-f0-9]{64}$')
    request_id: str = Field(min_length=1, max_length=80)
    metadata: dict[str, str] = Field(default_factory=dict)
    lines: list[InspectionLine] = Field(min_length=1, max_length=500)


class MutationRequest(BaseModel):
    model_config = {'extra': 'forbid'}
    expected_revision: int = Field(ge=0)
    request_id: str = Field(min_length=1, max_length=80)


class EditRequest(MutationRequest, ConfirmationFields):
    metadata: dict[str, str] = Field(default_factory=dict)
    lines: list[InspectionLine] = Field(min_length=1, max_length=500)


class MatchingPreviewRequest(Selection):
    document_id: int | None = Field(default=None,gt=0)
    expected_revision: int = Field(default=0,ge=0)
    expected_token: str = ''
    request_id: str = 'preview'
    metadata: dict[str,str] = Field(default_factory=dict)
    lines: list[InspectionLine] = Field(min_length=1,max_length=500)


META_FIELDS = {'date', 'reference_no', 'waybill', 'bill_of_lading', 'driver', 'plate', 'phone', 'receiver', 'note'}


def _hash(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False).encode()).hexdigest()


def _source_token(source):
    """Guard product definitions, not live stock or the import that delivered them.

    Outstanding quantities are checked by the allocation revision at confirmation.
    They are display context during worksheet preparation and initial saving.
    Keep catalogue order because image suggestions refer to its row positions.
    """
    fields = ('preorder_id','product_code','product_name','brand','conversion_rate',
              'manufacturer_product_code','barcode','barcode2','barcode_list','group_level3',
              'manufacturer_price','consumer_price')
    return _hash(dict(warehouse=source['warehouse'],supplier=source['supplier'],
                      order_ids=source['order_ids'],remaining_order_id=source.get('remaining_order_id'),
                      lines=[{k:row.get(k) for k in fields} for row in source['lines']],
                      catalog=[{k:row.get(k) for k in fields} for row in source['catalog']]))


def _compatible_payload(payload):
    if not payload.get('skip_order_matching'):
        payload = {k:v for k,v in payload.items() if k != 'skip_order_matching'}
    if not payload.get('worksheet_workflow'):
        payload = {k:v for k,v in payload.items() if k != 'worksheet_workflow'}
    if payload.get('remaining_order_id') is None:
        payload = {k:v for k,v in payload.items() if k != 'remaining_order_id'}
    # Omitted/default-off splitting must retain historical idempotency hashes.
    if not payload.get('split_by_brand'):
        payload = {k:v for k,v in payload.items() if k != 'split_by_brand'}
    if not payload.get('confirm_receipt'):
        return {k:v for k,v in payload.items() if k not in ('confirm_receipt','allocations','allocation_revision','accept_unallocated')}
    return payload


def _scope(warehouse):
    if warehouse not in WAREHOUSES:
        raise WarehouseAssistantError('انبار انتخاب‌شده معتبر نیست.')


def _orders(conn, warehouse):
    # Resolve identities from the sent order's snapshot, never from another warehouse.
    rows = supply_rows(conn, warehouse=warehouse)
    orders = {}
    for row in rows:
        remaining = max(0, row['order_quantity'] - row['received_qty'] - row['closed_qty'])
        if remaining <= 0:
            continue
        order = orders.setdefault(row['preorder_id'], dict(
            id=row['preorder_id'], preorder_number=row['preorder_number'],
            supplier=row['supplier'], sent_at=row['sent_at'], lines=[], remaining_cartons=0,
        ))
        line = dict(row, remaining_qty=remaining)
        order['lines'].append(line)
        order['remaining_cartons'] += remaining / max(1, row['conversion_rate'])
    return list(orders.values())


def _suppliers(conn, warehouse, orders):
    # Active supply is independent of automatic-order enablement or old deliveries.
    return [r[0] for r in conn.execute('''SELECT DISTINCT supplier
        FROM warehouse_supply_scope WHERE warehouse_code=? AND enabled=1
        AND TRIM(COALESCE(supplier,''))<>'' ORDER BY supplier COLLATE NOCASE''', (warehouse,))]


def context(settings, warehouse, supplier=''):
    _scope(warehouse)
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.execute('BEGIN')
        orders = _orders(conn, warehouse)
        return {'suppliers': _suppliers(conn, warehouse, orders),
                'orders': [o for o in orders if o['supplier'] == supplier]}


def _prepare(conn, warehouse, supplier, order_ids, remaining_order_id=None):
    _scope(warehouse)
    if len(order_ids) > 100 or len(order_ids) != len(set(order_ids)):
        raise WarehouseAssistantError('سفارش‌های انتخابی نباید تکراری یا بیش از حد مجاز باشند.')
    available_orders = _orders(conn, warehouse)
    if supplier not in _suppliers(conn, warehouse, available_orders):
        raise WarehouseAssistantError('تأمین‌کننده در تأمین فعال این انبار نیست؛ اطلاعات را بازخوانی کنید.')
    orders = [o for o in available_orders if o['supplier'] == supplier and o['id'] in order_ids]
    if {o['id'] for o in orders} != set(order_ids):
        raise WarehouseAssistantError('سفارش انتخابی تحویل شده یا متعلق به این انبار و تأمین‌کننده نیست؛ بازخوانی کنید.')
    if remaining_order_id is not None and order_ids != [remaining_order_id]:
        raise WarehouseAssistantError('چک‌بار مانده باید مربوط به یک سفارش انتخابی باشد.')
    marks = ','.join('?' for _ in order_ids)
    identity_rows = conn.execute(f'''SELECT l.preorder_id,l.product_code,l.brand,
        l.manufacturer_price,l.consumer_price,i.manufacturer_product_code,i.barcode,i.barcode2,i.barcode_list,i.group_level3,o.snapshot_id
        FROM warehouse_automatic_preorder_lines l
        JOIN warehouse_automatic_preorders o ON o.id=l.preorder_id
        LEFT JOIN warehouse_snapshot_items i ON i.snapshot_id=o.snapshot_id
          AND i.warehouse_code=l.warehouse_code AND i.product_code=l.product_code
        WHERE l.preorder_id IN ({marks})''', order_ids).fetchall() if order_ids else []
    identities = {(r['preorder_id'], r['product_code']): dict(r) for r in identity_rows}
    lines = [dict(line, **{k: v for k, v in identities[(line['preorder_id'], line['product_code'])].items()
                          if k not in ('preorder_id', 'product_code')})
             for order in orders for line in order['lines']]
    snapshot = conn.execute('SELECT id,imported_at FROM warehouse_snapshots ORDER BY id DESC LIMIT 1').fetchone()
    catalog = []
    if snapshot:
        catalog = [dict(r, preorder_id=None, snapshot_id=snapshot['id'], preorder_number='',
                        order_quantity=None, received_qty=None, remaining_qty=None)
                   for r in conn.execute('''SELECT product_code,product_name,brand,conversion_rate,
                       manufacturer_product_code,barcode,barcode2,barcode_list,group_level3,manufacturer_price,consumer_price,stock
                       FROM warehouse_snapshot_items WHERE snapshot_id=? AND warehouse_code=?
                       AND manufacturer=? COLLATE NOCASE ORDER BY product_name,product_code''',
                       (snapshot['id'], warehouse, supplier))]
    aggregate = {}
    catalog_by_code = {line['product_code']:line for line in catalog}
    for order in available_orders:
        if order['supplier'] != supplier:
            continue
        if remaining_order_id is not None and order['id'] != remaining_order_id:
            continue
        for line in order['lines']:
            code=line['product_code']
            if code not in catalog_by_code:
                identity=conn.execute('''SELECT l.brand,l.manufacturer_price,l.consumer_price,
                    i.manufacturer_product_code,i.barcode,i.barcode2,i.barcode_list,i.group_level3,o.snapshot_id
                    FROM warehouse_automatic_preorder_lines l JOIN warehouse_automatic_preorders o ON o.id=l.preorder_id
                    LEFT JOIN warehouse_snapshot_items i ON i.snapshot_id=o.snapshot_id AND i.warehouse_code=o.warehouse_code AND i.product_code=l.product_code
                    WHERE l.preorder_id=? AND l.product_code=?''',(order['id'],code)).fetchone()
                catalog_by_code[code]=dict(line,**dict(identity))
                catalog_by_code[code].update(preorder_id=None,preorder_number='',order_quantity=None,received_qty=None,remaining_qty=None)
                catalog.append(catalog_by_code[code])
            item=aggregate.setdefault(code,dict(catalog_by_code[code],order_quantity=0,received_qty=0,remaining_qty=0,
                                               aggregate_open_order=True,preorder_number='مجموع سفارش‌های در راه'))
            for key in ('order_quantity','received_qty','remaining_qty'):
                item[key]+=line[key]
    data = dict(warehouse=warehouse, warehouse_name=WAREHOUSES[warehouse]['name'], supplier=supplier,
                order_ids=sorted(order_ids), orders=orders, lines=lines, catalog=catalog,
                aggregate_open_lines=list(aggregate.values()),
                outstanding_orders=[dict(id=o['id'],number=o['preorder_number'],remaining_cartons=o['remaining_cartons'])
                                    for o in available_orders if o['supplier']==supplier and (remaining_order_id is None or o['id']==remaining_order_id)],
                catalog_snapshot_at=snapshot['imported_at'] if snapshot else None)
    if remaining_order_id is not None:
        data['remaining_order_id'] = remaining_order_id
    data['expected_token'] = _source_token(data)
    return data


def prepare(settings, warehouse, supplier, order_ids, remaining_order_id=None):
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.execute('BEGIN')
        return _prepare(conn, warehouse, supplier, order_ids, remaining_order_id)


def _document(row):
    return dict(json.loads(row['document_json']), id=row['id'], number=f"CB-{row['id']:06d}",
                created_at=row['created_at'], created_by=row['created_by'], revision=0)


def _current_document(conn, document_id, include_deleted=False):
    row=conn.execute('SELECT * FROM warehouse_checkbars WHERE id=?',(document_id,)).fetchone()
    if not row:
        raise WarehouseAssistantError('چک بار پیدا نشد.')
    doc=_document(row)
    revision=conn.execute('SELECT * FROM warehouse_checkbar_revisions WHERE document_id=? ORDER BY revision DESC LIMIT 1',
                          (document_id,)).fetchone()
    if revision:
        doc.update(json.loads(revision['document_json']))
        doc.update(revision=revision['revision'],updated_at=revision['created_at'],updated_by=revision['created_by'],
                   deleted=revision['operation']=='delete')
    if doc.get('deleted') and not include_deleted:
        raise WarehouseAssistantError('این چک بار حذف شده است؛ سوابق را بازخوانی کنید.')
    if not doc.get('receipt_confirmed') and conn.execute(
            'SELECT 1 FROM warehouse_checkbar_confirmations WHERE document_id=? AND active=1',(document_id,)).fetchone():
        prior=conn.execute('SELECT payload_json FROM warehouse_checkbar_transfer_history WHERE document_id=? ORDER BY rowid DESC LIMIT 1',(document_id,)).fetchone()
        if prior:
            plan=json.loads(prior[0]).get('order_matching')
            if plan:
                doc.update(receipt_confirmed=True,order_matching=plan)
    return doc


def _validate_metadata(metadata):
    if set(metadata)-META_FIELDS or any(len(v)>500 for v in metadata.values()):
        raise WarehouseAssistantError('مشخصات بار نامعتبر یا بیش از حد طولانی است.')
    reference = metadata.get('reference_no', '').strip().translate(
        str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789'))
    if not re.fullmatch(r'[0-9]{1,10}', reference) or not 1 <= int(reference) <= 2147483647:
        raise WarehouseAssistantError('شمارهٔ سند عطف تأمین‌کننده الزامی است؛ عدد صحیح بین ۱ و ۲۱۴۷۴۸۳۶۴۷ وارد کنید.')
    metadata['reference_no'] = reference


def _inspection_value(base, value):
    rate=float(base['conversion_rate'])
    if not math.isfinite(rate) or rate<1:
        raise WarehouseAssistantError('ضریب کارتن کالا معتبر نیست.')
    actual=None if value['cartons'] is None and value['units'] is None else (
        (value['cartons'] or 0)*rate+(value['units'] or 0))
    if actual is not None and (not math.isfinite(actual) or actual>1e9):
        raise WarehouseAssistantError('تعداد واقعی بار بیش از حد مجاز است.')
    return dict(base, **{k:v for k,v in value.items() if k not in ('preorder_id','product_code')},actual_qty=actual)


def _confirm_draft(conn, doc, payload, document_id=None):
    if not payload.get('confirm_receipt'):
        if doc.get('receipt_confirmed'):
            raise WarehouseAssistantError('اصلاح این چک‌بار نیازمند تأیید مجدد شمارش و تطبیق سفارش‌ها است.')
        return
    if doc.get('worksheet_workflow') and not doc.get('worksheet_approved_at'):
        raise WarehouseAssistantError('ابتدا برگه چک‌بار را ذخیره و برای شمارش انبار تأیید کنید.')
    from app.warehouse_order_receipts import matching
    if any(line.get('actual_qty') is None for line in doc['lines']):
        raise WarehouseAssistantError('تعدادهای خالی را تکمیل یا با تأیید از چک‌بار حذف کنید.')
    if document_id:
        transfer=conn.execute('SELECT status FROM warehouse_checkbar_transfers WHERE document_id=?',(document_id,)).fetchone()
        if transfer and transfer['status'] in ('sent','pending'):
            raise WarehouseAssistantError('چک‌بار منتقل‌شده یا در حال انتقال قابل تأیید مجدد دریافت نیست؛ ابتدا وضعیت رسید را تعیین تکلیف کنید.')
    if payload.get('skip_order_matching'):
        if payload.get('allocations') or payload.get('allocation_revision'):
            raise WarehouseAssistantError('در حالت بدون تطبیق، تخصیص سفارش نباید ارسال شود.')
        rows = [dict(source_row=i, product_code=line['product_code'], product_name=line.get('product_name',''),
                     actual_qty=line['actual_qty'], unallocated_qty=line['actual_qty'], orders=[])
                for i,line in enumerate(doc['lines'],1) if line['actual_qty']>0]
        doc.update(receipt_confirmed=True,order_matching=dict(version=1,skipped=True,revision='',
                   allocations=[],rows=rows,unallocated_qty=sum(row['actual_qty'] for row in rows)))
        return
    if not payload.get('allocation_revision') or payload.get('allocations') is None:
        raise WarehouseAssistantError('تطبیق سفارش‌ها را تأیید کنید یا گزینهٔ تطبیق را خاموش کنید.')
    plan=matching(conn,doc,payload['allocations'],payload['allocation_revision'],document_id)
    if plan['unallocated_qty']>0 and not payload.get('accept_unallocated'):
        raise WarehouseAssistantError('دریافت اضافه بر سفارش یا بدون سفارش را صریحاً تأیید کنید.')
    doc.update(receipt_confirmed=True,order_matching=plan)


def matching_preview(settings, payload):
    payload=MatchingPreviewRequest.model_validate(payload).model_dump()
    _validate_metadata(payload['metadata'])
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.execute('BEGIN')
        document_id=payload['document_id']
        if document_id:
            doc=_current_document(conn,document_id)
            if doc['revision']!=payload['expected_revision'] or doc['warehouse']!=payload['warehouse'] or doc['supplier']!=payload['supplier']:
                raise WarehouseAssistantError('چک‌بار تغییر کرده است؛ نسخهٔ جدید را باز کنید.')
            transfer=conn.execute('SELECT status FROM warehouse_checkbar_transfers WHERE document_id=?',(document_id,)).fetchone()
            if transfer and transfer['status'] in ('sent','pending'):
                raise WarehouseAssistantError('ابتدا وضعیت رسید منتقل‌شده را تعیین تکلیف کنید.')
            doc=_edited_draft(conn,doc,payload)
        else:
            doc=_new_draft(conn,payload)
        if any(line.get('actual_qty') is None for line in doc['lines']):
            raise WarehouseAssistantError('تعدادهای خالی را تکمیل یا با تأیید حذف کنید.')
        from app.warehouse_order_receipts import matching
        return matching(conn,doc,exclude_document_id=document_id)


def _new_draft(conn, payload):
    source = _prepare(conn, payload['warehouse'], payload['supplier'], payload['order_ids'], payload.get('remaining_order_id'))
    # Accept an unchanged legacy form during deployment as well as the stable token.
    legacy_token = _hash({k:v for k,v in source.items() if k != 'expected_token'})
    if payload['expected_token'] not in (source['expected_token'], legacy_token):
        raise WarehouseAssistantError('مشخصات مبنای کالا یا فهرست اقلام تغییر کرده است؛ اطلاعات فرم نیاز به بازخوانی دارد. تعدادهای واردشده را حفظ کنید.')
    by_key = {(l['preorder_id'], l['product_code']): l for l in source['lines']}
    manual = {l['product_code']: l for l in source['catalog']}
    manual.update({l['product_code']:l for l in source['aggregate_open_lines']})
    source_codes = {l['product_code'] for l in source['lines']}
    seen, lines = set(), []
    for value in payload['lines']:
        key = (value['preorder_id'], value['product_code'])
        if key in seen:
            raise WarehouseAssistantError('یک ردیف سفارش بیش از یک بار وارد شده است.')
        seen.add(key)
        if value['preorder_id'] is None:
            base = manual.get(value['product_code'])
            if value['product_code'] in source_codes:
                raise WarehouseAssistantError('این کالا در سفارش‌های انتخابی هست؛ همان ردیف را ویرایش کنید.')
        else:
            base = by_key.get(key)
        if base is None:
            raise WarehouseAssistantError('کالا متعلق به سفارش یا تأمین‌کنندهٔ انتخابی نیست.')
        lines.append(_inspection_value(base,value))
    document = dict(warehouse=source['warehouse'], warehouse_name=source['warehouse_name'],
                    supplier=source['supplier'], order_ids=source['order_ids'],
                    metadata=payload['metadata'], lines=lines, source_token=source['expected_token'])
    if payload.get('worksheet_workflow'):
        document['worksheet_workflow'] = True
    if source.get('remaining_order_id') is not None:
        document['remaining_order_id'] = source['remaining_order_id']
    return document


def _brand_documents(document):
    """Partition the already validated whole-shipment plan, without reallocating."""
    groups, code_brands = {}, {}
    for index,line in enumerate(document['lines'],1):
        brand = str(line.get('brand') or '').strip()
        if not brand:
            raise WarehouseAssistantError(f"برند کالای {line['product_code']} مشخص نیست؛ اطلاعات برند را اصلاح کنید یا تفکیک برند را خاموش کنید.")
        code = line['product_code']
        if code in code_brands and code_brands[code] != brand:
            raise WarehouseAssistantError(f'برند کالای {code} در اطلاعات سفارش‌ها یکسان نیست؛ ابتدا اطلاعات برند را اصلاح کنید.')
        code_brands[code] = brand
        groups.setdefault(brand,[]).append(index)
    documents=[]
    for brand,indices in groups.items():
        doc=deepcopy(document)
        doc.update(brand=brand,lines=[doc['lines'][i-1] for i in indices])
        if doc.get('receipt_confirmed'):
            remap={old:new for new,old in enumerate(indices,1)}
            plan=doc['order_matching']
            for field in ('allocations','rows'):
                plan[field]=[dict(row,source_row=remap[row['source_row']])
                             for row in plan[field] if row['source_row'] in remap]
            plan['unallocated_qty']=sum(row['unallocated_qty'] for row in plan['rows'])
        documents.append(doc)
    return documents


def issue(settings, username, payload):
    # Validate again at the service boundary, including direct callers.
    payload = InspectionRequest.model_validate(payload).model_dump()
    metadata = payload['metadata']
    _validate_metadata(metadata)
    request_hash = _hash(_compatible_payload(payload))
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.execute('BEGIN IMMEDIATE')
        old = conn.execute('SELECT * FROM warehouse_checkbars WHERE created_by=? AND request_id=?',
                           (username, payload['request_id'])).fetchone()
        if old:
            if old['request_hash'] != request_hash:
                raise WarehouseAssistantError('این درخواست قبلاً با اطلاعات دیگری ثبت شده است؛ فرم را دوباره باز کنید.')
            return _current_document(conn,old['id'])
        document = _new_draft(conn,payload)
        _confirm_draft(conn,document,payload)
        children=_brand_documents(document) if payload['split_by_brand'] else [document]
        ids=[]
        for index,child in enumerate(children):
            key=payload['request_id'] if index==0 else 'brand:'+_hash([username,payload['request_id'],index])
            cursor = conn.execute('''INSERT INTO warehouse_checkbars
                (created_at,created_by,warehouse_code,supplier,request_id,request_hash,document_json)
                VALUES(?,?,?,?,?,?,?)''', (_now(), username, child['warehouse'], child['supplier'],
                    key, request_hash, json.dumps(child, ensure_ascii=False, allow_nan=False)))
            ids.append(cursor.lastrowid)
            if child.get('receipt_confirmed'):
                from app.warehouse_order_receipts import confirm_checkbar
                confirm_checkbar(conn,cursor.lastrowid,child,username)
        if payload['split_by_brand']:
            batch=[dict(id=id,number=f'CB-{id:06d}',brand=child['brand']) for id,child in zip(ids,children)]
            for id,child in zip(ids,children):
                child['brand_batch']=batch
                conn.execute('UPDATE warehouse_checkbars SET document_json=? WHERE id=?',
                             (json.dumps(child,ensure_ascii=False,allow_nan=False),id))
        return _current_document(conn,ids[0])


def documents(settings, include_deleted=False):
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        result = [dict(id=r['id'], number=f"CB-{r['id']:06d}", created_at=r['created_at'],
                     receipt_number=json.loads(r['receipt_result'] or '{}').get('VocherNo'),
                     receipt_state=r['receipt_state'],
                     brand=json.loads(r['document_json']).get('brand',''),
                     supplier=r['supplier'], warehouse_name=WAREHOUSES[r['warehouse_code']]['name'],
                     deleted=bool(r['deleted']))
                for r in conn.execute('''SELECT c.*,
                    COALESCE(CASE WHEN t.status='sent' THEN t.result_json END,(SELECT h.result_json FROM warehouse_checkbar_transfer_history h WHERE h.document_id=c.id ORDER BY h.rowid DESC LIMIT 1)) AS receipt_result,
                    CASE WHEN t.status IN ('sent','pending') THEN t.status WHEN EXISTS(SELECT 1 FROM warehouse_checkbar_transfer_history h WHERE h.document_id=c.id) THEN 'deleted' ELSE 'not_sent' END AS receipt_state,
                    EXISTS
                    (SELECT 1 FROM warehouse_checkbar_revisions r WHERE r.document_id=c.id AND r.operation='delete') AS deleted
                    FROM warehouse_checkbars c LEFT JOIN warehouse_checkbar_transfers t ON t.document_id=c.id WHERE ? OR NOT EXISTS
                    (SELECT 1 FROM warehouse_checkbar_revisions r WHERE r.document_id=c.id AND r.operation='delete')
                    ORDER BY c.id DESC LIMIT 100''', (include_deleted,))]

        for row in result:
            doc = _current_document(conn, row["id"], include_deleted=True)
            row.update(worksheet_workflow=bool(doc.get("worksheet_workflow")),
                       worksheet_approved_at=doc.get("worksheet_approved_at"),
                       receipt_confirmed=bool(doc.get("receipt_confirmed")))
        return result


def document_history(settings, document_id):
    """Read retained snapshots, including deleted documents, without restoring live state."""
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.execute('BEGIN')
        current = _current_document(conn, document_id, include_deleted=True)
        original = _document(conn.execute('SELECT * FROM warehouse_checkbars WHERE id=?', (document_id,)).fetchone())
        versions = [dict(revision=0, operation='issue', created_at=original['created_at'],
                         created_by=original['created_by'], document=original)]
        for row in conn.execute('SELECT * FROM warehouse_checkbar_revisions WHERE document_id=? ORDER BY revision', (document_id,)):
            doc = dict(json.loads(row['document_json']), revision=row['revision'],
                       deleted=row['operation']=='delete', updated_at=row['created_at'], updated_by=row['created_by'])
            versions.append(dict(revision=row['revision'], operation=row['operation'],
                                 created_at=row['created_at'], created_by=row['created_by'], document=doc))
        return dict(document=current, versions=versions)


def historical_document(settings, document_id, revision):
    for version in document_history(settings, document_id)['versions']:
        if version['revision'] == revision:
            return version['document']
    raise WarehouseAssistantError('نسخهٔ چک بار پیدا نشد.')


def get_document(settings, document_id):
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.execute('BEGIN')
        return _current_document(conn,document_id)


def edit_context(settings,document_id):
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.execute('BEGIN IMMEDIATE')
        from app.warehouse_receipt_lifecycle import inspect
        inspect(conn,settings,document_id,strict=True)
        doc=_current_document(conn,document_id)
        # Old documents remain editable after delivery or supplier deactivation.
        try:
            source=_prepare(conn,doc['warehouse'],doc['supplier'],[])
        except WarehouseAssistantError:
            source={'catalog':[]}
        if doc.get('brand_batch'):
            source['catalog']=[line for line in source['catalog'] if str(line.get('brand') or '').strip()==doc['brand']]
        return dict(document=doc,catalog=source['catalog'])


def _edited_draft(conn, doc, payload):
    existing={(line['preorder_id'],line['product_code']):line for line in doc['lines']}
    # Only additions use current catalog. Existing identities and carton rates stay historical.
    catalog={}
    if any((v['preorder_id'],v['product_code']) not in existing for v in payload['lines']):
        catalog={l['product_code']:l for l in _prepare(conn,doc['warehouse'],doc['supplier'],[])['catalog']}
    source_codes={l['product_code'] for l in doc['lines'] if l['preorder_id'] is not None}
    seen,lines=set(),[]
    for value in payload['lines']:
        key=(value['preorder_id'],value['product_code'])
        if key in seen:
            raise WarehouseAssistantError('یک ردیف بیش از یک بار وارد شده است.')
        seen.add(key)
        base=existing.get(key)
        if base is None and value['preorder_id'] is None and value['product_code'] not in source_codes:
            base=catalog.get(value['product_code'])
        if base is None:
            raise WarehouseAssistantError('کالا متعلق به چک بار یا تأمین‌کنندهٔ همین انبار نیست.')
        lines.append(_inspection_value(base,value))
    if any(line['actual_qty'] is None for line in lines) and not (doc.get('worksheet_workflow') and not doc.get('receipt_confirmed')):
        raise WarehouseAssistantError('تعداد ردیف‌های خالی را تکمیل یا پیش از ذخیره حذف کنید.')
    if doc.get('brand_batch') and any(str(line.get('brand') or '').strip()!=doc['brand'] for line in lines):
        raise WarehouseAssistantError('این چک‌بار برای یک برند صادر شده است؛ کالای برند دیگر را در چک‌بار جدا ثبت کنید.')
    doc.update(metadata=payload['metadata'],lines=lines)
    return doc


def _mutate_document(settings,username,document_id,payload,operation):
    payload=(EditRequest if operation=='edit' else MutationRequest).model_validate(payload).model_dump()
    if operation=='edit':
        _validate_metadata(payload['metadata'])
    digest=_hash(dict(operation=operation,payload=_compatible_payload(payload)))
    init_warehouse_store(settings)
    with warehouse_connection(settings) as conn:
        conn.execute('BEGIN IMMEDIATE')
        previous=conn.execute('''SELECT * FROM warehouse_checkbar_revisions
            WHERE document_id=? AND created_by=? AND request_id=?''',
            (document_id,username,payload['request_id'])).fetchone()
        doc=_current_document(conn,document_id,include_deleted=True)
        if previous:
            if previous['request_hash']!=digest:
                raise WarehouseAssistantError('این درخواست قبلاً با اطلاعات دیگری ثبت شده است؛ فرم را دوباره باز کنید.')
            if operation=='delete':
                return dict(id=document_id,number=doc['number'],deleted=True)
            return _current_document(conn,document_id)
        if doc.get('deleted'):
            raise WarehouseAssistantError('این چک بار قبلاً حذف شده است.')
        if doc['revision']!=payload['expected_revision']:
            raise WarehouseAssistantError('چک بار توسط کاربر دیگری تغییر کرده است؛ فرم را ببندید و نسخهٔ جدید را باز کنید.')
        from app.warehouse_receipt_lifecycle import inspect
        inspect(conn,settings,document_id,strict=True)
        transfer = conn.execute('SELECT status FROM warehouse_checkbar_transfers WHERE document_id=?', (document_id,)).fetchone()
        doc=_current_document(conn,document_id)
        if transfer and transfer['status'] == 'pending':
            raise WarehouseAssistantError('ابتدا نتیجهٔ انتقال به ورانگر را با «پیگیری انتقال» مشخص کنید؛ تغییر سند تا آن زمان ممکن نیست.')
        if operation=='edit':
            doc=_edited_draft(conn,doc,payload)
            _confirm_draft(conn,doc,payload,document_id)
        elif operation=='approve_worksheet':
            if not doc.get('worksheet_workflow') or doc.get('receipt_confirmed'):
                raise WarehouseAssistantError('این سند در مرحله تأیید برگه شمارش نیست.')
            if doc.get('worksheet_approved_at'):
                return doc
            doc.update(worksheet_approved_at=_now(),worksheet_approved_by=username,
                       worksheet_approved_revision=doc['revision']+1)
        conn.execute('''INSERT INTO warehouse_checkbar_revisions
            (document_id,revision,operation,created_at,created_by,request_id,request_hash,document_json)
            VALUES(?,?,?,?,?,?,?,?)''',(document_id,doc['revision']+1,'edit' if operation=='approve_worksheet' else operation,_now(),username,
            payload['request_id'],digest,json.dumps(doc,ensure_ascii=False,allow_nan=False)))
        if doc.get('receipt_confirmed'):
            if operation=='edit':
                from app.warehouse_order_receipts import confirm_checkbar
                confirm_checkbar(conn,document_id,dict(doc,revision=doc['revision']+1),username)
            elif not transfer or transfer['status']!='sent':
                conn.execute('UPDATE warehouse_checkbar_confirmations SET active=0,revision=? WHERE document_id=?',(doc['revision']+1,document_id))
        return dict(id=document_id,number=doc['number'],deleted=True) if operation=='delete' else _current_document(conn,document_id)


def edit_document(settings,username,document_id,payload):
    return _mutate_document(settings,username,document_id,payload,'edit')


def delete_document(settings,username,document_id,payload):
    return _mutate_document(settings,username,document_id,payload,'delete')


def approve_worksheet(settings,username,document_id,payload):
    """Approve the retained worksheet revision without confirming any receipt."""
    return _mutate_document(settings,username,document_id,payload,'approve_worksheet')


def workbook(document):
    """First fourteen columns follow the legacy check-bar data contract; no VBA."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'چک بار'
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    meta = document['metadata']
    for cell, value in {'B1': 'تاریخ: '+meta.get('date', ''), 'B2': 'تأمین‌کننده: '+document['supplier'],
        'D1':'شماره حواله: '+meta.get('waybill',''), 'D2':'شماره بارنامه: '+meta.get('bill_of_lading',''),
        'G1':'راننده: '+meta.get('driver',''), 'G2':'پلاک: '+meta.get('plate',''),
        'K1':'تلفن: '+meta.get('phone',''), 'K2':'تحویل‌گیرنده: '+meta.get('receiver',''),
        'B3':'توضیحات: '+meta.get('note',''), 'D3':'انبار: '+document['warehouse_name'],
        'G3':'شماره سند عطف: '+meta.get('reference_no',''),
        'K3':'برند: '+document.get('brand',''),
        'O1':document['number'], 'O2':('دریافت چک‌بار تأیید شده؛ انتقال ورانگر جداگانه است.' if document.get('receipt_confirmed') else 'صدور چک بار، ثبت دریافت نیست.')}.items():
        sheet[cell] = _excel_safe(value)
    columns = ['ردیف','کد کالا','کد کالای تأمین‌کننده','بارکد','نام کالا','مبنا',
               'مصرف کننده','تولید کننده','تعداد به کارتن','تعداد جزء','تعداد کل','تولید جدید','مصرف جدید',
               'سفارش مبدأ','سفارش (عدد)','دریافت قبلی (عدد)','مانده (عدد)','مغایرت با مانده (عدد)']
    for col, title in enumerate(columns, 1):
        cell = sheet.cell(4, col, title)
        cell.fill = PatternFill('solid', fgColor='145C49')
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    widths = [7,14,18,21,38,9,16,16,12,12,14,16,16,25,14,14,14,17]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for n, line in enumerate(document['lines'], 5):
        values = [n-4, str(line['product_code']), str(line.get('manufacturer_product_code') or ''),
                  str(line.get('barcode') or ''), line['product_name'], line['conversion_rate'],
                  line['consumer_price'], line['manufacturer_price'], line['cartons'], line['units'],
                  None, line['manufacturer_price_new'], line['consumer_price_new'],
                  line['preorder_number'] or 'افزودهٔ دستی', line['order_quantity'], line['received_qty'],
                  line['remaining_qty'], None]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(n, col, _excel_safe(value))
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
            if isinstance(value, (int, float)):
                cell.number_format = '#,##0.####'
            if col in (9,10,12,13):
                cell.fill = PatternFill('solid', fgColor='FFF8DE')
        sheet.cell(n,11,f'=IF(COUNT(I{n}:J{n})=0,"",I{n}*F{n}+J{n})')
        sheet.cell(n,18,f'=IF(OR(K{n}="",Q{n}=""),"",K{n}-Q{n})')
        sheet.row_dimensions[n].height = 30
    sheet.row_dimensions[4].height = 32
    sheet.freeze_panes = 'F5'
    sheet.auto_filter.ref = f'A4:R{sheet.max_row}'
    sheet.print_title_rows = '1:4'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    info = book.create_sheet('مراجع سفارش')
    info.sheet_view.rightToLeft = True
    info.append(['چک بار', 'ثبت‌کننده', 'زمان صدور', 'سفارش مبدأ', 'شناسه سفارش', 'کد کالا', 'شناسه اطلاعات کالا'])
    for line in document['lines']:
        info.append([_excel_safe(v) for v in [document['number'], document['created_by'], document['created_at'],
            line['preorder_number'], line['preorder_id'], str(line['product_code']), line['snapshot_id']]])
    for i in range(1,8):
        info.column_dimensions[get_column_letter(i)].width = 25
    info.freeze_panes = 'A2'
    output = BytesIO()
    book.save(output)
    return output.getvalue()
